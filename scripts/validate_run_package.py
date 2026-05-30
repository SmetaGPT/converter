from __future__ import annotations

import argparse
import json
import sys
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
SRC = ROOT / "src"
if str(SRC) not in sys.path:
    sys.path.insert(0, str(SRC))

from openpyxl import load_workbook

from doc_converter.schema_validation import validate_json_file, validate_payload


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(description="Validate converter run package JSON artifacts against local schemas.")
    parser.add_argument("run_dir", type=Path, help="Path to converter run directory")
    args = parser.parse_args(argv)

    run_dir = args.run_dir.expanduser().resolve()
    legacy_agent_run_metadata = _validate_run_json(run_dir / "run.json")
    validate_json_file(run_dir / "summary.json", "summary.v1.schema.json")
    validate_json_file(run_dir / "queue-state.json", "queue-state.v1.schema.json")
    validate_json_file(run_dir / "processed-documents-catalog.json", "processed-documents-catalog.v1.schema.json")
    _validate_workbook(run_dir / "processed-documents-catalog.xlsx")
    _validate_jsonl(run_dir / "telemetry.jsonl", "log.v1.schema.json")
    _validate_jsonl(run_dir / "manifest.jsonl", "manifest.v1.schema.json")
    _validate_jsonl(run_dir / "review-required.jsonl", "review-required.v1.schema.json")

    documents_dir = run_dir / "documents"
    for document_path in sorted(documents_dir.glob("*/document.v1.json")):
        validate_json_file(document_path, "document.v1.schema.json")
        formula_recognition_path = document_path.parent / "formula-recognition.jsonl"
        if formula_recognition_path.exists():
            _validate_jsonl(formula_recognition_path, "formula-recognition.v1.schema.json")

    chunks_path = run_dir / "chunks.v1.jsonl"
    if chunks_path.exists():
        _validate_jsonl(chunks_path, "chunks.v1.schema.json")

    print(
        json.dumps(
            {
                "run_dir": str(run_dir),
                "status": "ok",
                "legacy_agent_run_metadata": legacy_agent_run_metadata,
            },
            ensure_ascii=False,
        )
    )
    return 0


def _validate_run_json(path: Path) -> bool:
    payload = json.loads(path.read_text(encoding="utf-8"))
    if "agent_run_metadata" in payload:
        validate_payload(payload, "run.v1.schema.json")
        return False

    legacy_payload = dict(payload)
    legacy_payload["agent_run_metadata"] = {
        "agent_id": "legacy",
        "agent_version": str(payload.get("converter_version") or "legacy"),
        "task_id": str(payload.get("run_id") or "legacy"),
    }
    validate_payload(legacy_payload, "run.v1.schema.json")
    return True


def _validate_jsonl(path: Path, schema_filename: str) -> None:
    if not path.exists():
        return
    for line in path.read_text(encoding="utf-8").splitlines():
        if not line.strip():
            continue
        validate_payload(json.loads(line), schema_filename)


def _validate_workbook(path: Path) -> None:
    workbook = load_workbook(path, read_only=True)
    try:
        if "Документы" not in workbook.sheetnames:
            raise ValueError(f"Workbook does not contain expected worksheet: {path}")
    finally:
        workbook.close()


if __name__ == "__main__":
    raise SystemExit(main())