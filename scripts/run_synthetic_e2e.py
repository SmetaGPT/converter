# pyright: reportMissingTypeStubs=false

from __future__ import annotations

import argparse
import json
import shutil
import sys
from pathlib import Path

from docx import Document

ROOT = Path(__file__).resolve().parents[1]
SRC = ROOT / "src"
if str(SRC) not in sys.path:
    sys.path.insert(0, str(SRC))

from doc_converter.chunking import build_chunks_from_document
from doc_converter.config import ConverterConfig
from doc_converter.runner import run_convert_folder
from doc_converter.schema_validation import validate_json_file, validate_payload


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(description="Run a synthetic end-to-end converter smoke on local generated documents.")
    parser.add_argument("--input", type=Path, default=ROOT / "runs" / "synthetic-e2e-input")
    parser.add_argument("--output", type=Path, default=ROOT / "runs" / "synthetic-e2e-output")
    parser.add_argument("--clean", action="store_true")
    args = parser.parse_args(argv)

    if args.clean:
        _remove_if_exists(args.input)
        _remove_if_exists(args.output)

    args.input.mkdir(parents=True, exist_ok=True)
    args.output.mkdir(parents=True, exist_ok=True)
    _build_synthetic_input(args.input)

    result = run_convert_folder(ConverterConfig(input_dir=args.input, output_dir=args.output))
    run_dir = result.run_dir
    _validate_run_dir(run_dir)
    chunks_path = run_dir / "chunks.v1.jsonl"
    _build_chunks(run_dir, chunks_path)
    _validate_run_dir(run_dir)

    payload = json.loads((run_dir / "summary.json").read_text(encoding="utf-8"))
    print(json.dumps({"run_dir": str(run_dir), "status": result.status, "summary": payload}, ensure_ascii=False))
    return 0 if result.status == "success" else 1


def _build_synthetic_input(input_dir: Path) -> None:
    document = Document()
    document.add_heading("Синтетический документ", level=1)
    document.add_paragraph("Это synthetic e2e smoke для Windows Document Converter.")
    document.save(str(input_dir / "synthetic.docx"))


def _build_chunks(run_dir: Path, chunks_path: Path) -> None:
    records: list[dict[str, object]] = []
    for document_path in sorted((run_dir / "documents").glob("*/document.v1.json")):
        payload = validate_json_file(document_path, "document.v1.schema.json")
        for chunk in build_chunks_from_document(payload):
            validate_payload(chunk, "chunks.v1.schema.json")
            records.append(chunk)
    with chunks_path.open("w", encoding="utf-8", newline="\n") as handle:
        for record in records:
            handle.write(json.dumps(record, ensure_ascii=False, separators=(",", ":")) + "\n")


def _validate_run_dir(run_dir: Path) -> None:
    validate_json_file(run_dir / "run.json", "run.v1.schema.json")
    validate_json_file(run_dir / "summary.json", "summary.v1.schema.json")
    validate_json_file(run_dir / "queue-state.json", "queue-state.v1.schema.json")
    _validate_jsonl(run_dir / "manifest.jsonl", "manifest.v1.schema.json")
    _validate_jsonl(run_dir / "review-required.jsonl", "review-required.v1.schema.json")
    if (run_dir / "chunks.v1.jsonl").exists():
        _validate_jsonl(run_dir / "chunks.v1.jsonl", "chunks.v1.schema.json")


def _validate_jsonl(path: Path, schema_filename: str) -> None:
    for line in path.read_text(encoding="utf-8").splitlines():
        if not line.strip():
            continue
        validate_payload(json.loads(line), schema_filename)


def _remove_if_exists(path: Path) -> None:
    if path.exists():
        shutil.rmtree(path)


if __name__ == "__main__":
    raise SystemExit(main())