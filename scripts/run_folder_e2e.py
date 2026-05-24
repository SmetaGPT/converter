# pyright: reportMissingTypeStubs=false

from __future__ import annotations

import argparse
import json
import shutil
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

from doc_converter.chunking import build_chunks_from_document
from doc_converter.config import ConverterConfig, ConverterOptions
from doc_converter.runner import run_convert_folder
from doc_converter.schema_validation import validate_json_file, validate_payload


ROOT = Path(__file__).resolve().parents[1]


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(description="Run end-to-end converter validation on an existing input folder.")
    parser.add_argument("input", type=Path, help="Existing folder with DOCX/PDF/XLSX files to convert.")
    parser.add_argument("--output", type=Path, default=ROOT / "runs" / "folder-e2e", help="Output folder for run packages.")
    parser.add_argument("--ocr-languages", default="rus,eng", help="Comma- or plus-separated OCR language codes.")
    parser.add_argument("--clean", action="store_true", help="Remove the output folder before running.")
    parser.add_argument("--include-originals", action="store_true", help="Copy source files into canonical document packages.")
    args = parser.parse_args(argv)

    input_dir = args.input.expanduser().resolve()
    output_dir = args.output.expanduser().resolve()
    if not input_dir.is_dir():
        raise FileNotFoundError(f"Input folder does not exist: {input_dir}")

    if args.clean:
        _remove_if_exists(output_dir)
    output_dir.mkdir(parents=True, exist_ok=True)

    options = ConverterOptions(
        ocr_languages=_parse_ocr_languages(args.ocr_languages),
        include_originals=args.include_originals,
    )
    result = run_convert_folder(ConverterConfig(input_dir=input_dir, output_dir=output_dir, options=options))
    run_dir = result.run_dir
    summary = validate_json_file(run_dir / "summary.json", "summary.v1.schema.json")
    _validate_run_dir(run_dir)
    chunks = _build_chunks(run_dir, run_dir / "chunks.v1.jsonl")
    _validate_run_dir(run_dir)

    passed = summary["failed_files"] == 0 and summary["status"] in {"success", "partial_success"}
    report = {
        "input_dir": str(input_dir),
        "run_dir": str(run_dir),
        "status": summary["status"],
        "passed": passed,
        "discovered_files": summary["discovered_files"],
        "supported_files": summary["supported_files"],
        "processed_files": summary["processed_files"],
        "partial_files": summary["partial_files"],
        "failed_files": summary["failed_files"],
        "review_required_files": summary["review_required_files"],
        "duplicate_groups": summary["duplicate_groups"],
        "chunks": chunks,
        "partial_reasons": summary["partial_reasons"],
        "failed_reasons": summary["failed_reasons"],
        "review_required_reasons": summary["review_required_reasons"],
    }
    print(json.dumps(report, ensure_ascii=False, indent=2))
    return 0 if passed else 1


def _parse_ocr_languages(value: str) -> tuple[str, ...]:
    normalized = value.replace("+", ",")
    languages = tuple(part.strip() for part in normalized.split(",") if part.strip())
    if not languages:
        raise ValueError("At least one OCR language must be provided.")
    return languages


def _build_chunks(run_dir: Path, chunks_path: Path) -> int:
    chunks: list[dict[str, Any]] = []
    for document_path in sorted((run_dir / "documents").glob("*/document.v1.json")):
        document_payload = validate_json_file(document_path, "document.v1.schema.json")
        for chunk in build_chunks_from_document(document_payload):
            validate_payload(chunk, "chunks.v1.schema.json")
            chunks.append(chunk)

    with chunks_path.open("w", encoding="utf-8", newline="\n") as handle:
        for chunk in chunks:
            handle.write(json.dumps(chunk, ensure_ascii=False, separators=(",", ":")) + "\n")
    return len(chunks)


def _validate_run_dir(run_dir: Path) -> None:
    validate_json_file(run_dir / "run.json", "run.v1.schema.json")
    validate_json_file(run_dir / "summary.json", "summary.v1.schema.json")
    validate_json_file(run_dir / "queue-state.json", "queue-state.v1.schema.json")
    validate_json_file(run_dir / "processed-documents-catalog.json", "processed-documents-catalog.v1.schema.json")
    _validate_workbook(run_dir / "processed-documents-catalog.xlsx")
    _validate_jsonl(run_dir / "manifest.jsonl", "manifest.v1.schema.json")
    _validate_jsonl(run_dir / "review-required.jsonl", "review-required.v1.schema.json")
    if (run_dir / "chunks.v1.jsonl").exists():
        _validate_jsonl(run_dir / "chunks.v1.jsonl", "chunks.v1.schema.json")


def _validate_jsonl(path: Path, schema_filename: str) -> None:
    for line in path.read_text(encoding="utf-8").splitlines():
        if line.strip():
            validate_payload(json.loads(line), schema_filename)


def _validate_workbook(path: Path) -> None:
    workbook = load_workbook(path, read_only=True)
    try:
        if "Документы" not in workbook.sheetnames:
            raise ValueError(f"Workbook does not contain expected worksheet: {path}")
    finally:
        workbook.close()


def _remove_if_exists(path: Path) -> None:
    if path.exists():
        shutil.rmtree(path)


if __name__ == "__main__":
    raise SystemExit(main())