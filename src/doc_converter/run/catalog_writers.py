from __future__ import annotations

from collections.abc import Callable, Sequence
from dataclasses import dataclass
from pathlib import Path
from typing import Any, Protocol, TypedDict

from openpyxl import Workbook
from openpyxl.styles import Font

from .paths import ConverterError


class ProcessedDocumentCatalogEntry(TypedDict):
    relative_input_path: str
    original_filename: str
    output_dir: str | None
    output_folder_name: str | None
    status: str
    status_label: str
    issue: str | None


WriteValidatedJson = Callable[[Path, dict[str, object], str], None]


@dataclass(frozen=True)
class CatalogWriteContext:
    run_dir: Path
    run_id: str
    documents: list[ProcessedDocumentCatalogEntry]
    write_validated_json: WriteValidatedJson


class CatalogWriter(Protocol):
    @property
    def writer_id(self) -> str: ...

    def write(self, context: CatalogWriteContext) -> str | None: ...


@dataclass(frozen=True)
class JsonCatalogWriter:
    writer_id: str = "json"

    def write(self, context: CatalogWriteContext) -> str | None:
        payload: dict[str, object] = {
            "schema_version": "processed-documents-catalog.v1",
            "run_id": context.run_id,
            "documents": context.documents,
        }
        path = context.run_dir / "processed-documents-catalog.json"
        context.write_validated_json(path, payload, "processed-documents-catalog.v1.schema.json")
        return path.name


@dataclass(frozen=True)
class XlsxCatalogWriter:
    writer_id: str = "xlsx"

    def write(self, context: CatalogWriteContext) -> str | None:
        workbook = Workbook()
        sheet = workbook.active
        if sheet is None:
            raise ConverterError("Workbook must have an active worksheet")
        sheet.title = "Документы"
        sheet.freeze_panes = "A2"

        headers = [
            "Исходный файл",
            "Относительный путь",
            "Папка документа",
            "Статус",
            "Пояснение",
            "document.v1.json",
            "search_text.txt",
        ]
        for column_index, header in enumerate(headers, start=1):
            cell = sheet.cell(row=1, column=column_index, value=header)
            cell.font = Font(bold=True)

        for row_index, document in enumerate(context.documents, start=2):
            output_dir = document["output_dir"]
            output_folder_name = document["output_folder_name"]
            folder_path = (context.run_dir / output_dir).resolve() if output_dir is not None else None
            document_path = folder_path / "document.v1.json" if folder_path is not None else None
            search_text_path = folder_path / "search_text.txt" if folder_path is not None else None

            sheet.cell(row=row_index, column=1, value=document["original_filename"])
            sheet.cell(row=row_index, column=2, value=document["relative_input_path"])
            sheet.cell(row=row_index, column=3, value=output_folder_name)
            sheet.cell(row=row_index, column=4, value=document["status_label"])
            sheet.cell(row=row_index, column=5, value=document["issue"])
            sheet.cell(row=row_index, column=6, value="document.v1.json" if document_path is not None else None)
            sheet.cell(row=row_index, column=7, value="search_text.txt" if search_text_path is not None else None)

            if folder_path is not None:
                _set_hyperlink(sheet.cell(row=row_index, column=3), folder_path)
            if document_path is not None and document_path.exists():
                _set_hyperlink(sheet.cell(row=row_index, column=6), document_path)
            if search_text_path is not None and search_text_path.exists():
                _set_hyperlink(sheet.cell(row=row_index, column=7), search_text_path)

        column_widths = {
            "A": 36,
            "B": 48,
            "C": 28,
            "D": 24,
            "E": 48,
            "F": 20,
            "G": 18,
        }
        for column_letter, width in column_widths.items():
            sheet.column_dimensions[column_letter].width = width

        path = context.run_dir / "processed-documents-catalog.xlsx"
        workbook.save(path)
        return path.name


def build_catalog_writers(writer_ids: Sequence[str] | None) -> tuple[CatalogWriter, ...]:
    normalized = tuple(writer_ids) if writer_ids is not None else ("json", "xlsx")
    if not normalized:
        raise RuntimeError("At least one catalog writer must be configured.")

    writers: list[CatalogWriter] = []
    seen: set[str] = set()
    for writer_id in normalized:
        key = str(writer_id).strip().lower()
        if not key or key in seen:
            continue
        seen.add(key)
        if key == "json":
            writers.append(JsonCatalogWriter())
            continue
        if key == "xlsx":
            writers.append(XlsxCatalogWriter())
            continue
        raise RuntimeError(f"Unsupported catalog writer: {writer_id}")
    if not writers:
        raise RuntimeError("At least one catalog writer must be configured.")
    return tuple(writers)


def _set_hyperlink(cell: Any, target_path: Path) -> None:
    cell.hyperlink = target_path.resolve().as_uri()
    cell.style = "Hyperlink"


__all__ = [
    "CatalogWriteContext",
    "CatalogWriter",
    "JsonCatalogWriter",
    "ProcessedDocumentCatalogEntry",
    "XlsxCatalogWriter",
    "build_catalog_writers",
]
