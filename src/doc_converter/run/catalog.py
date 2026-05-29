from __future__ import annotations

from pathlib import Path
from typing import Any, TypedDict

from openpyxl import Workbook
from openpyxl.styles import Font

from ..inventory import InventoryRecord, UnsupportedInventoryRecord
from .paths import ConverterError


class ProcessedDocumentCatalogEntry(TypedDict):
    relative_input_path: str
    original_filename: str
    output_dir: str | None
    output_folder_name: str | None
    status: str
    status_label: str
    issue: str | None


def _write_processed_documents_catalog(
    *,
    run_dir: Path,
    run_id: str,
    inventory_records: list[InventoryRecord],
    unsupported_records: list[UnsupportedInventoryRecord],
    manifest_records: list[dict[str, object]],
    error_details_by_relative_path: dict[str, str],
    write_validated_json: Any,
) -> None:
    documents = _build_processed_documents_catalog_documents(
        inventory_records=inventory_records,
        unsupported_records=unsupported_records,
        manifest_records=manifest_records,
        error_details_by_relative_path=error_details_by_relative_path,
    )

    write_validated_json(
        run_dir / "processed-documents-catalog.json",
        {
            "schema_version": "processed-documents-catalog.v1",
            "run_id": run_id,
            "documents": documents,
        },
        "processed-documents-catalog.v1.schema.json",
    )
    _write_processed_documents_catalog_xlsx(run_dir=run_dir, documents=documents)


def _build_processed_documents_catalog_documents(
    *,
    inventory_records: list[InventoryRecord],
    unsupported_records: list[UnsupportedInventoryRecord],
    manifest_records: list[dict[str, object]],
    error_details_by_relative_path: dict[str, str],
) -> list[ProcessedDocumentCatalogEntry]:
    manifest_by_relative_path = {
        str(record["relative_path"]): record
        for record in manifest_records
        if isinstance(record.get("relative_path"), str)
    }
    documents: list[ProcessedDocumentCatalogEntry] = []

    for inventory_record in inventory_records:
        manifest_record = manifest_by_relative_path.get(inventory_record.relative_path)
        if manifest_record is None:
            continue
        output_dir = manifest_record.get("output_dir")
        output_folder = output_dir if isinstance(output_dir, str) else None
        status = str(manifest_record.get("status", ""))
        documents.append(
            {
                "relative_input_path": inventory_record.relative_path,
                "original_filename": inventory_record.filename,
                "output_dir": output_folder,
                "output_folder_name": Path(output_folder).name if output_folder else None,
                "status": status,
                "status_label": _catalog_status_label(status),
                "issue": _catalog_issue(
                    status=status,
                    manifest_record=manifest_record,
                    error_details_by_relative_path=error_details_by_relative_path,
                ),
            }
        )

    for unsupported_record in unsupported_records:
        documents.append(
            {
                "relative_input_path": unsupported_record.relative_path,
                "original_filename": unsupported_record.filename,
                "output_dir": None,
                "output_folder_name": None,
                "status": "skipped_unsupported",
                "status_label": _catalog_status_label("skipped_unsupported"),
                "issue": "; ".join(unsupported_record.warnings) if unsupported_record.warnings else None,
            }
        )

    return documents


def _write_processed_documents_catalog_xlsx(
    *,
    run_dir: Path,
    documents: list[ProcessedDocumentCatalogEntry],
) -> None:
    workbook = Workbook()
    sheet = workbook.active
    if sheet is None:
        raise ConverterError("Workbook must have an active worksheet")
    sheet.title = "Документы"
    sheet.freeze_panes = "A2"

    headers = [
        "Исходный файл",
        "Относительный путь",
        "Папка документа",
        "Статус",
        "Пояснение",
        "document.v1.json",
        "search_text.txt",
    ]
    for column_index, header in enumerate(headers, start=1):
        cell = sheet.cell(row=1, column=column_index, value=header)
        cell.font = Font(bold=True)

    for row_index, document in enumerate(documents, start=2):
        output_dir = document["output_dir"]
        output_folder_name = document["output_folder_name"]
        folder_path = (run_dir / output_dir).resolve() if output_dir is not None else None
        document_path = folder_path / "document.v1.json" if folder_path is not None else None
        search_text_path = folder_path / "search_text.txt" if folder_path is not None else None

        sheet.cell(row=row_index, column=1, value=document["original_filename"])
        sheet.cell(row=row_index, column=2, value=document["relative_input_path"])
        sheet.cell(row=row_index, column=3, value=output_folder_name)
        sheet.cell(row=row_index, column=4, value=document["status_label"])
        sheet.cell(row=row_index, column=5, value=document["issue"])
        sheet.cell(row=row_index, column=6, value="document.v1.json" if document_path is not None else None)
        sheet.cell(row=row_index, column=7, value="search_text.txt" if search_text_path is not None else None)

        if folder_path is not None:
            _set_hyperlink(sheet.cell(row=row_index, column=3), folder_path)
        if document_path is not None and document_path.exists():
            _set_hyperlink(sheet.cell(row=row_index, column=6), document_path)
        if search_text_path is not None and search_text_path.exists():
            _set_hyperlink(sheet.cell(row=row_index, column=7), search_text_path)

    column_widths = {
        "A": 36,
        "B": 48,
        "C": 28,
        "D": 24,
        "E": 48,
        "F": 20,
        "G": 18,
    }
    for column_letter, width in column_widths.items():
        sheet.column_dimensions[column_letter].width = width

    workbook.save(run_dir / "processed-documents-catalog.xlsx")


def _set_hyperlink(cell: Any, target_path: Path) -> None:
    cell.hyperlink = target_path.resolve().as_uri()
    cell.style = "Hyperlink"


def _catalog_status_label(status: str) -> str:
    return {
        "success": "Успешная обработка",
        "partial_success": "Частичная обработка",
        "failed": "Ошибка",
        "skipped_duplicate": "Дубликат",
        "skipped_unsupported": "Неподдерживаемый формат",
    }.get(status, status or "Неизвестный статус")


def _catalog_issue(
    *,
    status: str,
    manifest_record: dict[str, object],
    error_details_by_relative_path: dict[str, str],
) -> str | None:
    relative_path = str(manifest_record.get("relative_path", ""))
    if status == "failed":
        return error_details_by_relative_path.get(relative_path) or str(manifest_record.get("error") or "") or None
    if status == "partial_success":
        warnings = manifest_record.get("warnings")
        if isinstance(warnings, list):
            normalized = [str(item).strip() for item in warnings if str(item).strip()]
            if normalized:
                return "; ".join(normalized)
    if status == "skipped_duplicate":
        duplicate_of = manifest_record.get("duplicate_of")
        if isinstance(duplicate_of, str) and duplicate_of:
            return f"duplicate_of:{duplicate_of}"
    return None