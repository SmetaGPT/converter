from __future__ import annotations

import json
import shutil
from dataclasses import dataclass
from datetime import UTC, datetime
from pathlib import Path
from typing import Any, Callable, TypedDict, Union

from openpyxl import Workbook
from openpyxl.styles import Font

from . import __version__
from .config import ConverterConfig, serialize_converter_options
from .converters.docx import ConversionResult, convert_docx
from .converters.pdf_scan import PdfScanConversionResult, convert_pdf_scan
from .converters.pdf_text import PdfTextConversionResult, convert_pdf_text
from .converters.xlsx import XlsxConversionResult, convert_xlsx
from .formula_recognition import FormulaRecognitionPostprocessResult, run_formula_recognition_postprocess
from .inventory import InventoryRecord, UnsupportedInventoryRecord, build_inventory
from .schema_validation import SchemaValidationError, validate_json_file, validate_payload


class ConverterError(RuntimeError):
    """Raised when a conversion run cannot be started."""


@dataclass(frozen=True)
class RunResult:
    run_id: str
    run_dir: Path
    discovered_files: int
    supported_files: int
    status: str


@dataclass(frozen=True)
class ResumeHit:
    run_id: str
    output_dir: str
    status: str
    manifest_record: dict[str, Any]


class ProcessedDocumentCatalogEntry(TypedDict):
    relative_input_path: str
    original_filename: str
    output_dir: str | None
    output_folder_name: str | None
    status: str
    status_label: str
    issue: str | None


def run_convert_folder(
    config: ConverterConfig,
    *,
    progress_callback: Callable[[dict[str, object]], None] | None = None,
    should_cancel: Callable[[], bool] | None = None,
) -> RunResult:
    input_dir, output_dir = validate_run_directories(config.input_dir, config.output_dir)

    run_id = _new_run_id()
    run_dir = _allocate_run_dir(output_dir / "runs", run_id)
    documents_dir = run_dir / "documents"
    documents_dir.mkdir(parents=True, exist_ok=False)

    run_metadata = _build_run_metadata(run_id, input_dir, output_dir, config)
    _write_validated_json(run_dir / "run.json", run_metadata, "run.v1.schema.json")

    inventory = build_inventory(input_dir)
    inventory_records = inventory.supported_records
    unsupported_records = inventory.unsupported_records

    manifest_path = run_dir / "manifest.jsonl"
    log_path = run_dir / "processing-log.jsonl"
    errors_path = run_dir / "errors.jsonl"
    queue_state_path = run_dir / "queue-state.json"
    review_required_path = run_dir / "review-required.jsonl"

    _write_text(manifest_path, "")
    _write_text(errors_path, "")
    _write_text(review_required_path, "")
    _append_jsonl(log_path, {"event": "run_started", "run_id": run_id, "status": "ok"})

    resume_index = _load_resume_index(output_dir / "runs", run_dir)
    total_files = len(inventory_records)
    error_details_by_relative_path: dict[str, str] = {}
    for unsupported_record in unsupported_records:
        _append_jsonl(
            log_path,
            {
                "event": "document_skipped_unsupported",
                "run_id": run_id,
                "relative_path": unsupported_record.relative_path,
                "filename": unsupported_record.filename,
                "format": unsupported_record.format,
                "size_bytes": unsupported_record.size_bytes,
                "warnings": list(unsupported_record.warnings),
                "status": "skipped_unsupported",
            },
        )
    _emit_progress(
        progress_callback,
        {
            "event": "inventory_built",
            "run_id": run_id,
            "total_files": total_files,
            "processed_files": 0,
        },
    )

    completed: list[str] = []
    partial: list[str] = []
    failed: list[str] = []
    final_manifest_records: list[dict[str, object]] = []
    manifest_by_relative_path: dict[str, dict[str, object]] = {}
    cancelled = False

    for record in inventory_records:
        if should_cancel is not None and should_cancel():
            cancelled = True
            _append_jsonl(
                log_path,
                {
                    "event": "run_cancelled",
                    "run_id": run_id,
                    "status": "cancelled",
                    "processed_files": len(completed) + len(partial) + len(failed),
                },
            )
            _emit_progress(
                progress_callback,
                {
                    "event": "run_cancelled",
                    "run_id": run_id,
                    "total_files": total_files,
                    "processed_files": len(completed) + len(partial) + len(failed),
                },
            )
            break

        manifest_record = record.to_manifest_record(run_id)
        if record.route in {"docx_native", "pdf_text", "pdf_scan", "xlsx_native"}:
            document_dir = _document_output_path(documents_dir, record.sha256)
            _emit_progress(
                progress_callback,
                {
                    "event": "document_started",
                    "run_id": run_id,
                    "relative_path": record.relative_path,
                    "total_files": total_files,
                    "processed_files": len(completed) + len(partial) + len(failed),
                },
            )
            if record.duplicate_of is not None:
                primary_manifest = manifest_by_relative_path.get(record.duplicate_of)
                if primary_manifest is None:
                    raise ConverterError(f"Primary duplicate record was not processed first: {record.duplicate_of}")
                manifest_record["status"] = "skipped_duplicate"
                if "output_dir" in primary_manifest:
                    manifest_record["output_dir"] = primary_manifest["output_dir"]
                _copy_result_metadata(primary_manifest, manifest_record)
                if config.options.include_originals and "output_dir" in manifest_record:
                    _copy_original_file(
                        input_dir / record.relative_path,
                        run_dir / str(manifest_record["output_dir"]),
                        record.relative_path,
                    )
                completed.append(record.relative_path)
                _append_jsonl(
                    log_path,
                    {
                        "event": "document_skipped_duplicate",
                        "run_id": run_id,
                        "relative_path": record.relative_path,
                        "duplicate_of": record.duplicate_of,
                        "status": "skipped_duplicate",
                    },
                )
                _emit_progress(
                    progress_callback,
                    {
                        "event": "document_finished",
                        "run_id": run_id,
                        "relative_path": record.relative_path,
                        "status": "skipped_duplicate",
                        "total_files": total_files,
                        "processed_files": len(completed) + len(partial) + len(failed),
                    },
                )
                final_manifest_records.append(manifest_record)
                manifest_by_relative_path[record.relative_path] = manifest_record
                continue

            resume_hit = resume_index.get((record.sha256, record.relative_path))
            if resume_hit is not None:
                source_document_dir = (output_dir / "runs" / resume_hit.run_id / resume_hit.output_dir).resolve()
                _reuse_document_dir(source_document_dir, document_dir)
                manifest_record["status"] = resume_hit.status
                manifest_record["output_dir"] = document_dir.relative_to(run_dir).as_posix()
                manifest_record["resumed_from_run_id"] = resume_hit.run_id
                manifest_record["reused_previous_output"] = True
                _copy_result_metadata(resume_hit.manifest_record, manifest_record)
                if config.options.include_originals:
                    _copy_original_file(input_dir / record.relative_path, document_dir, record.relative_path)
                if resume_hit.status == "partial_success":
                    partial.append(record.relative_path)
                else:
                    completed.append(record.relative_path)
                _append_jsonl(
                    log_path,
                    {
                        "event": "document_reused",
                        "run_id": run_id,
                        "relative_path": record.relative_path,
                        "route": record.route,
                        "status": resume_hit.status,
                        "resumed_from_run_id": resume_hit.run_id,
                    },
                )
                _emit_progress(
                    progress_callback,
                    {
                        "event": "document_finished",
                        "run_id": run_id,
                        "relative_path": record.relative_path,
                        "status": resume_hit.status,
                        "total_files": total_files,
                        "processed_files": len(completed) + len(partial) + len(failed),
                    },
                )
                final_manifest_records.append(manifest_record)
                manifest_by_relative_path[record.relative_path] = manifest_record
                continue

            try:
                result: Union[ConversionResult, PdfTextConversionResult, PdfScanConversionResult, XlsxConversionResult]
                if record.route == "docx_native":
                    result = convert_docx(
                        input_dir / record.relative_path,
                        document_dir,
                        record.sha256,
                        relative_source_path=record.relative_path,
                    )
                    manifest_record["assets_count"] = result.assets_count
                    manifest_record["search_text_chars"] = result.search_text_chars
                elif record.route == "pdf_text":
                    result = convert_pdf_text(
                        input_dir / record.relative_path,
                        document_dir,
                        record.sha256,
                        relative_source_path=record.relative_path,
                    )
                    manifest_record["pages"] = result.pages
                    manifest_record["search_text_chars"] = result.text_chars
                elif record.route == "xlsx_native":
                    result = convert_xlsx(
                        input_dir / record.relative_path,
                        document_dir,
                        record.sha256,
                        relative_source_path=record.relative_path,
                    )
                    manifest_record["sheets"] = result.sheets
                    manifest_record["search_text_chars"] = result.text_chars
                    manifest_record["formula_cells"] = result.formula_cells
                    manifest_record["warnings"] = list(result.warnings)
                else:
                    result = convert_pdf_scan(
                        input_dir / record.relative_path,
                        document_dir,
                        record.sha256,
                        config.options.ocr_languages,
                        relative_source_path=record.relative_path,
                    )
                    manifest_record["pages"] = result.pages
                    manifest_record["search_text_chars"] = result.text_chars
                    manifest_record["warnings"] = list(result.warnings)
                manifest_record["status"] = result.status
                manifest_record["output_dir"] = document_dir.relative_to(run_dir).as_posix()
                manifest_record["units_count"] = result.units_count
                formula_result = _run_formula_recognition_stage(document_dir, config)
                _merge_formula_recognition_manifest_data(manifest_record, formula_result)
                if config.options.include_originals:
                    manifest_record["original_copy_path"] = _copy_original_file(
                        input_dir / record.relative_path,
                        document_dir,
                        record.relative_path,
                    )
                if result.status == "partial_success":
                    partial.append(record.relative_path)
                else:
                    completed.append(record.relative_path)
                _append_jsonl(
                    log_path,
                    {
                        "event": "document_converted",
                        "run_id": run_id,
                        "relative_path": record.relative_path,
                        "route": record.route,
                        "status": result.status,
                    },
                )
                _emit_progress(
                    progress_callback,
                    {
                        "event": "document_finished",
                        "run_id": run_id,
                        "relative_path": record.relative_path,
                        "status": result.status,
                        "total_files": total_files,
                        "processed_files": len(completed) + len(partial) + len(failed),
                    },
                )
            except Exception as exc:  # noqa: BLE001 - batch run records per-document failures.
                manifest_record["status"] = "failed"
                manifest_record["error"] = type(exc).__name__
                error_details_by_relative_path[record.relative_path] = str(exc)
                failed.append(record.relative_path)
                _append_jsonl(
                    errors_path,
                    {
                        "event": "document_failed",
                        "run_id": run_id,
                        "relative_path": record.relative_path,
                        "route": record.route,
                        "error_type": type(exc).__name__,
                        "error": str(exc),
                    },
                )
                _emit_progress(
                    progress_callback,
                    {
                        "event": "document_finished",
                        "run_id": run_id,
                        "relative_path": record.relative_path,
                        "status": "failed",
                        "total_files": total_files,
                        "processed_files": len(completed) + len(partial) + len(failed),
                    },
                )
        manifest_by_relative_path[record.relative_path] = manifest_record
        final_manifest_records.append(manifest_record)

    for manifest_record in final_manifest_records:
        _append_validated_jsonl(manifest_path, manifest_record, "manifest.v1.schema.json")

    review_required_records = _build_review_required_records(run_dir, run_id, final_manifest_records)
    for review_required_record in review_required_records:
        _append_validated_jsonl(review_required_path, review_required_record, "review-required.v1.schema.json")

    _write_processed_documents_catalog(
        run_dir=run_dir,
        run_id=run_id,
        inventory_records=inventory_records,
        unsupported_records=unsupported_records,
        manifest_records=final_manifest_records,
        error_details_by_relative_path=error_details_by_relative_path,
    )

    duplicate_groups = sorted({record.duplicate_group_id for record in inventory_records if record.duplicate_group_id})
    queue_state = {
        "schema_version": "queue-state.v1",
        "run_id": run_id,
        "status": _queue_state_status(cancelled, completed, partial, failed),
        "queued": [
            record.relative_path
            for record in inventory_records
            if record.relative_path not in completed and record.relative_path not in partial and record.relative_path not in failed
        ],
        "completed": completed,
        "partial": partial,
        "failed": failed,
    }
    _write_validated_json(queue_state_path, queue_state, "queue-state.v1.schema.json")

    summary = {
        "schema_version": "summary.v1",
        "run_id": run_id,
        "status": _summary_status(cancelled, completed, partial, failed),
        "input_dir": str(input_dir),
        "output_dir": str(output_dir),
        "run_dir": str(run_dir),
        "discovered_files": inventory.scanned_files,
        "supported_files": len(inventory_records),
        "unsupported_files": len(unsupported_records),
        "duplicate_groups": len(duplicate_groups),
        "processed_files": len(completed) + len(partial),
        "partial_files": len(partial),
        "failed_files": len(failed),
        "review_required_files": len(review_required_records),
        "review_required_reasons": _count_reason_map(record["flags"] for record in review_required_records),
        "partial_reasons": _count_reason_map(
            record["flags"] for record in review_required_records if record["status"] == "partial_success"
        ),
        "failed_reasons": _count_reason_map(
            [str(record["error_type"])]
            for record in review_required_records
            if record["status"] == "failed" and record.get("error_type")
        ),
    }
    _write_validated_json(run_dir / "summary.json", summary, "summary.v1.schema.json")
    final_status = _summary_status(cancelled, completed, partial, failed)
    _append_jsonl(log_path, {"event": "run_completed", "run_id": run_id, "status": final_status})
    _emit_progress(
        progress_callback,
        {
            "event": "run_completed",
            "run_id": run_id,
            "status": final_status,
            "total_files": total_files,
            "processed_files": len(completed) + len(partial) + len(failed),
        },
    )

    return RunResult(
        run_id=run_id,
        run_dir=run_dir,
        discovered_files=inventory.scanned_files,
        supported_files=len(inventory_records),
        status=final_status,
    )


def validate_run_directories(input_dir: Path, output_dir: Path) -> tuple[Path, Path]:
    resolved_input_dir = input_dir.expanduser().resolve()
    resolved_output_dir = output_dir.expanduser().resolve()

    _validate_startup_paths(resolved_input_dir, resolved_output_dir)
    return resolved_input_dir, resolved_output_dir


def _validate_startup_paths(input_dir: Path, output_dir: Path) -> None:
    if not input_dir.exists():
        raise ConverterError(f"Input directory does not exist: {input_dir}")
    if not input_dir.is_dir():
        raise ConverterError(f"Input path is not a directory: {input_dir}")
    if output_dir.exists() and not output_dir.is_dir():
        raise ConverterError(f"Output path is not a directory: {output_dir}")
    if _paths_overlap(input_dir, output_dir):
        raise ConverterError(
            "Input and output directories must be different and must not be nested inside each other: "
            f"input={input_dir}, output={output_dir}"
        )
    output_dir.mkdir(parents=True, exist_ok=True)


def _paths_overlap(first_path: Path, second_path: Path) -> bool:
    return _is_relative_to(first_path, second_path) or _is_relative_to(second_path, first_path)


def _is_relative_to(path: Path, other_path: Path) -> bool:
    try:
        path.relative_to(other_path)
    except ValueError:
        return False
    return True


def _new_run_id() -> str:
    return datetime.now(UTC).strftime("%Y%m%dT%H%M%SZ")


def _allocate_run_dir(runs_dir: Path, base_run_id: str) -> Path:
    runs_dir.mkdir(parents=True, exist_ok=True)
    candidate = runs_dir / base_run_id
    if not candidate.exists():
        candidate.mkdir()
        return candidate

    for index in range(1, 1000):
        candidate = runs_dir / f"{base_run_id}-{index:03d}"
        if not candidate.exists():
            candidate.mkdir()
            return candidate

    raise ConverterError(f"Unable to allocate run directory under: {runs_dir}")


def _build_run_metadata(
    run_id: str,
    input_dir: Path,
    output_dir: Path,
    config: ConverterConfig,
) -> dict[str, Any]:
    return {
        "schema_version": "run.v1",
        "run_id": run_id,
        "converter_version": __version__,
        "created_at": datetime.now(UTC).isoformat(),
        "input_dir": str(input_dir),
        "output_dir": str(output_dir),
        "options": serialize_converter_options(config.options),
    }


def _document_output_dir(documents_dir: Path, sha256: str) -> Path:
    document_dir = _document_output_path(documents_dir, sha256)
    document_dir.mkdir(parents=True, exist_ok=True)
    return document_dir


def _document_output_path(documents_dir: Path, sha256: str) -> Path:
    return documents_dir / f"sha256_{sha256}"


def _run_formula_recognition_stage(
    document_dir: Path,
    config: ConverterConfig,
) -> FormulaRecognitionPostprocessResult:
    try:
        return run_formula_recognition_postprocess(document_dir, config.options.formula_recognition)
    except Exception:  # noqa: BLE001 - best-effort network stage must not fail document conversion.
        return FormulaRecognitionPostprocessResult(warnings=("formula_recognition_postprocess_failed",))


def _merge_formula_recognition_manifest_data(
    manifest_record: dict[str, object],
    formula_result: FormulaRecognitionPostprocessResult,
) -> None:
    if formula_result.attempted > 0:
        manifest_record["formula_recognition_attempted"] = formula_result.attempted
        manifest_record["formula_recognition_recognized"] = formula_result.recognized
        manifest_record["formula_recognition_provider_calls"] = formula_result.provider_calls
    if formula_result.artifact_path is not None:
        manifest_record["formula_recognition_results_path"] = formula_result.artifact_path
    if formula_result.warnings:
        warnings = manifest_record.get("warnings")
        if isinstance(warnings, list):
            warning_list = [str(item) for item in warnings]
        else:
            warning_list = []
        for warning in formula_result.warnings:
            if warning not in warning_list:
                warning_list.append(warning)
        manifest_record["warnings"] = warning_list


def _write_json(path: Path, payload: dict[str, Any]) -> None:
    path.write_text(json.dumps(payload, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")


def _write_validated_json(path: Path, payload: dict[str, Any], schema_filename: str) -> None:
    validate_payload(payload, schema_filename)
    _write_json(path, payload)


def _append_validated_jsonl(path: Path, payload: dict[str, Any], schema_filename: str) -> None:
    validate_payload(payload, schema_filename)
    _append_jsonl(path, payload)


def _copy_original_file(source_path: Path, document_dir: Path, relative_path: str) -> str:
    target = document_dir / "originals" / Path(relative_path)
    target.parent.mkdir(parents=True, exist_ok=True)
    shutil.copy2(source_path, target)
    return target.relative_to(document_dir).as_posix()


def _reuse_document_dir(source_dir: Path, target_dir: Path) -> None:
    if target_dir.exists():
        shutil.rmtree(target_dir)
    shutil.copytree(source_dir, target_dir)


def _copy_result_metadata(source_record: dict[str, object], target_record: dict[str, object]) -> None:
    for key in ("assets_count", "pages", "search_text_chars", "units_count", "warnings", "original_copy_path"):
        if key in source_record:
            target_record[key] = source_record[key]


def _load_resume_index(runs_dir: Path, current_run_dir: Path) -> dict[tuple[str, str], ResumeHit]:
    if not runs_dir.exists():
        return {}

    index: dict[tuple[str, str], ResumeHit] = {}
    for candidate in sorted((path for path in runs_dir.iterdir() if path.is_dir()), reverse=True):
        if candidate.resolve() == current_run_dir.resolve():
            continue
        manifest_path = candidate / "manifest.jsonl"
        if not manifest_path.exists():
            continue
        for line in manifest_path.read_text(encoding="utf-8").splitlines():
            if not line.strip():
                continue
            try:
                manifest_record = json.loads(line)
                validate_payload(manifest_record, "manifest.v1.schema.json")
            except (json.JSONDecodeError, SchemaValidationError):
                continue

            status = str(manifest_record.get("status", ""))
            output_dir = manifest_record.get("output_dir")
            sha256 = manifest_record.get("sha256")
            relative_path = manifest_record.get("relative_path")
            if status not in {"success", "partial_success"}:
                continue
            if not isinstance(output_dir, str) or not isinstance(sha256, str) or not isinstance(relative_path, str):
                continue

            document_path = candidate / output_dir / "document.v1.json"
            if not document_path.exists():
                continue
            try:
                validate_json_file(document_path, "document.v1.schema.json")
            except (json.JSONDecodeError, OSError, SchemaValidationError):
                continue

            key = (sha256, relative_path)
            if key in index:
                continue
            index[key] = ResumeHit(
                run_id=candidate.name,
                output_dir=output_dir,
                status=status,
                manifest_record=manifest_record,
            )
    return index


def _emit_progress(
    progress_callback: Callable[[dict[str, object]], None] | None,
    payload: dict[str, object],
) -> None:
    if progress_callback is not None:
        progress_callback(payload)


def _summary_status(cancelled: bool, completed: list[str], partial: list[str], failed: list[str]) -> str:
    if cancelled:
        return "cancelled"
    if failed and not completed and not partial:
        return "failed"
    if failed or partial:
        return "partial_success"
    return "success"


def _queue_state_status(cancelled: bool, completed: list[str], partial: list[str], failed: list[str]) -> str:
    if cancelled:
        return "cancelled"
    if failed and not completed and not partial:
        return "failed"
    if failed or partial:
        return "partial_success"
    return "completed"


def _build_review_required_records(
    run_dir: Path,
    run_id: str,
    manifest_records: list[dict[str, object]],
) -> list[dict[str, object]]:
    records: list[dict[str, object]] = []
    for manifest_record in manifest_records:
        status = str(manifest_record.get("status", ""))
        output_dir = manifest_record.get("output_dir")
        flags: list[str] = []
        warnings: list[str] = []
        document_id: str | None = None
        if isinstance(output_dir, str):
            document_path = run_dir / output_dir / "document.v1.json"
            if document_path.exists():
                try:
                    payload = validate_json_file(document_path, "document.v1.schema.json")
                    document_id = str(payload.get("document_id"))
                    flags = [str(flag) for flag in payload.get("quality", {}).get("flags", [])]
                    warnings = [str(warning) for warning in payload.get("quality", {}).get("warnings", [])]
                except (OSError, json.JSONDecodeError, SchemaValidationError) as exc:
                    warnings.append(f"document_validation_failed:{type(exc).__name__}")
        error_type = str(manifest_record.get("error", "")) or None
        if status == "failed" and error_type is not None:
            warnings.append(error_type)
        if status in {"partial_success", "failed"} or "review_required" in flags or warnings:
            records.append(
                {
                    "schema_version": "review-required.v1",
                    "run_id": run_id,
                    "relative_path": str(manifest_record["relative_path"]),
                    "status": status,
                    "document_id": document_id,
                    "output_dir": output_dir if isinstance(output_dir, str) else None,
                    "flags": flags,
                    "warnings": warnings,
                    "error_type": error_type,
                }
            )
    return records


def _write_processed_documents_catalog(
    *,
    run_dir: Path,
    run_id: str,
    inventory_records: list[InventoryRecord],
    unsupported_records: list[UnsupportedInventoryRecord],
    manifest_records: list[dict[str, object]],
    error_details_by_relative_path: dict[str, str],
) -> None:
    documents = _build_processed_documents_catalog_documents(
        inventory_records=inventory_records,
        unsupported_records=unsupported_records,
        manifest_records=manifest_records,
        error_details_by_relative_path=error_details_by_relative_path,
    )

    _write_validated_json(
        run_dir / "processed-documents-catalog.json",
        {
            "schema_version": "processed-documents-catalog.v1",
            "run_id": run_id,
            "documents": documents,
        },
        "processed-documents-catalog.v1.schema.json",
    )
    _write_processed_documents_catalog_xlsx(run_dir=run_dir, documents=documents)


def _build_processed_documents_catalog_documents(
    *,
    inventory_records: list[InventoryRecord],
    unsupported_records: list[UnsupportedInventoryRecord],
    manifest_records: list[dict[str, object]],
    error_details_by_relative_path: dict[str, str],
) -> list[ProcessedDocumentCatalogEntry]:
    manifest_by_relative_path = {
        str(record["relative_path"]): record
        for record in manifest_records
        if isinstance(record.get("relative_path"), str)
    }
    documents: list[ProcessedDocumentCatalogEntry] = []

    for inventory_record in inventory_records:
        manifest_record = manifest_by_relative_path.get(inventory_record.relative_path)
        if manifest_record is None:
            continue
        output_dir = manifest_record.get("output_dir")
        output_folder = output_dir if isinstance(output_dir, str) else None
        status = str(manifest_record.get("status", ""))
        documents.append(
            {
                "relative_input_path": inventory_record.relative_path,
                "original_filename": inventory_record.filename,
                "output_dir": output_folder,
                "output_folder_name": Path(output_folder).name if output_folder else None,
                "status": status,
                "status_label": _catalog_status_label(status),
                "issue": _catalog_issue(
                    status=status,
                    manifest_record=manifest_record,
                    error_details_by_relative_path=error_details_by_relative_path,
                ),
            }
        )

    for unsupported_record in unsupported_records:
        documents.append(
            {
                "relative_input_path": unsupported_record.relative_path,
                "original_filename": unsupported_record.filename,
                "output_dir": None,
                "output_folder_name": None,
                "status": "skipped_unsupported",
                "status_label": _catalog_status_label("skipped_unsupported"),
                "issue": "; ".join(unsupported_record.warnings) if unsupported_record.warnings else None,
            }
        )

    return documents


def _write_processed_documents_catalog_xlsx(
    *,
    run_dir: Path,
    documents: list[ProcessedDocumentCatalogEntry],
) -> None:
    workbook = Workbook()
    sheet = workbook.active
    if sheet is None:
        raise ConverterError("Workbook must have an active worksheet")
    sheet.title = "Документы"
    sheet.freeze_panes = "A2"

    headers = [
        "Исходный файл",
        "Относительный путь",
        "Папка документа",
        "Статус",
        "Пояснение",
        "document.v1.json",
        "search_text.txt",
    ]
    for column_index, header in enumerate(headers, start=1):
        cell = sheet.cell(row=1, column=column_index, value=header)
        cell.font = Font(bold=True)

    for row_index, document in enumerate(documents, start=2):
        output_dir = document["output_dir"]
        output_folder_name = document["output_folder_name"]
        folder_path = (run_dir / output_dir).resolve() if output_dir is not None else None
        document_path = folder_path / "document.v1.json" if folder_path is not None else None
        search_text_path = folder_path / "search_text.txt" if folder_path is not None else None

        sheet.cell(row=row_index, column=1, value=document["original_filename"])
        sheet.cell(row=row_index, column=2, value=document["relative_input_path"])
        sheet.cell(row=row_index, column=3, value=output_folder_name)
        sheet.cell(row=row_index, column=4, value=document["status_label"])
        sheet.cell(row=row_index, column=5, value=document["issue"])
        sheet.cell(row=row_index, column=6, value="document.v1.json" if document_path is not None else None)
        sheet.cell(row=row_index, column=7, value="search_text.txt" if search_text_path is not None else None)

        if folder_path is not None:
            _set_hyperlink(sheet.cell(row=row_index, column=3), folder_path)
        if document_path is not None and document_path.exists():
            _set_hyperlink(sheet.cell(row=row_index, column=6), document_path)
        if search_text_path is not None and search_text_path.exists():
            _set_hyperlink(sheet.cell(row=row_index, column=7), search_text_path)

    column_widths = {
        "A": 36,
        "B": 48,
        "C": 28,
        "D": 24,
        "E": 48,
        "F": 20,
        "G": 18,
    }
    for column_letter, width in column_widths.items():
        sheet.column_dimensions[column_letter].width = width

    workbook.save(run_dir / "processed-documents-catalog.xlsx")


def _set_hyperlink(cell: Any, target_path: Path) -> None:
    cell.hyperlink = target_path.resolve().as_uri()
    cell.style = "Hyperlink"


def _catalog_status_label(status: str) -> str:
    return {
        "success": "Успешная обработка",
        "partial_success": "Частичная обработка",
        "failed": "Ошибка",
        "skipped_duplicate": "Дубликат",
        "skipped_unsupported": "Неподдерживаемый формат",
    }.get(status, status or "Неизвестный статус")


def _catalog_issue(
    *,
    status: str,
    manifest_record: dict[str, object],
    error_details_by_relative_path: dict[str, str],
) -> str | None:
    relative_path = str(manifest_record.get("relative_path", ""))
    if status == "failed":
        return error_details_by_relative_path.get(relative_path) or str(manifest_record.get("error") or "") or None
    if status == "partial_success":
        warnings = manifest_record.get("warnings")
        if isinstance(warnings, list):
            normalized = [str(item).strip() for item in warnings if str(item).strip()]
            if normalized:
                return "; ".join(normalized)
    if status == "skipped_duplicate":
        duplicate_of = manifest_record.get("duplicate_of")
        if isinstance(duplicate_of, str) and duplicate_of:
            return f"duplicate_of:{duplicate_of}"
    return None


def _count_reason_map(reason_groups: Any) -> dict[str, int]:
    counts: dict[str, int] = {}
    for reasons in reason_groups:
        for reason in reasons:
            normalized = str(reason).strip()
            if not normalized:
                continue
            counts[normalized] = counts.get(normalized, 0) + 1
    return dict(sorted(counts.items()))


def _append_jsonl(path: Path, payload: dict[str, Any]) -> None:
    with path.open("a", encoding="utf-8", newline="\n") as handle:
        handle.write(json.dumps(payload, ensure_ascii=False, separators=(",", ":")) + "\n")


def _write_text(path: Path, text: str) -> None:
    path.write_text(text, encoding="utf-8")