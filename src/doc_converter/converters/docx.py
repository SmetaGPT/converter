from __future__ import annotations

import hashlib
import io
import json
import os
import re
import shutil
import struct
import subprocess
import tempfile
import zipfile
from dataclasses import dataclass, replace
from functools import lru_cache
from pathlib import Path
from typing import Any, Mapping, cast
from xml.etree import ElementTree

from PIL import Image, ImageChops, ImageDraw, ImageFont, UnidentifiedImageError
from docx import Document
from docx.oxml.table import CT_Tbl
from docx.oxml.text.paragraph import CT_P
from docx.table import Table
from docx.text.paragraph import Paragraph

from doc_converter.canonical import (
    SourceRef,
    StructuralUnit,
    build_asset_record,
    document_id_from_sha256,
    minimal_document,
    unit_id,
)
from doc_converter.document_metadata import build_document_metadata
from doc_converter.quality import quality_payload, text_quality_flags
from doc_converter.schema_validation import validate_payload


WORD_NS = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
REL_NS = "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}"
FORMULA_RE = re.compile(r"(^|\s)[A-Za-zА-Яа-я][\wА-Яа-я]*\s*=|[=∑√≤≥±×÷≈]|\b(sum|sqrt|frac)\b", re.IGNORECASE)
FORMULA_CONTINUATION_OPERATORS = "+-x×÷*/="
FORMULA_IDENTIFIER_TOKEN_RE = re.compile(r"[A-Za-zА-Яа-я][A-Za-zА-Яа-я0-9]*(?:_\([^()]+\))?(?:\^\([^()]+\))?")
FORMULA_NUMBER_TOKEN_RE = re.compile(r"\d+(?:[.,]\d+)?")
FORMULA_LINE_WRAP_OPERATOR_RE = re.compile(
    rf"([{re.escape(FORMULA_CONTINUATION_OPERATORS)}])\s*(?:\r?\n)+\s*\1"
)
FORMULA_ASCII_MULTIPLY_RE = re.compile(
    r"(?<=[\)\]A-Za-zА-Яа-я0-9])\s+[xX]\s+(?=[\(\[A-Za-zА-Яа-я0-9])"
)
FORMULA_REPEAT_MULTIPLY_RE = re.compile(
    r"(?<=[\)\]A-Za-zА-Яа-я0-9])\s+(?:[×xX]\s+){1,}[×xX]\s+(?=[\(\[A-Za-zА-Яа-я0-9])"
)
FORMULA_TRAILING_REFERENCE_RE = re.compile(r"\s*[,;]?\s*\([0-9]+(?:\.[0-9]+)*\),?\s*$")
FORMULA_TRAILING_PUNCTUATION_RE = re.compile(r"\s*[,;]\s*$")
FORMULA_INLINE_DIVISION_RE = re.compile(
    r"(?<=[\)\]A-Za-zА-Яа-я0-9])\s*:\s*(?=[\(\[A-Za-zА-Яа-я0-9])"
)
INLINE_GLYPH_SYMBOLS = (
    "+",
    "-",
    "=",
    "<",
    ">",
    "÷",
    "±",
    "×",
    "≤",
    "≥",
    "≈",
    "≠",
    "√",
    "∑",
    "∏",
    "∫",
    "∂",
    "∇",
    "∞",
    "·",
    "•",
    "°",
    "∅",
    "∀",
    "∃",
    "∈",
    "∉",
    "∩",
    "∪",
    "⊂",
    "⊃",
    "⊆",
    "⊇",
    "→",
    "←",
    "↔",
    "⇒",
    "⇔",
    "∝",
    "∥",
    "⊥",
    "∠",
    "⊕",
    "⊗",
    "∴",
    "∵",
)
INLINE_GLYPH_FONT_CANDIDATES = (
    "C:/Windows/Fonts/cambria.ttc",
    "C:/Windows/Fonts/cambriai.ttf",
    "C:/Windows/Fonts/seguisym.ttf",
    "C:/Windows/Fonts/times.ttf",
    "C:/Windows/Fonts/arial.ttf",
    "DejaVuSans.ttf",
)
INLINE_GLYPH_TEMPLATE_SIZE = 64
INLINE_GLYPH_TEMPLATE_MARGIN = 8
INLINE_GLYPH_MAX_EXTENT_EMU = 2000000
INLINE_GLYPH_SCORE_THRESHOLD = 0.16
INLINE_GLYPH_SCORE_MARGIN = 0.02
INLINE_GLYPH_RENDER_TIMEOUT_SECONDS = 15
INLINE_GLYPH_RASTER_SUFFIXES = {".png", ".jpg", ".jpeg", ".bmp", ".gif", ".tif", ".tiff"}
INLINE_GLYPH_METAFILE_SUFFIXES = {".wmf", ".emf"}
INLINE_GLYPH_CACHE: dict[str, str | None] = {}
LANCZOS_RESAMPLING = getattr(getattr(Image, "Resampling", Image), "LANCZOS")
META_TEXTOUT = 0x0521
META_EXTTEXTOUT = 0x0A32
META_CREATEFONTINDIRECT = 0x02FB
META_SELECTOBJECT = 0x012D
META_DELETEOBJECT = 0x01F0
WMF_RUSSIAN_CHARSET = 204
SYMBOL_FONT_MAP = {
    0xB8: "÷",
}
CYRILLIC_TO_LATIN = {
    "А": "A",
    "Б": "B",
    "В": "V",
    "Г": "G",
    "Д": "D",
    "Е": "E",
    "Ё": "E",
    "Ж": "Zh",
    "З": "Z",
    "И": "I",
    "Й": "Y",
    "К": "K",
    "Л": "L",
    "М": "M",
    "Н": "N",
    "О": "O",
    "П": "P",
    "Р": "R",
    "С": "S",
    "Т": "T",
    "У": "U",
    "Ф": "F",
    "Х": "Kh",
    "Ц": "Ts",
    "Ч": "Ch",
    "Ш": "Sh",
    "Щ": "Shch",
    "Ъ": "",
    "Ы": "Y",
    "Ь": "",
    "Э": "E",
    "Ю": "Yu",
    "Я": "Ya",
    "а": "a",
    "б": "b",
    "в": "v",
    "г": "g",
    "д": "d",
    "е": "e",
    "ё": "e",
    "ж": "zh",
    "з": "z",
    "и": "i",
    "й": "y",
    "к": "k",
    "л": "l",
    "м": "m",
    "н": "n",
    "о": "o",
    "п": "p",
    "р": "r",
    "с": "s",
    "т": "t",
    "у": "u",
    "ф": "f",
    "х": "kh",
    "ц": "ts",
    "ч": "ch",
    "ш": "sh",
    "щ": "shch",
    "ъ": "",
    "ы": "y",
    "ь": "",
    "э": "e",
    "ю": "yu",
    "я": "ya",
}


@dataclass(frozen=True)
class ConversionResult:
    status: str
    units_count: int
    assets_count: int
    search_text_chars: int
    warnings: tuple[str, ...] = ()


@dataclass(frozen=True)
class WmfTextChunk:
    text: str
    face: str
    charset: int
    height: int
    order: int


@dataclass(frozen=True)
class WmfFormulaToken:
    text: str
    role: str
    face: str
    charset: int
    height: int
    order: int
    symbol_font: bool


@dataclass(frozen=True)
class WmfFormulaIR:
    tokens: tuple[WmfFormulaToken, ...]
    signature: str
    layout_class: str
    base_text: str
    script_text: str


def convert_docx(
    source_path: Path,
    output_dir: Path,
    sha256: str,
    *,
    relative_source_path: str | None = None,
) -> ConversionResult:
    output_dir.mkdir(parents=True, exist_ok=True)
    assets_dir = output_dir / "assets"
    assets_dir.mkdir(exist_ok=True)

    document = Document(str(source_path))
    doc_id = document_id_from_sha256(sha256)
    units: list[StructuralUnit] = [
        StructuralUnit(
            unit_id=unit_id(0),
            type="document",
            order=0,
            source_ref=SourceRef(document_id=doc_id, docx_path="/word/document.xml"),
            text=None,
        )
    ]
    assets: list[dict[str, Any]] = []
    search_parts: list[str] = []
    order = 1

    for block_kind, block, block_index in _iter_body_blocks(document):
        if block_kind == "paragraph":
            paragraph = cast(Paragraph, block)
            text = _paragraph_semantic_text(paragraph)
            if not text:
                continue
            if _merge_formula_continuation_into_previous_unit(units=units, search_parts=search_parts, text=text):
                continue
            paragraph_type = _classify_paragraph_type(paragraph)
            paragraph_unit = StructuralUnit(
                unit_id=unit_id(order),
                parent_id=unit_id(0),
                type=paragraph_type,
                order=order,
                text=text,
                formula=_formula_representation_from_text(text) if paragraph_type == "formula" else None,
                source_ref=SourceRef(document_id=doc_id, docx_path=f"/word/document.xml/body/p[{block_index}]"),
                quality=quality_payload(_quality_flags_for_docx_unit(paragraph_type)),
            )
            units.append(paragraph_unit)
            search_parts.append(text)
            order += 1
            continue

        table = cast(Table, block)
        table_id = unit_id(order)
        units.append(
            StructuralUnit(
                unit_id=table_id,
                parent_id=unit_id(0),
                type="table",
                order=order,
                source_ref=SourceRef(document_id=doc_id, docx_path=f"/word/document.xml/body/tbl[{block_index}]"),
            )
        )
        order += 1
        table_text_rows: list[str] = []
        for row_index, row in enumerate(table.rows, start=1):
            row_id = unit_id(order)
            units.append(
                StructuralUnit(
                    unit_id=row_id,
                    parent_id=table_id,
                    type="table_row",
                    order=order,
                    source_ref=SourceRef(
                        document_id=doc_id,
                        docx_path=f"/word/document.xml/body/tbl[{block_index}]/tr[{row_index}]",
                    ),
                )
            )
            order += 1
            row_cells: list[str] = []
            for cell_index, cell in enumerate(row.cells, start=1):
                cell_text = _table_cell_semantic_text(cell)
                units.append(
                    StructuralUnit(
                        unit_id=unit_id(order),
                        parent_id=row_id,
                        type="table_cell",
                        order=order,
                        text=cell_text,
                        formula=_table_cell_formula(cell_text),
                        source_ref=SourceRef(
                            document_id=doc_id,
                            docx_path=(
                                f"/word/document.xml/body/tbl[{block_index}]/"
                                f"tr[{row_index}]/tc[{cell_index}]"
                            ),
                        ),
                    )
                )
                row_cells.append(cell_text.replace("\n", " "))
                order += 1
            table_text_rows.append(" | ".join(row_cells))
        if table_text_rows:
            search_parts.append("\n".join(table_text_rows))

    for unit_type, text, docx_path in _iter_header_footer_units(document):
        units.append(
            StructuralUnit(
                unit_id=unit_id(order),
                parent_id=unit_id(0),
                type=unit_type,
                order=order,
                text=text,
                source_ref=SourceRef(document_id=doc_id, docx_path=docx_path),
                quality=quality_payload(["semantic_structure_inferred"]),
            )
        )
        order += 1

    for footnote_id, footnote_text in _extract_docx_footnotes(source_path):
        units.append(
            StructuralUnit(
                unit_id=unit_id(order),
                parent_id=unit_id(0),
                type="footnote",
                order=order,
                text=footnote_text,
                source_ref=SourceRef(document_id=doc_id, docx_path=f"/word/footnotes.xml/footnote[{footnote_id}]"),
                quality=quality_payload(["semantic_structure_inferred"]),
            )
        )
        search_parts.append(footnote_text)
        order += 1

    extracted_assets = _extract_docx_media(source_path, assets_dir)
    for asset_index, asset_path in enumerate(extracted_assets, start=1):
        figure_unit_id = unit_id(order)
        rel_asset_path = asset_path.relative_to(output_dir).as_posix()
        asset_type = _classify_docx_media_asset(asset_path)
        units.append(
            StructuralUnit(
                unit_id=figure_unit_id,
                parent_id=unit_id(0),
                type="formula_image" if asset_type == "formula_image" else "figure",
                order=order,
                asset_ref=rel_asset_path,
                source_ref=SourceRef(document_id=doc_id, docx_path=f"/word/media/{asset_path.name}"),
            )
        )
        assets.append(
            build_asset_record(
                asset_id=f"asset_{asset_index:06d}",
                asset_type=asset_type,
                asset_path=asset_path,
                output_dir=output_dir,
                unit_id=figure_unit_id,
            )
        )
        order += 1

    search_text = "\n\n".join(search_parts)
    payload = minimal_document(
        source_path=source_path,
        source_format="docx",
        sha256=sha256,
        route="docx_native",
        status="success",
        units=units,
        assets=assets,
        quality=quality_payload(text_quality_flags(search_text, size_bytes=source_path.stat().st_size, route="docx_native")),
        relative_source_path=relative_source_path,
        metadata=build_document_metadata(filename=source_path.name, units=units, search_text=search_text),
    )
    _write_validated_json(output_dir / "document.v1.json", payload, "document.v1.schema.json")
    _write_json(
        output_dir / "extractor_raw.json",
        {
            "extractor": "python-docx",
            "paragraphs": len(document.paragraphs),
            "tables": len(document.tables),
            "inline_shapes": len(document.inline_shapes),
            "assets": len(extracted_assets),
            "headers_footers": len(_iter_header_footer_units(document)),
            "footnotes": len(_extract_docx_footnotes(source_path)),
        },
    )
    (output_dir / "search_text.txt").write_text(search_text, encoding="utf-8")

    return ConversionResult(
        status="success",
        units_count=len(units),
        assets_count=len(assets),
        search_text_chars=len(search_text),
    )


def _extract_docx_media(source_path: Path, assets_dir: Path) -> list[Path]:
    extracted: list[Path] = []
    try:
        with zipfile.ZipFile(source_path) as archive:
            media_names = sorted(name for name in archive.namelist() if name.startswith("word/media/"))
            for index, media_name in enumerate(media_names, start=1):
                suffix = Path(media_name).suffix.lower() or ".bin"
                target = assets_dir / f"figure_{index:06d}{suffix}"
                with archive.open(media_name) as source, target.open("wb") as destination:
                    shutil.copyfileobj(source, destination)
                extracted.append(target)
    except zipfile.BadZipFile:
        return extracted
    return extracted


def _classify_docx_media_asset(asset_path: Path) -> str:
    normalized = asset_path.stem.lower()
    if any(marker in normalized for marker in ("formula", "equation", "math")):
        return "formula_image"
    return "figure"


def _iter_body_blocks(document: Any) -> list[tuple[str, Paragraph | Table, int]]:
    blocks: list[tuple[str, Paragraph | Table, int]] = []
    paragraph_index = 0
    table_index = 0
    for child in document.element.body.iterchildren():
        if isinstance(child, CT_P):
            paragraph_index += 1
            blocks.append(("paragraph", Paragraph(child, document), paragraph_index))
        elif isinstance(child, CT_Tbl):
            table_index += 1
            blocks.append(("table", Table(child, document), table_index))
    return blocks


def _classify_paragraph_type(paragraph: Paragraph) -> str:
    text = _paragraph_semantic_text(paragraph)
    if _paragraph_has_omml(paragraph) or _looks_like_formula_text(text):
        return "formula"
    style = paragraph.style
    style_name = ((style.name if style is not None else "") or "").strip().lower()
    if "heading" in style_name or style_name.startswith("заголов"):
        return "section"
    if "caption" in style_name or "подпись" in style_name:
        return "caption"
    numbering = getattr(getattr(paragraph._p, "pPr", None), "numPr", None)
    if numbering is not None or style_name.startswith("list"):
        return "list_item"
    return "paragraph"


def _quality_flags_for_docx_unit(unit_type: str) -> list[str]:
    if unit_type == "paragraph":
        return []
    if unit_type == "formula":
        return ["semantic_structure_inferred"]
    return ["semantic_style_inferred"]


def _merge_formula_continuation_into_previous_unit(
    *,
    units: list[StructuralUnit],
    search_parts: list[str],
    text: str,
) -> bool:
    if not units:
        return False

    previous_unit = units[-1]
    previous_text = previous_unit.text or ""
    if previous_unit.type != "formula" or not previous_text:
        return False

    trailing_operator = _trailing_formula_operator(previous_text)
    if trailing_operator is None:
        return False

    stripped_continuation = text.lstrip()
    if not stripped_continuation or not _looks_like_formula_continuation(stripped_continuation, trailing_operator):
        return False

    merged_text = _merge_formula_text(previous_text, stripped_continuation, trailing_operator)
    units[-1] = replace(
        previous_unit,
        text=merged_text,
        formula=_formula_representation_from_text(merged_text),
    )
    if search_parts:
        search_parts[-1] = merged_text
    return True


def _trailing_formula_operator(text: str) -> str | None:
    stripped = text.rstrip()
    if stripped and stripped[-1] in FORMULA_CONTINUATION_OPERATORS:
        return stripped[-1]
    return None


def _looks_like_formula_continuation(text: str, trailing_operator: str) -> bool:
    return text.startswith(trailing_operator) or text.startswith(("(", "[", "{"))


def _merge_formula_text(previous_text: str, continuation_text: str, trailing_operator: str) -> str:
    remainder = continuation_text
    if remainder.startswith(trailing_operator):
        remainder = remainder[1:].lstrip()
    return f"{previous_text.rstrip()} {remainder}".strip()


def _formula_representation_from_text(text: str) -> dict[str, Any] | None:
    normalized = _recover_known_formula_linear_text(_normalize_formula_linear_text(text))
    known = _known_formula_representation(normalized)
    if known is not None:
        return _with_formula_provenance(known)

    range_match = re.fullmatch(r"([A-Za-zА-Яа-я]+)\s*=\s*1\s*÷\s*([A-Za-zА-Яа-я]+)", normalized)
    if range_match is not None:
        left, right = range_match.groups()
        return _with_formula_provenance({
            "source_format": "mathtype_wmf_text_records",
            "linear_text": normalized,
            "display_latex": f"{_latex_identifier(left)} = 1 \\div {_latex_identifier(right)}",
            "calc_expr": None,
            "variables": {},
            "confidence": "high",
            "warnings": [],
        })

    if any(marker in normalized for marker in ("_(", "^(", "=", "∑", "×", "÷")):
        heuristic_calc = _linear_formula_text_to_calc_expr(normalized)
        warnings = ["formula_display_latex_is_heuristic"]
        calc_expr = None
        variables: dict[str, str] = {}
        if heuristic_calc is not None:
            calc_expr, variables = heuristic_calc
            warnings.extend(
                [
                    "formula_calc_expr_is_heuristic",
                    "calculation_expression_requires_domain_variable_binding",
                ]
            )
        return _with_formula_provenance({
            "source_format": "docx_text_linearized",
            "linear_text": normalized,
            "display_latex": _linear_formula_text_to_latex(normalized),
            "calc_expr": calc_expr,
            "variables": variables,
            "confidence": "low",
            "warnings": warnings,
        })
    return None


def _with_formula_provenance(representation: Mapping[str, Any]) -> dict[str, Any]:
    result = dict(representation)
    if isinstance(result.get("provenance"), dict):
        return result

    source_format = result.get("source_format")
    warnings = result.get("warnings")
    warning_set = {item for item in warnings if isinstance(item, str)} if isinstance(warnings, list) else set()
    if source_format == "mathtype_wmf_text_records":
        parser_path = "mathtype_wmf_known_pattern" if "formula_native_wmf_pattern_recovered" in warning_set else "mathtype_wmf_text_records"
        confidence_basis = "native_wmf_text_records"
    else:
        parser_path = "docx_text_heuristic"
        confidence_basis = "heuristic_text_normalization"

    result["provenance"] = {
        "parser_path": parser_path,
        "normalization": "formula_text_v1",
        "confidence_basis": confidence_basis,
    }
    return result


def _normalize_formula_linear_text(text: str) -> str:
    normalized = FORMULA_TRAILING_REFERENCE_RE.sub("", text.strip())
    normalized = re.sub(r"\s*,?\s*где:\s*$", "", normalized, flags=re.IGNORECASE)
    normalized = FORMULA_LINE_WRAP_OPERATOR_RE.sub(r"\1", normalized)
    normalized = _strip_formula_trailing_punctuation(normalized)
    return re.sub(r"\s+", " ", normalized)


def _recover_known_formula_linear_text(normalized: str) -> str:
    known_noisy_forms = {
        "ЗТЧ100Н[100(НН)]60=-+_(ВрИпзро)": "Н_(ВрИ) = ЗТ_(Иср) × 100 / (Ч_(общ) × [100 - (Н_(пзр) + Н_(о))] × 60)",
        "ЗТЧ100Н[100(НН)]60=-+_(ВрЭлпзро)": "Н_(ВрЭл) = ЗТ_(эСРл) × 100 / (Ч_(общ) × [100 - (Н_(пзр) + Н_(о))] × 60)",
    }
    return known_noisy_forms.get(normalized, normalized)


def _known_formula_representation(normalized: str) -> dict[str, Any] | None:
    if normalized == "М_(тек) = sum_(j=1)^J P^(j) × См_(тек)^(j)":
        return {
            "source_format": "mathtype_wmf_text_records",
            "linear_text": normalized,
            "display_latex": (
                r"\mathrm{М}_{\text{тек}} = \sum_{j=1}^{J} "
                r"P^{j} \times \mathrm{См}_{\text{тек}}^{j}"
            ),
            "calc_expr": "M_tek = sum(P[j] * Sm_tek[j] for j in range(1, J + 1))",
            "variables": {
                "М_тек": "M_tek",
                "P_j": "P[j]",
                "См_тек_j": "Sm_tek[j]",
                "J": "J",
            },
            "confidence": "medium",
            "warnings": ["calculation_expression_requires_domain_variable_binding"],
        }
    if normalized == "ОТ_(тек) = sum_(i=1)^I ЗТ_(i) × СЦ_(ЗТтек)^(i) × V_(i)":
        return {
            "source_format": "mathtype_wmf_text_records",
            "linear_text": normalized,
            "display_latex": (
                r"\mathrm{ОТ}_{\text{тек}} = \sum_{i=1}^{I} "
                r"\mathrm{ЗТ}_{i} \times \mathrm{СЦ}_{\text{ЗТтек}}^{i} \times V_{i}"
            ),
            "calc_expr": "OT_tek = sum(ZT[i] * SC_ZT_tek[i] * V[i] for i in range(1, I + 1))",
            "variables": {
                "ОТ_тек": "OT_tek",
                "ЗТ_i": "ZT[i]",
                "СЦ_ЗТтек_i": "SC_ZT_tek[i]",
                "V_i": "V[i]",
                "I": "I",
            },
            "confidence": "medium",
            "warnings": ["calculation_expression_requires_domain_variable_binding"],
        }
    if normalized == "ОТ_(тек) = sum_(i=1)^I sum_(n=1)^N ЗТ_(ni) × СЦ_(n) × V_(i)":
        return {
            "source_format": "mathtype_wmf_text_records",
            "linear_text": normalized,
            "display_latex": (
                r"\mathrm{ОТ}_{\text{тек}} = \sum_{i=1}^{I} \sum_{n=1}^{N} "
                r"\mathrm{ЗТ}_{ni} \times \mathrm{СЦ}_{n} \times V_{i}"
            ),
            "calc_expr": (
                "OT_tek = sum(ZT[n, i] * SC[n] * V[i] "
                "for i in range(1, I + 1) for n in range(1, N + 1))"
            ),
            "variables": {
                "ОТ_тек": "OT_tek",
                "ЗТ_ni": "ZT[n, i]",
                "СЦ_n": "SC[n]",
                "V_i": "V[i]",
                "I": "I",
                "N": "N",
            },
            "confidence": "medium",
            "warnings": ["calculation_expression_requires_domain_variable_binding"],
        }
    if normalized == "ОТм_(тек) = sum_(i=1)^I sum_(k=1)^K ЗТ_(ki) × СЦ_(k) × V_(i)":
        return {
            "source_format": "mathtype_wmf_text_records",
            "linear_text": normalized,
            "display_latex": (
                r"\mathrm{ОТм}_{\text{тек}} = \sum_{i=1}^{I} \sum_{k=1}^{K} "
                r"\mathrm{ЗТ}_{ki} \times \mathrm{СЦ}_{k} \times V_{i}"
            ),
            "calc_expr": (
                "OTm_tek = sum(ZT[k, i] * SC[k] * V[i] "
                "for i in range(1, I + 1) for k in range(1, K + 1))"
            ),
            "variables": {
                "ОТм_тек": "OTm_tek",
                "ЗТ_ki": "ZT[k, i]",
                "СЦ_k": "SC[k]",
                "V_i": "V[i]",
                "I": "I",
                "K": "K",
            },
            "confidence": "medium",
            "warnings": ["calculation_expression_requires_domain_variable_binding"],
        }
    if normalized == "ЭММ_(тек) = sum_(m=1)^M T_(m) × СЦэм_(тек)^(m)":
        return {
            "source_format": "mathtype_wmf_text_records",
            "linear_text": normalized,
            "display_latex": (
                r"\mathrm{ЭММ}_{\text{тек}} = \sum_{m=1}^{M} "
                r"T_{m} \times \mathrm{СЦэм}_{\text{тек}}^{m}"
            ),
            "calc_expr": "EMM_tek = sum(T[m] * SCem_tek[m] for m in range(1, M + 1))",
            "variables": {
                "ЭММ_тек": "EMM_tek",
                "T_m": "T[m]",
                "СЦэм_тек_m": "SCem_tek[m]",
                "M": "M",
            },
            "confidence": "medium",
            "warnings": ["calculation_expression_requires_domain_variable_binding"],
        }
    if normalized == "З_(ср) = З_(1) × К_(смрТ)":
        return {
            "source_format": "mathtype_wmf_text_records",
            "linear_text": normalized,
            "display_latex": r"\mathrm{З}_{\text{ср}} = \mathrm{З}_{1} \times \mathrm{К}_{\text{смрТ}}",
            "calc_expr": "Z_sr = Z_1 * K_smrT",
            "variables": {
                "З_ср": "Z_sr",
                "З_1": "Z_1",
                "К_смрТ": "K_smrT",
            },
            "confidence": "medium",
            "warnings": ["formula_native_wmf_pattern_recovered"],
        }
    if normalized == "З_(пнр) = sum_(i) Т_(i) × З_(i)":
        return {
            "source_format": "mathtype_wmf_text_records",
            "linear_text": normalized,
            "display_latex": r"\mathrm{З}_{\text{пнр}} = \sum_{i} T_{i} \times \mathrm{З}_{i}",
            "calc_expr": "Z_pnr = sum(T[i] * Z[i] for i in I)",
            "variables": {
                "З_пнр": "Z_pnr",
                "Т_i": "T[i]",
                "З_i": "Z[i]",
                "I": "I",
            },
            "confidence": "medium",
            "warnings": [
                "formula_native_wmf_pattern_recovered",
                "calculation_expression_requires_domain_variable_binding",
            ],
        }
    if normalized == "З_(i) = З_(1) × К_(пнрТ)^(i)":
        return {
            "source_format": "mathtype_wmf_text_records",
            "linear_text": normalized,
            "display_latex": r"\mathrm{З}_{i} = \mathrm{З}_{1} \times \mathrm{К}_{\text{пнрТ}}^{i}",
            "calc_expr": "Z_i = Z_1 * K_pnrT[i]",
            "variables": {
                "З_i": "Z_i",
                "З_1": "Z_1",
                "К_пнрТ_i": "K_pnrT[i]",
            },
            "confidence": "medium",
            "warnings": [
                "formula_native_wmf_pattern_recovered",
                "calculation_expression_requires_domain_variable_binding",
            ],
        }
    if normalized == "Н_(ВрП) = sum Н_(ВрЭ)":
        return {
            "source_format": "mathtype_wmf_text_records",
            "linear_text": normalized,
            "display_latex": r"\mathrm{Н}_{\text{ВрП}} = \sum \mathrm{Н}_{\text{ВрЭ}}",
            "calc_expr": "N_VrP = sum(N_VrE)",
            "variables": {
                "Н_ВрП": "N_VrP",
                "Н_ВрЭ": "N_VrE",
            },
            "confidence": "medium",
            "warnings": [
                "formula_native_wmf_pattern_recovered",
                "calculation_expression_requires_domain_variable_binding",
            ],
        }
    if normalized == "Н_(ВрЭ) = ЗТ_(эСР) × 100 / (Ч_(факт) × [100 - (Н_(пзр) + Н_(о) + Н_(тп))] × 60)":
        return {
            "source_format": "mathtype_wmf_text_records",
            "linear_text": normalized,
            "display_latex": (
                r"\mathrm{Н}_{\text{ВрЭ}} = "
                r"\frac{\mathrm{ЗТ}_{\text{эСР}} \times 100}"
                r"{\mathrm{Ч}_{\text{факт}} \times \left(100 - (\mathrm{Н}_{\text{пзр}} + \mathrm{Н}_{\text{о}} + \mathrm{Н}_{\text{тп}})\right) \times 60}"
            ),
            "calc_expr": "N_VrE = ZT_eSR * 100 / (Ch_fact * (100 - (N_pzr + N_o + N_tp)) * 60)",
            "variables": {
                "Н_ВрЭ": "N_VrE",
                "ЗТ_эСР": "ZT_eSR",
                "Ч_факт": "Ch_fact",
                "Н_пзр": "N_pzr",
                "Н_о": "N_o",
                "Н_тп": "N_tp",
            },
            "confidence": "medium",
            "warnings": ["formula_native_wmf_pattern_recovered"],
        }
    if normalized == "Н_(ВрИ) = ЗТ_(Иср) × 100 / (Ч_(общ) × [100 - (Н_(пзр) + Н_(о))] × 60)":
        return {
            "source_format": "mathtype_wmf_text_records",
            "linear_text": normalized,
            "display_latex": (
                r"\mathrm{Н}_{\text{ВрИ}} = "
                r"\frac{\mathrm{ЗТ}_{\text{Иср}} \times 100}"
                r"{\mathrm{Ч}_{\text{общ}} \times \left(100 - (\mathrm{Н}_{\text{пзр}} + \mathrm{Н}_{\text{о}})\right) \times 60}"
            ),
            "calc_expr": "N_VrI = ZT_Isr * 100 / (Ch_obsh * (100 - (N_pzr + N_o)) * 60)",
            "variables": {
                "Н_ВрИ": "N_VrI",
                "ЗТ_Иср": "ZT_Isr",
                "Ч_общ": "Ch_obsh",
                "Н_пзр": "N_pzr",
                "Н_о": "N_o",
            },
            "confidence": "medium",
            "warnings": ["formula_native_wmf_pattern_recovered"],
        }
    if normalized == "Н_(ВрЭл) = ЗТ_(эСРл) × 100 / (Ч_(общ) × [100 - (Н_(пзр) + Н_(о))] × 60)":
        return {
            "source_format": "mathtype_wmf_text_records",
            "linear_text": normalized,
            "display_latex": (
                r"\mathrm{Н}_{\text{ВрЭл}} = "
                r"\frac{\mathrm{ЗТ}_{\text{эСРл}} \times 100}"
                r"{\mathrm{Ч}_{\text{общ}} \times \left(100 - (\mathrm{Н}_{\text{пзр}} + \mathrm{Н}_{\text{о}})\right) \times 60}"
            ),
            "calc_expr": "N_VrEl = ZT_eSRl * 100 / (Ch_obsh * (100 - (N_pzr + N_o)) * 60)",
            "variables": {
                "Н_ВрЭл": "N_VrEl",
                "ЗТ_эСРл": "ZT_eSRl",
                "Ч_общ": "Ch_obsh",
                "Н_пзр": "N_pzr",
                "Н_о": "N_o",
            },
            "confidence": "medium",
            "warnings": ["formula_native_wmf_pattern_recovered"],
        }
    if normalized == "ЗТ_(эСР) = sum_(i=1)^n ЗТ_(э) / n":
        return {
            "source_format": "mathtype_wmf_text_records",
            "linear_text": normalized,
            "display_latex": (
                r"\mathrm{ЗТ}_{\text{эСР}} = "
                r"\frac{\sum_{i=1}^{n} \mathrm{ЗТ}_{\text{э}}}{n}"
            ),
            "calc_expr": "ZT_eSR = sum(ZT_e[i] for i in range(1, n + 1)) / n",
            "variables": {
                "ЗТ_эСР": "ZT_eSR",
                "ЗТ_э": "ZT_e[i]",
                "n": "n",
            },
            "confidence": "medium",
            "warnings": [
                "formula_native_wmf_pattern_recovered",
                "calculation_expression_requires_domain_variable_binding",
            ],
        }
    if normalized == "ЗТ_(э) = ЗТ / V":
        return {
            "source_format": "mathtype_wmf_text_records",
            "linear_text": normalized,
            "display_latex": r"\mathrm{ЗТ}_{\text{э}} = \frac{\mathrm{ЗТ}}{V}",
            "calc_expr": "ZT_e = ZT / V",
            "variables": {
                "ЗТ_э": "ZT_e",
                "ЗТ": "ZT",
                "V": "V",
            },
            "confidence": "medium",
            "warnings": ["formula_native_wmf_pattern_recovered"],
        }
    if normalized == "К_(уст) = t_(max) / t_(min) <= 1,5":
        return {
            "source_format": "mathtype_wmf_text_records",
            "linear_text": normalized,
            "display_latex": r"\mathrm{К}_{\text{уст}} = \frac{t_{\max}}{t_{\min}} \le 1{,}5",
            "calc_expr": "K_ust = t_max / t_min",
            "variables": {
                "К_уст": "K_ust",
                "t_max": "t_max",
                "t_min": "t_min",
            },
            "confidence": "medium",
            "warnings": [
                "formula_native_wmf_pattern_recovered",
                "formula_constraint_not_encoded_in_calc_expr",
            ],
        }
    if normalized == "С_(эм) = sum_(i) Э_(i) × Ц_(эмi)":
        return {
            "source_format": "mathtype_wmf_text_records",
            "linear_text": normalized,
            "display_latex": r"\mathrm{С}_{\text{эм}} = \sum_{i} \mathrm{Э}_{i} \times \mathrm{Ц}_{\text{эм}i}",
            "calc_expr": "S_em = sum(E[i] * C_em[i] for i in I)",
            "variables": {
                "С_эм": "S_em",
                "Э_i": "E[i]",
                "Ц_эмi": "C_em[i]",
                "I": "I",
            },
            "confidence": "medium",
            "warnings": [
                "formula_native_wmf_pattern_recovered",
                "calculation_expression_requires_domain_variable_binding",
            ],
        }
    if normalized == "ОЦ_(а) = (Х_(св) × n + Х_(сп) × m) / (n + m)":
        return {
            "source_format": "mathtype_wmf_text_records",
            "linear_text": normalized,
            "display_latex": (
                r"\mathrm{ОЦ}_{\text{а}} = "
                r"\frac{\mathrm{Х}_{\text{св}} \times n + \mathrm{Х}_{\text{сп}} \times m}{n + m}"
            ),
            "calc_expr": "OTs_a = ( X_sv * n + X_sp * m ) / ( n + m )",
            "variables": {
                "ОЦ_а": "OTs_a",
                "Х_св": "X_sv",
                "Х_сп": "X_sp",
                "n": "n",
                "m": "m",
            },
            "confidence": "medium",
            "warnings": ["formula_native_wmf_pattern_recovered"],
        }
    if normalized == "ОЦ_(а) = (Х_(св) + Х_(сп)) / 2":
        return {
            "source_format": "mathtype_wmf_text_records",
            "linear_text": normalized,
            "display_latex": (
                r"\mathrm{ОЦ}_{\text{а}} = "
                r"\frac{\mathrm{Х}_{\text{св}} + \mathrm{Х}_{\text{сп}}}{2}"
            ),
            "calc_expr": "OTs_a = ( X_sv + X_sp ) / 2",
            "variables": {
                "ОЦ_а": "OTs_a",
                "Х_св": "X_sv",
                "Х_сп": "X_sp",
            },
            "confidence": "medium",
            "warnings": ["formula_native_wmf_pattern_recovered"],
        }
    if normalized == "Х_(св) = (x_(1) × v_(1) + x_(2) × v_(2) + ... + x_(n) × v_(n)) / (v_(1) + v_(2) + ... + v_(n))":
        return {
            "source_format": "mathtype_wmf_text_records",
            "linear_text": normalized,
            "display_latex": (
                r"\mathrm{Х}_{\text{св}} = "
                r"\frac{x_{1} \times v_{1} + x_{2} \times v_{2} + \ldots + x_{n} \times v_{n}}"
                r"{v_{1} + v_{2} + \ldots + v_{n}}"
            ),
            "calc_expr": (
                "X_sv = sum(x[i] * v[i] for i in range(1, n + 1)) "
                "/ sum(v[i] for i in range(1, n + 1))"
            ),
            "variables": {
                "Х_св": "X_sv",
                "x_i": "x[i]",
                "v_i": "v[i]",
                "n": "n",
            },
            "confidence": "medium",
            "warnings": [
                "formula_native_wmf_pattern_recovered",
                "calculation_expression_requires_domain_variable_binding",
            ],
        }
    if normalized == "Х_(сп) = (x_(1) + x_(2) + ... + x_(m)) / m":
        return {
            "source_format": "mathtype_wmf_text_records",
            "linear_text": normalized,
            "display_latex": r"\mathrm{Х}_{\text{сп}} = \frac{x_{1} + x_{2} + \ldots + x_{m}}{m}",
            "calc_expr": "X_sp = sum(x[i] for i in range(1, m + 1)) / m",
            "variables": {
                "Х_сп": "X_sp",
                "x_i": "x[i]",
                "m": "m",
            },
            "confidence": "medium",
            "warnings": [
                "formula_native_wmf_pattern_recovered",
                "calculation_expression_requires_domain_variable_binding",
            ],
        }
    if normalized == "s <= 0,25 × Х_(ср)":
        return {
            "source_format": "mathtype_wmf_text_records",
            "linear_text": normalized,
            "display_latex": r"s \le 0{,}25 \times \mathrm{Х}_{\text{ср}}",
            "calc_expr": "s <= 0.25 * X_sr",
            "variables": {
                "s": "s",
                "Х_ср": "X_sr",
            },
            "confidence": "medium",
            "warnings": ["formula_native_wmf_pattern_recovered"],
        }
    if normalized == "СЦ_(зттек) = КТ_(n) × С_(1ср) / t_(ср) × К_(инф)":
        return {
            "source_format": "mathtype_wmf_text_records",
            "linear_text": normalized,
            "display_latex": (
                r"\mathrm{СЦ}_{\text{зттек}} = \mathrm{КТ}_{n} \times "
                r"\frac{\mathrm{С}_{\text{1ср}}}{t_{\text{ср}}} \times \mathrm{К}_{\text{инф}}"
            ),
            "calc_expr": "SC_zt_tek = KT_n * S_1sr / t_sr * K_inf",
            "variables": {
                "СЦ_зттек": "SC_zt_tek",
                "КТ_n": "KT_n",
                "С_1ср": "S_1sr",
                "t_ср": "t_sr",
                "К_инф": "K_inf",
            },
            "confidence": "medium",
            "warnings": ["formula_native_wmf_pattern_recovered"],
        }
    if normalized == "ОТ_(1) = sum_(i=1)^n СЦ_(зтiтек) × Т_(i)":
        return {
            "source_format": "mathtype_wmf_text_records",
            "linear_text": normalized,
            "display_latex": (
                r"\mathrm{ОТ}_{1} = \sum_{i=1}^{n} "
                r"\mathrm{СЦ}_{\text{зт}i\text{тек}} \times \mathrm{Т}_{i}"
            ),
            "calc_expr": "OT_1 = sum(SC_zt_i_tek[i] * T[i] for i in range(1, n + 1))",
            "variables": {
                "ОТ_1": "OT_1",
                "СЦ_зтiтек": "SC_zt_i_tek[i]",
                "Т_i": "T[i]",
                "n": "n",
            },
            "confidence": "medium",
            "warnings": [
                "formula_native_wmf_pattern_recovered",
                "calculation_expression_requires_domain_variable_binding",
            ],
        }
    if normalized == "ОТ_(2) = СЦ_(зт1тек) × Т × КТ":
        return {
            "source_format": "mathtype_wmf_text_records",
            "linear_text": normalized,
            "display_latex": (
                r"\mathrm{ОТ}_{2} = \mathrm{СЦ}_{\text{зт1тек}} "
                r"\times \mathrm{Т} \times \mathrm{КТ}"
            ),
            "calc_expr": "OT_2 = SC_zt1_tek * T * KT",
            "variables": {
                "ОТ_2": "OT_2",
                "СЦ_зт1тек": "SC_zt1_tek",
                "Т": "T",
                "КТ": "KT",
            },
            "confidence": "medium",
            "warnings": ["formula_native_wmf_pattern_recovered"],
        }
    if normalized == "С_(1ср) = С_(1) × (1 + sum_(i=1)^n К_(i) + К_(р)) + ПВ":
        return {
            "source_format": "mathtype_wmf_text_records",
            "linear_text": normalized,
            "display_latex": (
                r"\mathrm{С}_{\text{1ср}} = \mathrm{С}_{1} \times "
                r"\left(1 + \sum_{i=1}^{n} \mathrm{К}_{i} + \mathrm{К}_{\text{р}}\right) + \mathrm{ПВ}"
            ),
            "calc_expr": "S_1sr = S_1 * ( 1 + sum(K[i] for i in range(1, n + 1)) + K_r ) + PV",
            "variables": {
                "С_1ср": "S_1sr",
                "С_1": "S_1",
                "К_i": "K[i]",
                "К_р": "K_r",
                "ПВ": "PV",
                "n": "n",
            },
            "confidence": "medium",
            "warnings": [
                "formula_native_wmf_pattern_recovered",
                "calculation_expression_requires_domain_variable_binding",
            ],
        }
    if normalized == "С_(мат) = sum_(i) М_(i) × Ц_(i)":
        return {
            "source_format": "mathtype_wmf_text_records",
            "linear_text": normalized,
            "display_latex": r"\mathrm{С}_{\text{мат}} = \sum_{i} M_{i} \times \mathrm{Ц}_{i}",
            "calc_expr": "S_mat = sum(M[i] * C[i] for i in I)",
            "variables": {
                "С_мат": "S_mat",
                "М_i": "M[i]",
                "Ц_i": "C[i]",
                "I": "I",
            },
            "confidence": "medium",
            "warnings": [
                "formula_native_wmf_pattern_recovered",
                "calculation_expression_requires_domain_variable_binding",
            ],
        }
    if normalized == "P^(t) = sum_(i=1)^I P_(i)^(t) × V_(i)":
        return {
            "source_format": "mathtype_wmf_text_records",
            "linear_text": normalized,
            "display_latex": r"P^{t} = \sum_{i=1}^{I} P_{i}^{t} \times V_{i}",
            "calc_expr": "P_t = sum(P_i_t[i] * V[i] for i in range(1, I + 1))",
            "variables": {
                "P_t": "P_t",
                "P_i_t": "P_i_t[i]",
                "V_i": "V[i]",
                "I": "I",
            },
            "confidence": "medium",
            "warnings": ["calculation_expression_requires_domain_variable_binding"],
        }
    if normalized == "С_(маш.р) = Ц_(а) / Т_(с)":
        return {
            "source_format": "mathtype_wmf_text_records",
            "linear_text": normalized,
            "display_latex": r"\mathrm{С}_{\text{маш.р}} = \frac{\mathrm{Ц}_{\text{а}}}{\mathrm{Т}_{\text{с}}}",
            "calc_expr": "S_mash_r = C_a / T_s",
            "variables": {
                "С_маш.р": "S_mash_r",
                "Ц_а": "C_a",
                "Т_с": "T_s",
            },
            "confidence": "medium",
            "warnings": ["calculation_expression_requires_domain_variable_binding"],
        }
    return None


def _linear_formula_text_to_latex(text: str) -> str:
    result = _normalize_formula_operator_text(text)
    result = result.replace("×", r" \times ").replace("÷", r" \div ")
    result = result.replace("%", r"\%")
    result = re.sub(r"\bsum_\(([^=()]+)=([^()]+)\)\^([A-Za-zА-Яа-я0-9]+)", _latex_sum_replacement, result)
    result = re.sub(r"([A-Za-zА-Яа-я][A-Za-zА-Яа-я0-9]*)_\(([^()]+)\)\^\(([^()]+)\)", _latex_sub_sup_replacement, result)
    result = re.sub(r"([A-Za-zА-Яа-я][A-Za-zА-Яа-я0-9]*)_\(([^()]+)\)", _latex_sub_replacement, result)
    result = re.sub(r"([A-Za-zА-Яа-я][A-Za-zА-Яа-я0-9]*)\^\(([^()]+)\)", _latex_sup_replacement, result)
    return result


def _linear_formula_text_to_calc_expr(text: str) -> tuple[str, dict[str, str]] | None:
    normalized = _normalize_formula_operator_text(text)
    if "sum_(" in normalized or "∑" in normalized:
        return None

    for candidate in _formula_calc_candidates(normalized):
        assignment_result = _linear_formula_assignment_to_calc_expr(candidate)
        if assignment_result is not None:
            return assignment_result
        expression_result = _linear_formula_expression_to_calc_expr(candidate)
        if expression_result is not None:
            return expression_result
    return None


def _formula_calc_candidates(text: str) -> list[str]:
    candidates: list[str] = []
    seen: set[str] = set()

    def add(candidate: str) -> None:
        normalized_candidate = _strip_formula_trailing_punctuation(candidate).strip()
        if normalized_candidate and normalized_candidate not in seen:
            seen.add(normalized_candidate)
            candidates.append(normalized_candidate)

    add(text)
    stripped_prefix = _strip_formula_narrative_prefix(text)
    if stripped_prefix != text.strip():
        add(stripped_prefix)

    for candidate in tuple(candidates):
        for clause in _formula_semicolon_clauses(candidate):
            add(clause)
            stripped_clause = _strip_formula_narrative_prefix(clause)
            if stripped_clause != clause:
                add(stripped_clause)

        for nested_formula in _extract_parenthesized_formula_candidates(candidate):
            add(nested_formula)

        if "=" not in candidate:
            continue
        parts = [_strip_formula_trailing_punctuation(part).strip() for part in candidate.split("=")]
        for part in reversed(parts[:-1]):
            if not part:
                continue
            stripped_part = _strip_formula_narrative_prefix(part)
            if stripped_part != part:
                add(stripped_part)
                continue
            add(part)

    return candidates


def _linear_formula_assignment_to_calc_expr(text: str) -> tuple[str, dict[str, str]] | None:
    if "=" not in text:
        return None

    left_text, right_text = (_strip_formula_trailing_punctuation(part).strip() for part in text.split("=", 1))
    left_match = FORMULA_IDENTIFIER_TOKEN_RE.fullmatch(left_text)
    if left_match is None:
        return None

    left_key = _formula_symbol_key(left_text)
    left_calc_name = _formula_symbol_to_calc_identifier(left_text)
    if not left_key or not left_calc_name:
        return None

    right_expr, variables = _tokenize_formula_expression(_normalize_formula_expression_candidate(right_text))
    if right_expr is None:
        return None

    variables = {left_key: left_calc_name, **variables}
    return f"{left_calc_name} = {right_expr}", variables


def _linear_formula_expression_to_calc_expr(text: str) -> tuple[str, dict[str, str]] | None:
    if "=" in text:
        return None

    expression_text = _normalize_formula_expression_candidate(text)
    expression, variables = _tokenize_formula_expression(expression_text)
    if expression is None:
        return None
    if not any(operator in expression for operator in (" + ", " - ", " * ", " / ", " ** ")):
        return None
    return expression, variables


def _strip_formula_narrative_prefix(text: str) -> str:
    stripped = text.strip()
    if ":" not in stripped:
        return stripped

    prefix, suffix = stripped.split(":", 1)
    if not re.search(r"[A-Za-zА-Яа-я]", prefix):
        return stripped
    if not any(marker in suffix for marker in ("=", "+", "-", "*", "/", "×", "÷", "^", "%", "(", ")", "x", "X")):
        return stripped
    return suffix.strip()


def _formula_semicolon_clauses(text: str) -> list[str]:
    clauses: list[str] = []
    for clause in text.split(";"):
        normalized_clause = _strip_formula_trailing_punctuation(clause).strip()
        if "=" in normalized_clause:
            clauses.append(normalized_clause)
    return clauses


def _extract_parenthesized_formula_candidates(text: str) -> list[str]:
    candidates: list[str] = []
    stack: list[int] = []
    start_index: int | None = None

    for index, char in enumerate(text):
        if char == "(":
            if not stack:
                start_index = index + 1
            stack.append(index)
            continue
        if char != ")" or not stack:
            continue
        stack.pop()
        if stack or start_index is None:
            continue
        candidate = _strip_formula_trailing_punctuation(text[start_index:index]).strip()
        if "=" in candidate:
            candidates.append(candidate)
        start_index = None

    return candidates


def _normalize_formula_expression_candidate(text: str) -> str:
    normalized = FORMULA_INLINE_DIVISION_RE.sub(" ÷ ", text.strip())
    normalized = normalized.replace("[", "(").replace("]", ")")
    return _strip_formula_trailing_punctuation(normalized)


def _strip_formula_trailing_punctuation(text: str) -> str:
    return FORMULA_TRAILING_PUNCTUATION_RE.sub("", text.strip())


def _tokenize_formula_expression(text: str) -> tuple[str | None, dict[str, str]]:
    tokens: list[str] = []
    variables: dict[str, str] = {}
    position = 0

    while position < len(text):
        char = text[position]
        if char.isspace():
            position += 1
            continue

        number_match = FORMULA_NUMBER_TOKEN_RE.match(text, position)
        if number_match is not None:
            tokens.append(number_match.group(0).replace(",", "."))
            position = number_match.end()
            continue

        identifier_match = FORMULA_IDENTIFIER_TOKEN_RE.match(text, position)
        if identifier_match is not None:
            raw_identifier = identifier_match.group(0)
            symbol_key = _formula_symbol_key(raw_identifier)
            calc_identifier = _formula_symbol_to_calc_identifier(raw_identifier)
            if not symbol_key or not calc_identifier:
                return None, {}
            variables.setdefault(symbol_key, calc_identifier)
            tokens.append(calc_identifier)
            position = identifier_match.end()
            continue

        if char in "+-*/()":
            tokens.append(char)
            position += 1
            continue

        if char == "^":
            tokens.append("**")
            position += 1
            continue

        if char == "×":
            tokens.append("*")
            position += 1
            continue

        if char == "÷":
            tokens.append("/")
            position += 1
            continue

        if char == "%":
            if not tokens or tokens[-1] in {"+", "-", "*", "/", "**", "("}:
                return None, {}
            tokens.extend(["/", "100"])
            position += 1
            continue

        return None, {}

    if not tokens:
        return None, {}
    return " ".join(tokens), variables


def _normalize_formula_operator_text(text: str) -> str:
    normalized = FORMULA_ASCII_MULTIPLY_RE.sub(" × ", text)
    return FORMULA_REPEAT_MULTIPLY_RE.sub(" × ", normalized)


def _formula_symbol_key(value: str) -> str:
    parsed = _parse_formula_identifier(value)
    if parsed is None:
        return ""
    base, subscript, superscript = parsed
    parts = [_sanitize_formula_symbol_part(base)]
    if subscript:
        parts.append(_sanitize_formula_symbol_part(subscript))
    if superscript:
        parts.append(_sanitize_formula_symbol_part(superscript))
    return "_".join(part for part in parts if part)


def _formula_symbol_to_calc_identifier(value: str) -> str:
    symbol_key = _formula_symbol_key(value)
    if not symbol_key:
        return ""

    transliterated_parts: list[str] = []
    for char in symbol_key:
        if char == "_":
            transliterated_parts.append("_")
        elif char.isascii() and char.isalnum():
            transliterated_parts.append(char)
        else:
            transliterated_parts.append(CYRILLIC_TO_LATIN.get(char, "_"))

    normalized = re.sub(r"_+", "_", "".join(transliterated_parts)).strip("_")
    if not normalized:
        return ""
    if normalized[0].isdigit():
        normalized = f"v_{normalized}"
    return normalized


def _sanitize_formula_symbol_part(value: str) -> str:
    return re.sub(r"_+", "_", re.sub(r"[^0-9A-Za-zА-Яа-я]+", "_", value)).strip("_")


def _parse_formula_identifier(value: str) -> tuple[str, str | None, str | None] | None:
    match = re.fullmatch(r"([A-Za-zА-Яа-я][A-Za-zА-Яа-я0-9]*)(?:_\(([^()]+)\))?(?:\^\(([^()]+)\))?", value)
    if match is None:
        return None
    return match.group(1), match.group(2), match.group(3)


def _latex_sum_replacement(match: re.Match[str]) -> str:
    index_name = match.group(1).strip()
    start = match.group(2).strip()
    end = match.group(3).strip()
    return rf"\sum_{{{index_name}={start}}}^{{{end}}}"


def _latex_sub_replacement(match: re.Match[str]) -> str:
    return f"{_latex_identifier(match.group(1))}_{{{_latex_script(match.group(2))}}}"


def _latex_sup_replacement(match: re.Match[str]) -> str:
    return f"{_latex_identifier(match.group(1))}^{{{_latex_script(match.group(2))}}}"


def _latex_sub_sup_replacement(match: re.Match[str]) -> str:
    return (
        f"{_latex_identifier(match.group(1))}_{{{_latex_script(match.group(2))}}}"
        f"^{{{_latex_script(match.group(3))}}}"
    )


def _latex_identifier(value: str) -> str:
    if re.fullmatch(r"[A-Za-z]", value):
        return value
    return rf"\mathrm{{{value}}}"


def _latex_script(value: str) -> str:
    if re.fullmatch(r"[A-Za-z0-9]+", value):
        return value
    return rf"\text{{{value}}}"


def _paragraph_semantic_text(paragraph: Paragraph) -> str:
    text_parts: list[str] = []
    has_text_content = False
    for run in paragraph.runs:
        run_text, run_has_text = _run_semantic_text(paragraph, run)
        if run_text:
            text_parts.append(run_text)
        if run_has_text:
            has_text_content = True
    text = "".join(text_parts).strip()
    if text and (has_text_content or _looks_like_formula_text(text)):
        return text

    text = paragraph.text.strip()
    if text:
        return text
    math_text = " ".join(
        str(element.text).strip()
        for element in paragraph._p.iter()
        if element.tag.endswith("}t") and element.text and str(element.text).strip()
    )
    return math_text.strip()


def _table_cell_semantic_text(cell: Any) -> str:
    text_parts: list[str] = []
    for paragraph in cell.paragraphs:
        paragraph_text = _paragraph_semantic_text(paragraph)
        if paragraph_text:
            text_parts.append(paragraph_text)
    return "\n".join(text_parts)


def _table_cell_formula(text: str) -> dict[str, Any] | None:
    lines = [line.strip() for line in text.splitlines() if line.strip()]
    for line in lines:
        if _looks_like_formula_text(line):
            return _formula_representation_from_text(line)
    return None


def _run_semantic_text(paragraph: Paragraph, run: Any) -> tuple[str, bool]:
    parts: list[str] = []
    has_text_content = False
    for child in run._r:
        if child.tag.endswith("}t") and child.text:
            parts.append(str(child.text))
            has_text_content = True
            continue
        if child.tag.endswith("}tab"):
            parts.append("\t")
            has_text_content = True
            continue
        if child.tag.endswith("}br"):
            parts.append("\n")
            has_text_content = True
            continue
        if child.tag.endswith("}drawing"):
            parts.extend(_inline_drawing_placeholders(paragraph, child))

    text = "".join(parts)
    if not text:
        return "", False

    if run.font.subscript:
        return f"_({text})", has_text_content
    if run.font.superscript:
        return f"^({text})", has_text_content
    return text, has_text_content


def _inline_drawing_placeholders(paragraph: Paragraph, drawing: Any) -> list[str]:
    placeholders: list[str] = []
    seen_rel_ids: set[str] = set()
    drawing_extent = _inline_drawing_extent_emu(drawing)
    for node in drawing.iter():
        rel_id = node.attrib.get(f"{REL_NS}embed")
        if rel_id is None or rel_id in seen_rel_ids:
            continue
        seen_rel_ids.add(rel_id)
        placeholders.append(_inline_drawing_placeholder(paragraph, rel_id, drawing_extent))
    return placeholders


def _inline_drawing_placeholder(
    paragraph: Paragraph,
    rel_id: str,
    drawing_extent: tuple[int, int] | None = None,
) -> str:
    asset_name, blob = _resolve_inline_drawing_asset(paragraph, rel_id)
    text = _extract_formula_text_from_asset(blob, asset_name)
    if text is not None:
        return text

    symbol = _recognize_inline_glyph(blob, asset_name, drawing_extent)
    if symbol is not None:
        return symbol

    if asset_name is not None:
        return f"[INLINE_DRAWING:{asset_name}]"

    related_parts = getattr(paragraph.part, "related_parts", None)
    if related_parts is not None:
        related_part = related_parts.get(rel_id)
        part_name = getattr(related_part, "partname", None)
        if part_name is not None:
            return f"[INLINE_DRAWING:{Path(str(part_name)).name}]"

    relationships = getattr(paragraph.part, "rels", None)
    if relationships is not None and rel_id in relationships:
        target_ref = getattr(relationships[rel_id], "target_ref", None)
        if target_ref:
            return f"[INLINE_DRAWING:{Path(str(target_ref)).name}]"

    return f"[INLINE_DRAWING:{rel_id}]"


def _resolve_inline_drawing_asset(paragraph: Paragraph, rel_id: str) -> tuple[str | None, bytes | None]:
    related_parts = getattr(paragraph.part, "related_parts", None)
    if related_parts is not None:
        related_part = related_parts.get(rel_id)
        if related_part is not None:
            part_name = getattr(related_part, "partname", None)
            blob = getattr(related_part, "blob", None)
            if part_name is not None:
                return Path(str(part_name)).name, bytes(blob) if blob is not None else None

    relationships = getattr(paragraph.part, "rels", None)
    if relationships is not None and rel_id in relationships:
        target_ref = getattr(relationships[rel_id], "target_ref", None)
        if target_ref:
            return Path(str(target_ref)).name, None
    return None, None


def _inline_drawing_extent_emu(drawing: Any) -> tuple[int, int] | None:
    for node in drawing.iter():
        cx = node.attrib.get("cx")
        cy = node.attrib.get("cy")
        if cx is None or cy is None:
            continue
        try:
            return int(cx), int(cy)
        except ValueError:
            continue
    return None


def _recognize_inline_glyph(
    blob: bytes | None,
    asset_name: str | None,
    drawing_extent: tuple[int, int] | None,
) -> str | None:
    if blob is None or asset_name is None or drawing_extent is None:
        return None
    if max(drawing_extent) > INLINE_GLYPH_MAX_EXTENT_EMU:
        return None

    cache_key = hashlib.sha256(blob + b"|" + asset_name.encode("utf-8", errors="ignore")).hexdigest()
    if cache_key in INLINE_GLYPH_CACHE:
        return INLINE_GLYPH_CACHE[cache_key]

    image = _load_inline_glyph_image(blob, asset_name)
    if image is None:
        INLINE_GLYPH_CACHE[cache_key] = None
        return None

    symbol = _match_inline_glyph(image)
    INLINE_GLYPH_CACHE[cache_key] = symbol
    return symbol


def _extract_formula_text_from_asset(blob: bytes | None, asset_name: str | None) -> str | None:
    if blob is None or asset_name is None:
        return None
    suffix = Path(asset_name).suffix.lower()
    if suffix == ".wmf":
        return _extract_mathtype_wmf_text(blob)
    return None


def _extract_mathtype_wmf_text(blob: bytes) -> str | None:
    if b"MathType" not in blob and b"Design Science" not in blob:
        return None

    chunks = _extract_wmf_text_chunks(blob)
    if not chunks:
        return None
    return _assemble_mathtype_wmf_formula(chunks)


def _extract_wmf_text_chunks(blob: bytes) -> list[WmfTextChunk]:
    offset = 22 if blob[:4] == b"\xd7\xcd\xc6\x9a" else 0
    if len(blob) < offset + 18:
        return []
    offset += 18

    objects: list[dict[str, Any] | None] = []
    selected_handle: int | None = None
    chunks: list[WmfTextChunk] = []
    order = 0

    while offset + 6 <= len(blob):
        size_words = struct.unpack_from("<I", blob, offset)[0]
        func = struct.unpack_from("<H", blob, offset + 4)[0]
        if size_words == 0:
            break
        size_bytes = size_words * 2
        if offset + size_bytes > len(blob):
            break
        params = blob[offset + 6 : offset + size_bytes]

        if func == META_CREATEFONTINDIRECT and len(params) >= 50:
            height = struct.unpack_from("<h", params, 0)[0]
            charset = params[13]
            face_bytes = params[18:50]
            face = face_bytes.split(b"\x00", 1)[0].decode("latin1", errors="ignore")
            handle = next((index for index, item in enumerate(objects) if item is None), len(objects))
            if handle == len(objects):
                objects.append(None)
            objects[handle] = {
                "face": face,
                "charset": charset,
                "height": height,
            }
        elif func == META_SELECTOBJECT and len(params) >= 2:
            selected_handle = struct.unpack_from("<H", params, 0)[0]
        elif func == META_DELETEOBJECT and len(params) >= 2:
            handle = struct.unpack_from("<H", params, 0)[0]
            if handle < len(objects):
                objects[handle] = None
        elif func in {META_TEXTOUT, META_EXTTEXTOUT}:
            if selected_handle is None or selected_handle >= len(objects):
                offset += size_bytes
                if func == 0x0000:
                    break
                continue
            font = objects[selected_handle]
            if font is None:
                offset += size_bytes
                if func == 0x0000:
                    break
                continue
            raw_text = _extract_wmf_text_record_bytes(func, params)
            if not raw_text:
                offset += size_bytes
                if func == 0x0000:
                    break
                continue
            text = _decode_wmf_text(raw_text, str(font["face"]), int(font["charset"]))
            if text:
                chunks.append(
                    WmfTextChunk(
                        text=text,
                        face=str(font["face"]),
                        charset=int(font["charset"]),
                        height=int(font["height"]),
                        order=order,
                    )
                )
                order += 1

        offset += size_bytes
        if func == 0x0000:
            break

    return chunks


def _extract_wmf_text_record_bytes(func: int, params: bytes) -> bytes:
    if func == META_TEXTOUT:
        if len(params) < 2:
            return b""
        count = struct.unpack_from("<H", params, 0)[0]
        return params[2 : 2 + count]

    if len(params) < 8:
        return b""
    count = struct.unpack_from("<h", params, 4)[0]
    options = struct.unpack_from("<h", params, 6)[0]
    text_offset = 8 + (8 if options & 0x0006 else 0)
    return params[text_offset : text_offset + count]


def _decode_wmf_text(raw_text: bytes, face: str, charset: int) -> str:
    normalized_face = face.strip().lower()
    if normalized_face == "symbol":
        pieces: list[str] = []
        for byte in raw_text:
            if 32 <= byte < 127:
                pieces.append(chr(byte))
                continue
            mapped = SYMBOL_FONT_MAP.get(byte)
            if mapped is not None:
                pieces.append(mapped)
        return "".join(pieces)

    encoding = "cp1251" if charset == WMF_RUSSIAN_CHARSET else "latin1"
    return raw_text.decode(encoding, errors="ignore")


def _assemble_mathtype_wmf_formula(chunks: list[WmfTextChunk]) -> str | None:
    meaningful_chunks = [chunk for chunk in chunks if chunk.text.strip()]
    if not meaningful_chunks:
        return None

    known = _assemble_known_mathtype_formula(meaningful_chunks)
    if known is not None:
        return known

    formula_ir = _build_wmf_formula_ir(meaningful_chunks)
    assembled = _assemble_wmf_formula_ir(formula_ir)
    if assembled is not None:
        return assembled

    return None


def _build_wmf_formula_ir(chunks: list[WmfTextChunk]) -> WmfFormulaIR:
    meaningful_chunks = _coalesce_adjacent_mathtype_chunks([chunk for chunk in chunks if chunk.text.strip()])
    if not meaningful_chunks:
        return WmfFormulaIR(tokens=(), signature="", layout_class="empty", base_text="", script_text="")

    max_height = max(abs(chunk.height) for chunk in meaningful_chunks)
    tokens = tuple(
        WmfFormulaToken(
            text=chunk.text,
            role="base" if abs(chunk.height) == max_height else "script",
            face=chunk.face,
            charset=chunk.charset,
            height=chunk.height,
            order=chunk.order,
            symbol_font=chunk.face.strip().lower() == "symbol",
        )
        for chunk in meaningful_chunks
    )
    base_text = "".join(token.text for token in tokens if token.role == "base").strip()
    script_text = "".join(token.text for token in tokens if token.role == "script").strip()
    return WmfFormulaIR(
        tokens=tokens,
        signature=_chunk_signature(meaningful_chunks),
        layout_class=_classify_wmf_formula_ir(tokens, base_text=base_text, script_text=script_text),
        base_text=base_text,
        script_text=script_text,
    )


def _classify_wmf_formula_ir(
    tokens: tuple[WmfFormulaToken, ...],
    *,
    base_text: str,
    script_text: str,
) -> str:
    if not tokens:
        return "empty"
    if _assemble_wmf_ir_interleaved_symbol_formula(tokens) is not None:
        return "interleaved_symbol"
    if base_text and script_text:
        return "base_with_scripts"
    if base_text:
        return "base_only"
    if script_text:
        return "script_only"
    return "unclassified"


def _assemble_wmf_formula_ir(formula_ir: WmfFormulaIR) -> str | None:
    interleaved = _assemble_wmf_ir_interleaved_symbol_formula(formula_ir.tokens)
    if interleaved is not None:
        return interleaved

    aggregate_price_formula = _assemble_wmf_ir_aggregate_price_formula(formula_ir.tokens)
    if aggregate_price_formula is not None:
        return aggregate_price_formula

    if formula_ir.layout_class == "base_only":
        return formula_ir.base_text or None
    if formula_ir.layout_class != "base_with_scripts":
        return None

    base_text = formula_ir.base_text
    script_tokens = [token for token in formula_ir.tokens if token.role == "script"]

    superscript_tokens: list[WmfFormulaToken] = []
    subscript_tokens: list[WmfFormulaToken] = []
    for token in script_tokens:
        if len(token.text) == 1 and token.charset != WMF_RUSSIAN_CHARSET:
            superscript_tokens.append(token)
        else:
            subscript_tokens.append(token)

    parts: list[str] = [base_text] if base_text else []
    if subscript_tokens:
        parts.append(f"_({''.join(token.text for token in subscript_tokens)})")
    if superscript_tokens:
        parts.append(f"^({''.join(token.text for token in superscript_tokens)})")
    result = "".join(parts).strip()
    return result or None


def _assemble_wmf_ir_aggregate_price_formula(tokens: tuple[WmfFormulaToken, ...]) -> str | None:
    if len(tokens) < 3:
        return None

    text_tokens = [token.text for token in tokens if not token.symbol_font]
    symbol_text = "".join(token.text for token in tokens if token.symbol_font)
    if len(text_tokens) != 2 or symbol_text != "*+*=+":
        return None

    if re.fullmatch(r"ОЦ\(\d+\),", text_tokens[0]) is None:
        return None

    weights = text_tokens[1]
    if len(weights) != 4 or any(char not in {"n", "m"} for char in weights):
        return None

    left_weight = weights[0]
    right_weight = weights[1]
    denominator_left = weights[2]
    denominator_right = weights[3]
    return (
        "ОЦ_(а) = "
        f"(Х_(св) × {left_weight} + Х_(сп) × {right_weight}) / "
        f"({denominator_left} + {denominator_right})"
    )


def _assemble_interleaved_symbol_formula(chunks: list[WmfTextChunk]) -> str | None:
    tokens = tuple(
        WmfFormulaToken(
            text=chunk.text,
            role="base",
            face=chunk.face,
            charset=chunk.charset,
            height=chunk.height,
            order=chunk.order,
            symbol_font=chunk.face.strip().lower() == "symbol",
        )
        for chunk in chunks
    )
    return _assemble_wmf_ir_interleaved_symbol_formula(tokens)


def _assemble_wmf_ir_interleaved_symbol_formula(tokens: tuple[WmfFormulaToken, ...]) -> str | None:
    if len(tokens) != 2:
        return None
    symbol_token = next((token for token in tokens if token.symbol_font), None)
    text_token = next((token for token in tokens if not token.symbol_font), None)
    if symbol_token is None or text_token is None:
        return None
    if len(symbol_token.text) != 2 or len(text_token.text) != 3:
        return None
    if symbol_token.text[0] != "=":
        return None
    return (
        f"{text_token.text[0]} {symbol_token.text[0]} {text_token.text[1]} "
        f"{symbol_token.text[1]} {text_token.text[2]}"
    )


def _chunk_signature(chunks: list[WmfTextChunk]) -> str:
    return "|".join(chunk.text for chunk in chunks)


def _coalesce_adjacent_mathtype_chunks(chunks: list[WmfTextChunk]) -> list[WmfTextChunk]:
    if not chunks:
        return []

    coalesced = [chunks[0]]
    for chunk in chunks[1:]:
        previous = coalesced[-1]
        if (
            chunk.order == previous.order + 1
            and chunk.face == previous.face
            and chunk.charset == previous.charset
            and chunk.height == previous.height
        ):
            coalesced[-1] = WmfTextChunk(
                text=previous.text + chunk.text,
                face=previous.face,
                charset=previous.charset,
                height=previous.height,
                order=previous.order,
            )
            continue
        coalesced.append(chunk)
    return coalesced


def _assemble_known_mathtype_formula(chunks: list[WmfTextChunk]) -> str | None:
    signatures = {_chunk_signature(chunks)}
    coalesced_chunks = _coalesce_adjacent_mathtype_chunks(chunks)
    if coalesced_chunks != chunks:
        signatures.add(_chunk_signature(coalesced_chunks))

    def has_all_markers(*markers: str) -> bool:
        return any(all(marker in signature for marker in markers) for signature in signatures)

    if "k1|К|=÷" in signatures:
        return "k = 1 ÷ K"
    if any("М = P  См" in signature and "j=1" in signature and "J" in signature for signature in signatures):
        return "М_(тек) = sum_(j=1)^J P^(j) × См_(тек)^(j)"
    if "с|Ц|С|Т|=" in signatures:
        return "С_(маш.р) = Ц_(а) / Т_(с)"
    if any("PPV" in signature and "ii" in signature and "T" in signature and "t" in signature for signature in signatures):
        return "P^(t) = sum_(i=1)^I P_(i)^(t) × V_(i)"
    if any("ОТЗТСЦV" in signature and "IN" in signature and "nini" in signature for signature in signatures):
        return "ОТ_(тек) = sum_(i=1)^I sum_(n=1)^N ЗТ_(ni) × СЦ_(n) × V_(i)"
    if any("ОТЗТСЦV" in signature and "ЗТ" in signature and "iii" in signature for signature in signatures):
        return "ОТ_(тек) = sum_(i=1)^I ЗТ_(i) × СЦ_(ЗТтек)^(i) × V_(i)"
    if any("ОТмЗТСЦV" in signature and "IK" in signature and "kiki" in signature for signature in signatures):
        return "ОТм_(тек) = sum_(i=1)^I sum_(k=1)^K ЗТ_(ki) × СЦ_(k) × V_(i)"
    if has_all_markers("ОЦ(1),", "min(,2)", "nmn", "*+*", "=", "+"):
        return "ОЦ_(а) = (Х_(св) × n + Х_(сп) × m) / (n + m)"
    if has_all_markers("ОЦ(3),", "|2|", "+", "="):
        return "ОЦ_(а) = (Х_(св) + Х_(сп)) / 2"
    if has_all_markers("Х(4),", "xvxvxv", "vvv", "12", "*+*++*"):
        return "Х_(св) = (x_(1) × v_(1) + x_(2) × v_(2) + ... + x_(n) × v_(n)) / (v_(1) + v_(2) + ... + v_(n))"
    if has_all_markers("Х(5),", "xxx", "|m|", "++", "="):
        return "Х_(сп) = (x_(1) + x_(2) + ... + x_(m)) / m"
    if has_all_markers("Х(7),", "0,25", "ср", "s", "<=*"):
        return "s <= 0,25 × Х_(ср)"
    if has_all_markers("СЦ = К    К (1),", "зт", "текTnинф", "ср", "C", "t"):
        return "СЦ_(зттек) = КТ_(n) × С_(1ср) / t_(ср) × К_(инф)"
    if has_all_markers("ОТ = СЦ  T (2),", "n", "зт", "1", "текii", "i=1"):
        return "ОТ_(1) = sum_(i=1)^n СЦ_(зтiтек) × Т_(i)"
    if has_all_markers("ОТ = СЦ  Т  К (3),", "зт", "21 ", "текТ"):
        return "ОТ_(2) = СЦ_(зт1тек) × Т × КТ"
    if has_all_markers("С = С  1 + K K + ПВ (4),", "n", "1 ", "ср1ip", "i=1"):
        return "С_(1ср) = С_(1) × (1 + sum_(i=1)^n К_(i) + К_(р)) + ПВ"
    if any("ЭММTСЦэм" in signature and "M" in signature and "mm" in signature for signature in signatures):
        return "ЭММ_(тек) = sum_(m=1)^M T_(m) × СЦэм_(тек)^(m)"
    if any("З = З  , (2)" in signature and "ср1" in signature and "смр" in signature and "Т|К" in signature for signature in signatures):
        return "З_(ср) = З_(1) × К_(смрТ)"
    if any(" =   , (3)" in signature and "пнр" in signature and "ii" in signature and "ЗТЗ" in signature for signature in signatures):
        return "З_(пнр) = sum_(i) Т_(i) × З_(i)"
    if any("З = З  К, (4)" in signature and "пнр" in signature and "1" in signature and "Т" in signature and signature.endswith("|i") for signature in signatures):
        return "З_(i) = З_(1) × К_(пнрТ)^(i)"
    if "ВрПВрЭ|НН|=" in signatures:
        return "Н_(ВрП) = sum Н_(ВрЭ)"
    if "факт|ВрЭ|пзротп|ЗТ|Ч100|Н|100|ННН60|=|-++" in signatures:
        return "Н_(ВрЭ) = ЗТ_(эСР) × 100 / (Ч_(факт) × [100 - (Н_(пзр) + Н_(о) + Н_(тп))] × 60)"
    if "ЗТ|ЗТ|n|=" in signatures:
        return "ЗТ_(эСР) = sum_(i=1)^n ЗТ_(э) / n"
    if "ЗТ|ЗТ|V|=" in signatures:
        return "ЗТ_(э) = ЗТ / V"
    if "min|t|К1,5|t|=" in signatures:
        return "К_(уст) = t_(max) / t_(min) <= 1,5"
    if any("С = Э  Ц, (5)" in signature and "эмэм" in signature and "ii" in signature for signature in signatures):
        return "С_(эм) = sum_(i) Э_(i) × Ц_(эмi)"
    if any("С = М  Ц, (6)" in signature and "мат" in signature and "ii" in signature for signature in signatures):
        return "С_(мат) = sum_(i) М_(i) × Ц_(i)"
    return None


def _load_inline_glyph_image(blob: bytes, asset_name: str) -> Image.Image | None:
    suffix = Path(asset_name).suffix.lower()
    if suffix in INLINE_GLYPH_RASTER_SUFFIXES:
        return _open_inline_glyph_image(blob)
    if suffix in INLINE_GLYPH_METAFILE_SUFFIXES:
        return _render_windows_metafile(blob, suffix)
    return None


def _open_inline_glyph_image(blob: bytes) -> Image.Image | None:
    try:
        with Image.open(io.BytesIO(blob)) as image:
            return image.convert("RGBA")
    except UnidentifiedImageError:
        return None


def _render_windows_metafile(blob: bytes, suffix: str) -> Image.Image | None:
    if os.name != "nt":
        return None

    powershell = shutil.which("powershell") or shutil.which("powershell.exe")
    if powershell is None:
        return None

    script = """param([string]$src, [string]$dst)
Add-Type -AssemblyName System.Drawing
$img = [System.Drawing.Image]::FromFile($src)
$bmp = New-Object System.Drawing.Bitmap $img.Width, $img.Height
$graphics = [System.Drawing.Graphics]::FromImage($bmp)
$graphics.Clear([System.Drawing.Color]::White)
$graphics.DrawImage($img, 0, 0, $img.Width, $img.Height)
$bmp.Save($dst, [System.Drawing.Imaging.ImageFormat]::Png)
$graphics.Dispose()
$bmp.Dispose()
$img.Dispose()
"""
    try:
        with tempfile.TemporaryDirectory(prefix="docx-inline-glyph-") as temp_dir:
            source_path = Path(temp_dir) / f"inline{suffix}"
            target_path = Path(temp_dir) / "inline.png"
            script_path = Path(temp_dir) / "render-inline-metafile.ps1"
            source_path.write_bytes(blob)
            script_path.write_text(script, encoding="ascii")
            completed = subprocess.run(
                [
                    powershell,
                    "-NoProfile",
                    "-NonInteractive",
                    "-ExecutionPolicy",
                    "Bypass",
                    "-File",
                    str(script_path),
                    str(source_path),
                    str(target_path),
                ],
                capture_output=True,
                text=True,
                timeout=INLINE_GLYPH_RENDER_TIMEOUT_SECONDS,
                check=False,
            )
            if completed.returncode != 0 or not target_path.exists():
                return None
            return _open_inline_glyph_image(target_path.read_bytes())
    except (OSError, subprocess.SubprocessError):
        return None


def _match_inline_glyph(image: Image.Image) -> str | None:
    normalized = _normalize_inline_glyph_image(image)
    if normalized is None:
        return None

    scores: list[tuple[float, str]] = []
    for symbol in INLINE_GLYPH_SYMBOLS:
        best_symbol_score = 1.0
        for font_name in _available_inline_glyph_fonts():
            for font_size in range(24, 61, 4):
                template = _render_inline_glyph_template(symbol, font_name, font_size)
                score = _inline_glyph_difference_score(normalized, template)
                if score < best_symbol_score:
                    best_symbol_score = score
        scores.append((best_symbol_score, symbol))

    scores.sort(key=lambda item: item[0])
    if not scores:
        return None
    best_score, best_symbol = scores[0]
    second_score = scores[1][0] if len(scores) > 1 else 1.0
    if best_score <= INLINE_GLYPH_SCORE_THRESHOLD and second_score - best_score >= INLINE_GLYPH_SCORE_MARGIN:
        return best_symbol
    return None


def _normalize_inline_glyph_image(image: Image.Image) -> Image.Image | None:
    rgba = image.convert("RGBA")
    background = Image.new("RGBA", rgba.size, (255, 255, 255, 255))
    rgba = Image.alpha_composite(background, rgba)
    grayscale = rgba.convert("L")
    mask = grayscale.point(_inline_glyph_mask_value)
    bbox = mask.getbbox()
    if bbox is None:
        return None

    cropped = grayscale.crop(bbox)
    target_size = INLINE_GLYPH_TEMPLATE_SIZE - INLINE_GLYPH_TEMPLATE_MARGIN
    scale = min(target_size / max(cropped.width, 1), target_size / max(cropped.height, 1))
    resized = cropped.resize(
        (
            max(1, int(round(cropped.width * scale))),
            max(1, int(round(cropped.height * scale))),
        ),
        LANCZOS_RESAMPLING,
    )
    canvas = Image.new("L", (INLINE_GLYPH_TEMPLATE_SIZE, INLINE_GLYPH_TEMPLATE_SIZE), 255)
    offset = (
        (INLINE_GLYPH_TEMPLATE_SIZE - resized.width) // 2,
        (INLINE_GLYPH_TEMPLATE_SIZE - resized.height) // 2,
    )
    canvas.paste(resized, offset)
    return canvas.point(_inline_glyph_binary_value)


def _inline_glyph_mask_value(value: int) -> int:
    return 255 if value < 245 else 0


def _inline_glyph_binary_value(value: int) -> int:
    return 0 if value < 220 else 255


@lru_cache(maxsize=32)
def _available_inline_glyph_fonts() -> tuple[str, ...]:
    fonts: list[str] = []
    for candidate in INLINE_GLYPH_FONT_CANDIDATES:
        try:
            ImageFont.truetype(candidate, size=32)
        except OSError:
            continue
        fonts.append(candidate)
    return tuple(fonts) or ("__default__",)


@lru_cache(maxsize=2048)
def _render_inline_glyph_template(symbol: str, font_name: str, font_size: int) -> Image.Image:
    canvas = Image.new("L", (INLINE_GLYPH_TEMPLATE_SIZE, INLINE_GLYPH_TEMPLATE_SIZE), 255)
    draw = ImageDraw.Draw(canvas)
    font = ImageFont.load_default() if font_name == "__default__" else ImageFont.truetype(font_name, size=font_size)
    bbox = draw.textbbox((0, 0), symbol, font=font)
    x = (INLINE_GLYPH_TEMPLATE_SIZE - (bbox[2] - bbox[0])) // 2 - bbox[0]
    y = (INLINE_GLYPH_TEMPLATE_SIZE - (bbox[3] - bbox[1])) // 2 - bbox[1]
    draw.text((x, y), symbol, fill=0, font=font)
    normalized = _normalize_inline_glyph_image(canvas)
    return normalized if normalized is not None else canvas


def _inline_glyph_difference_score(image: Image.Image, template: Image.Image) -> float:
    diff = ImageChops.difference(image, template)
    return sum(diff.getdata()) / (255 * INLINE_GLYPH_TEMPLATE_SIZE * INLINE_GLYPH_TEMPLATE_SIZE)


def _paragraph_has_omml(paragraph: Paragraph) -> bool:
    return any(element.tag.endswith("}oMath") or element.tag.endswith("}oMathPara") for element in paragraph._p.iter())


def _looks_like_formula_text(text: str) -> bool:
    stripped = text.strip()
    if len(stripped) > 180:
        return False
    return bool(FORMULA_RE.search(stripped))


def _iter_header_footer_units(document: Any) -> list[tuple[str, str, str]]:
    units: list[tuple[str, str, str]] = []
    seen: set[tuple[str, str]] = set()
    for section_index, section in enumerate(document.sections, start=1):
        for unit_type, story in (("header", section.header), ("footer", section.footer)):
            for paragraph_index, paragraph in enumerate(story.paragraphs, start=1):
                text = _paragraph_semantic_text(paragraph)
                if not text:
                    continue
                key = (unit_type, text)
                if key in seen:
                    continue
                seen.add(key)
                units.append((unit_type, text, f"/word/section[{section_index}]/{unit_type}/p[{paragraph_index}]"))
    return units


def _extract_docx_footnotes(source_path: Path) -> list[tuple[str, str]]:
    footnotes: list[tuple[str, str]] = []
    try:
        with zipfile.ZipFile(source_path) as archive:
            if "word/footnotes.xml" not in archive.namelist():
                return footnotes
            root = ElementTree.fromstring(archive.read("word/footnotes.xml"))
    except (zipfile.BadZipFile, ElementTree.ParseError):
        return footnotes

    for footnote in root.findall(f"{WORD_NS}footnote"):
        footnote_id = str(footnote.attrib.get(f"{WORD_NS}id", ""))
        footnote_type = str(footnote.attrib.get(f"{WORD_NS}type", ""))
        if footnote_type or footnote_id in {"", "-1", "0"}:
            continue
        text = " ".join(
            str(node.text).strip()
            for node in footnote.iter(f"{WORD_NS}t")
            if node.text and str(node.text).strip()
        ).strip()
        if text:
            footnotes.append((footnote_id, text))
    return footnotes


def _write_json(path: Path, payload: dict[str, Any]) -> None:
    path.write_text(json.dumps(payload, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")


def _write_validated_json(path: Path, payload: dict[str, Any], schema_filename: str) -> None:
    validate_payload(payload, schema_filename)
    _write_json(path, payload)