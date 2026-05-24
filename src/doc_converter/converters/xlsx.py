from __future__ import annotations

import json
from dataclasses import dataclass
from datetime import date, datetime, time
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

from doc_converter.canonical import SourceRef, StructuralUnit, document_id_from_sha256, minimal_document, unit_id
from doc_converter.document_metadata import build_document_metadata
from doc_converter.quality import quality_payload, text_quality_flags
from doc_converter.schema_validation import validate_payload


@dataclass(frozen=True)
class XlsxConversionResult:
    status: str
    sheets: int
    units_count: int
    text_chars: int
    formula_cells: int
    warnings: tuple[str, ...] = ()


def convert_xlsx(
    source_path: Path,
    output_dir: Path,
    sha256: str,
    *,
    relative_source_path: str | None = None,
) -> XlsxConversionResult:
    output_dir.mkdir(parents=True, exist_ok=True)
    formula_workbook = load_workbook(source_path, data_only=False, read_only=False)
    values_workbook = load_workbook(source_path, data_only=True, read_only=False)
    try:
        doc_id = document_id_from_sha256(sha256)
        units: list[StructuralUnit] = [
            StructuralUnit(
                unit_id=unit_id(0),
                type="document",
                order=0,
                source_ref=SourceRef(document_id=doc_id, docx_path="/xl/workbook.xml"),
            )
        ]
        search_parts: list[str] = []
        sheets_payload: list[dict[str, Any]] = []
        order = 1
        total_formula_cells = 0
        missing_cached_formula_values = 0

        for sheet_index, formula_sheet in enumerate(formula_workbook.worksheets, start=1):
            values_sheet = values_workbook[formula_sheet.title]
            sheet_path = f"/xl/worksheets/sheet{sheet_index}.xml"
            sheet_id = unit_id(order)
            sheet_flags = [] if formula_sheet.sheet_state == "visible" else ["sheet_hidden"]
            units.append(
                StructuralUnit(
                    unit_id=sheet_id,
                    parent_id=unit_id(0),
                    type="section",
                    order=order,
                    text=formula_sheet.title,
                    source_ref=SourceRef(document_id=doc_id, docx_path=sheet_path),
                    quality=quality_payload(sheet_flags),
                )
            )
            search_parts.append(f"[{formula_sheet.title}]")
            order += 1

            table_id = unit_id(order)
            units.append(
                StructuralUnit(
                    unit_id=table_id,
                    parent_id=sheet_id,
                    type="table",
                    order=order,
                    text=_sheet_table_text(formula_sheet.title, formula_sheet.calculate_dimension()),
                    source_ref=SourceRef(document_id=doc_id, docx_path=sheet_path),
                    quality=quality_payload(["spreadsheet_table_inferred"]),
                )
            )
            order += 1

            non_empty_cells = 0
            formula_cells = 0
            sheet_search_rows: list[str] = []
            for row_index, row in enumerate(formula_sheet.iter_rows(), start=1):
                cell_units: list[tuple[Any, dict[str, Any], str]] = []
                for formula_cell in row:
                    values_cell = values_sheet[formula_cell.coordinate]
                    cell_payload = _spreadsheet_cell_payload(formula_sheet.title, formula_cell, values_cell)
                    if not _cell_has_content(cell_payload):
                        continue
                    non_empty_cells += 1
                    if cell_payload["formula"] is not None:
                        formula_cells += 1
                        if cell_payload["value"] is None:
                            missing_cached_formula_values += 1
                    cell_units.append((formula_cell, cell_payload, _cell_search_text(cell_payload)))

                if not cell_units:
                    continue

                row_id = unit_id(order)
                units.append(
                    StructuralUnit(
                        unit_id=row_id,
                        parent_id=table_id,
                        type="table_row",
                        order=order,
                        source_ref=SourceRef(document_id=doc_id, docx_path=f"{sheet_path}/row[{row_index}]"),
                    )
                )
                order += 1

                sheet_search_rows.append("\t".join(cell_text for _cell, _payload, cell_text in cell_units))
                for formula_cell, cell_payload, cell_text in cell_units:
                    units.append(
                        StructuralUnit(
                            unit_id=unit_id(order),
                            parent_id=row_id,
                            type="table_cell",
                            order=order,
                            text=cell_text,
                            cell=cell_payload,
                            source_ref=SourceRef(
                                document_id=doc_id,
                                docx_path=f"{sheet_path}/{formula_cell.coordinate}",
                            ),
                        )
                    )
                    order += 1

            total_formula_cells += formula_cells
            if sheet_search_rows:
                search_parts.append("\n".join(sheet_search_rows))
            sheets_payload.append(
                {
                    "title": formula_sheet.title,
                    "state": formula_sheet.sheet_state,
                    "dimension": formula_sheet.calculate_dimension(),
                    "max_row": formula_sheet.max_row,
                    "max_column": formula_sheet.max_column,
                    "non_empty_cells": non_empty_cells,
                    "formula_cells": formula_cells,
                    "merged_ranges": [str(merged_range) for merged_range in formula_sheet.merged_cells.ranges],
                    "freeze_panes": str(formula_sheet.freeze_panes) if formula_sheet.freeze_panes else None,
                    "print_area": str(formula_sheet.print_area) if formula_sheet.print_area else None,
                }
            )

        search_text = "\n\n".join(part for part in search_parts if part.strip())
        flags = text_quality_flags(search_text, size_bytes=source_path.stat().st_size, route="xlsx_native")
        warnings = []
        if missing_cached_formula_values:
            warnings.append("formula_cached_values_missing")
        status = "success" if search_text else "partial_success"
        payload = minimal_document(
            source_path=source_path,
            source_format="xlsx",
            sha256=sha256,
            route="xlsx_native",
            status=status,
            units=units,
            assets=[],
            quality=quality_payload(flags, warnings=warnings),
            relative_source_path=relative_source_path,
            metadata=build_document_metadata(filename=source_path.name, units=units, search_text=search_text),
        )
        _write_validated_json(output_dir / "document.v1.json", payload, "document.v1.schema.json")
        _write_json(
            output_dir / "extractor_raw.json",
            {
                "extractor": "openpyxl",
                "sheets": sheets_payload,
                "formula_cells": total_formula_cells,
                "formula_values_source": "workbook_cached_values",
                "missing_cached_formula_values": missing_cached_formula_values,
            },
        )
        (output_dir / "search_text.txt").write_text(search_text, encoding="utf-8")

        return XlsxConversionResult(
            status=status,
            sheets=len(formula_workbook.worksheets),
            units_count=len(units),
            text_chars=len(search_text),
            formula_cells=total_formula_cells,
            warnings=tuple(warnings),
        )
    finally:
        formula_workbook.close()
        values_workbook.close()


def _spreadsheet_cell_payload(sheet_title: str, formula_cell: Any, values_cell: Any) -> dict[str, Any]:
    formula = _cell_formula(formula_cell)
    value = _json_cell_value(values_cell.value if formula is not None else formula_cell.value)
    return {
        "worksheet": sheet_title,
        "address": formula_cell.coordinate,
        "row": formula_cell.row,
        "column": formula_cell.column,
        "value": value,
        "formula": formula,
        "data_type": str(formula_cell.data_type) if formula_cell.data_type is not None else None,
        "number_format": formula_cell.number_format or None,
        "hyperlink": str(formula_cell.hyperlink.target) if formula_cell.hyperlink is not None else None,
    }


def _cell_formula(cell: Any) -> str | None:
    value = cell.value
    if isinstance(value, str) and value.startswith("="):
        return value
    if getattr(cell, "data_type", None) == "f" and value is not None:
        text = str(value)
        return text if text.startswith("=") else f"={text}"
    return None


def _cell_has_content(cell_payload: dict[str, Any]) -> bool:
    return cell_payload["value"] is not None or cell_payload["formula"] is not None or cell_payload["hyperlink"] is not None


def _cell_search_text(cell_payload: dict[str, Any]) -> str:
    value_text = _value_text(cell_payload["value"])
    formula = cell_payload["formula"]
    if formula is not None and value_text:
        return f"{cell_payload['address']}: {formula} -> {value_text}"
    if formula is not None:
        return f"{cell_payload['address']}: {formula}"
    return f"{cell_payload['address']}: {value_text}"


def _value_text(value: object) -> str:
    if value is None:
        return ""
    if isinstance(value, bool):
        return "TRUE" if value else "FALSE"
    return str(value)


def _json_cell_value(value: object) -> str | int | float | bool | None:
    if value is None or isinstance(value, (str, int, float, bool)):
        return value
    if isinstance(value, (datetime, date, time)):
        return value.isoformat()
    return str(value)


def _sheet_table_text(title: str, dimension: str) -> str:
    return f"Worksheet {title}: {dimension}"


def _write_json(path: Path, payload: dict[str, Any]) -> None:
    path.write_text(json.dumps(payload, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")


def _write_validated_json(path: Path, payload: dict[str, Any], schema_filename: str) -> None:
    validate_payload(payload, schema_filename)
    _write_json(path, payload)