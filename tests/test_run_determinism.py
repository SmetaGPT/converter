from __future__ import annotations

import json
import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch

from docx import Document

from doc_converter.config import ConverterConfig, ConverterOptions
from doc_converter.runner import run_convert_folder


class RunDeterminismTests(unittest.TestCase):
    def test_clean_reruns_produce_identical_manifest_with_different_iterator_order(self) -> None:
        with tempfile.TemporaryDirectory() as input_dir, tempfile.TemporaryDirectory() as output_root:
            input_root = Path(input_dir).resolve()
            first_output = Path(output_root) / "first"
            second_output = Path(output_root) / "second"
            first_output.mkdir()
            second_output.mkdir()

            primary_path = input_root / "a.docx"
            duplicate_path = input_root / "z.docx"
            unique_path = input_root / "m.docx"
            unsupported_path = input_root / "notes.txt"

            primary_document = Document()
            primary_document.add_paragraph("Duplicate candidate")
            primary_document.save(str(primary_path))
            duplicate_path.write_bytes(primary_path.read_bytes())

            unique_document = Document()
            unique_document.add_paragraph("Unique candidate")
            unique_document.save(str(unique_path))
            unsupported_path.write_text("ignored", encoding="utf-8")

            first_order = [duplicate_path, unsupported_path, unique_path, primary_path]
            second_order = [primary_path, unique_path, unsupported_path, duplicate_path]

            with patch("doc_converter.inventory._iter_input_files", side_effect=[first_order, second_order]):
                with patch("doc_converter.run.orchestration._new_run_id", return_value="20260531T000000Z"):
                    first_result = run_convert_folder(
                        ConverterConfig(input_dir=input_root, output_dir=first_output, options=ConverterOptions())
                    )
                    second_result = run_convert_folder(
                        ConverterConfig(input_dir=input_root, output_dir=second_output, options=ConverterOptions())
                    )

            first_manifest_lines = _manifest_lines(first_result.run_dir / "manifest.jsonl")
            second_manifest_lines = _manifest_lines(second_result.run_dir / "manifest.jsonl")

            self.assertEqual(first_manifest_lines, second_manifest_lines)

            manifest_records = [json.loads(line) for line in first_manifest_lines]
            self.assertEqual([record["relative_path"] for record in manifest_records], ["a.docx", "m.docx", "z.docx"])
            self.assertEqual(manifest_records[0]["status"], "success")
            self.assertEqual(manifest_records[1]["status"], "success")
            self.assertEqual(manifest_records[2]["status"], "skipped_duplicate")
            self.assertEqual(manifest_records[2]["duplicate_of"], "a.docx")
            self.assertNotIn("reused_previous_output", manifest_records[2])


def _manifest_lines(path: Path) -> list[str]:
    return [line for line in path.read_text(encoding="utf-8").splitlines() if line.strip()]


if __name__ == "__main__":
    unittest.main()
