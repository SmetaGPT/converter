from __future__ import annotations

import base64
import json
import tempfile
import unittest
import zipfile
from pathlib import Path
from unittest.mock import patch

from PIL import Image, ImageDraw
from PIL import ImageFont
from docx import Document

from doc_converter.converters.docx import (
    INLINE_GLYPH_CACHE,
    WmfTextChunk,
    _assemble_mathtype_wmf_formula,
    _build_wmf_formula_ir,
    _formula_representation_from_text,
)
from doc_converter.converters.docx.inline_glyph import _available_inline_glyph_fonts
from doc_converter.converters.docx.formulas.wmf import WmfParseLimitError, _extract_wmf_text_chunks
from doc_converter.config import ConverterConfig, ConverterOptions, FormulaRecognitionConfig
from doc_converter.font_bundle import bundled_font_paths
from doc_converter.runner import run_convert_folder
from doc_converter.schema_validation import validate_payload


class DocxConverterTests(unittest.TestCase):
    def setUp(self) -> None:
        INLINE_GLYPH_CACHE.clear()

    def test_runner_converts_docx_to_document_package(self) -> None:
        with tempfile.TemporaryDirectory() as input_dir, tempfile.TemporaryDirectory() as output_dir:
            source_path = Path(input_dir) / "sample.docx"
            document = Document()
            document.add_paragraph("Заголовок")
            document.add_paragraph("Первый абзац")
            table = document.add_table(rows=1, cols=2)
            table.rows[0].cells[0].text = "A"
            table.rows[0].cells[1].text = "B"
            document.save(str(source_path))

            INLINE_GLYPH_CACHE.clear()
            with patch("doc_converter.converters.docx.inline_glyph._recognize_inline_glyph", return_value="÷"):
                result = run_convert_folder(_docx_converter_config(input_dir, output_dir))

            self.assertEqual(result.status, "success")
            manifest_records = [
                json.loads(line)
                for line in (result.run_dir / "manifest.jsonl").read_text(encoding="utf-8").splitlines()
                if line.strip()
            ]
            self.assertEqual(len(manifest_records), 1)
            self.assertEqual(manifest_records[0]["status"], "success")
            validate_payload(manifest_records[0], "manifest.v1.schema.json")
            document_dir = result.run_dir / manifest_records[0]["output_dir"]
            payload = json.loads((document_dir / "document.v1.json").read_text(encoding="utf-8"))
            search_text = (document_dir / "search_text.txt").read_text(encoding="utf-8")

            validate_payload(payload, "document.v1.schema.json")
            self.assertEqual(payload["schema_version"], "document.v1")
            self.assertEqual(payload["processing"]["route"], "docx_native")
            self.assertEqual(payload["source"]["relative_input_path"], "sample.docx")
            self.assertEqual(payload["metadata"]["method"], "rule_based_title_extraction")
            self.assertIn("Заголовок", payload["metadata"]["title"])
            self.assertIn("Первый абзац", search_text)
            self.assertIn("paragraph", {unit["type"] for unit in payload["units"]})
            self.assertIn("table", {unit["type"] for unit in payload["units"]})
            self.assertIn("table_cell", {unit["type"] for unit in payload["units"]})

    def test_docx_preserves_body_order_and_basic_semantics(self) -> None:
        with tempfile.TemporaryDirectory() as input_dir, tempfile.TemporaryDirectory() as output_dir:
            source_path = Path(input_dir) / "ordered.docx"
            document = Document()
            document.add_heading("Раздел 1", level=1)
            document.add_paragraph("Абзац до таблицы")
            table = document.add_table(rows=1, cols=2)
            table.rows[0].cells[0].text = "A"
            table.rows[0].cells[1].text = "B"
            document.add_paragraph("Пункт списка", style="List Bullet")
            document.add_paragraph("Абзац после таблицы")
            document.save(str(source_path))

            result = run_convert_folder(_docx_converter_config(input_dir, output_dir))

            self.assertEqual(result.status, "success")
            manifest_records = [
                json.loads(line)
                for line in (result.run_dir / "manifest.jsonl").read_text(encoding="utf-8").splitlines()
                if line.strip()
            ]
            document_dir = result.run_dir / manifest_records[0]["output_dir"]
            payload = json.loads((document_dir / "document.v1.json").read_text(encoding="utf-8"))
            semantic_units = [
                unit
                for unit in payload["units"]
                if unit["type"] in {"section", "paragraph", "table", "list_item"}
            ]

            self.assertEqual(semantic_units[0]["type"], "section")
            self.assertEqual(semantic_units[0]["text"], "Раздел 1")
            self.assertIn("semantic_style_inferred", semantic_units[0]["quality"]["flags"])
            self.assertEqual(semantic_units[1]["text"], "Абзац до таблицы")
            self.assertEqual(semantic_units[2]["type"], "table")
            self.assertEqual(semantic_units[3]["type"], "list_item")
            self.assertIn("semantic_style_inferred", semantic_units[3]["quality"]["flags"])
            self.assertEqual(semantic_units[4]["text"], "Абзац после таблицы")
            self.assertLess(semantic_units[2]["order"], semantic_units[4]["order"])

    def test_docx_extracts_formula_header_footer_and_footnotes(self) -> None:
        with tempfile.TemporaryDirectory() as input_dir, tempfile.TemporaryDirectory() as output_dir:
            source_path = Path(input_dir) / "semantic.docx"
            document = Document()
            document.sections[0].header.paragraphs[0].text = "Верхний колонтитул"
            document.sections[0].footer.paragraphs[0].text = "Нижний колонтитул"
            document.add_paragraph("E = mc^2")
            document.add_paragraph("Основной текст")
            document.save(str(source_path))
            _add_footnotes_xml(source_path, "Текст сноски")

            result = run_convert_folder(_docx_converter_config(input_dir, output_dir))

            self.assertEqual(result.status, "success")
            manifest_records = [
                json.loads(line)
                for line in (result.run_dir / "manifest.jsonl").read_text(encoding="utf-8").splitlines()
                if line.strip()
            ]
            document_dir = result.run_dir / manifest_records[0]["output_dir"]
            payload = json.loads((document_dir / "document.v1.json").read_text(encoding="utf-8"))
            validate_payload(payload, "document.v1.schema.json")

            units_by_type = {unit["type"]: unit for unit in payload["units"]}
            self.assertEqual(units_by_type["formula"]["text"], "E = mc^2")
            self.assertIn("semantic_structure_inferred", units_by_type["formula"]["quality"]["flags"])
            self.assertEqual(units_by_type["header"]["text"], "Верхний колонтитул")
            self.assertEqual(units_by_type["footer"]["text"], "Нижний колонтитул")
            self.assertEqual(units_by_type["footnote"]["text"], "Текст сноски")
            search_text = (document_dir / "search_text.txt").read_text(encoding="utf-8")
            self.assertIn("E = mc^2", search_text)
            self.assertIn("Текст сноски", search_text)

    def test_runner_rejects_docx_with_excessive_archive_entries(self) -> None:
        with tempfile.TemporaryDirectory() as input_dir, tempfile.TemporaryDirectory() as output_dir:
            source_path = Path(input_dir) / "oversized-entries.docx"
            document = Document()
            document.add_paragraph("normal docx")
            document.save(str(source_path))

            with patch("doc_converter.converters.docx.pipeline.MAX_DOCX_ARCHIVE_ENTRIES", 1):
                result = run_convert_folder(_docx_converter_config(input_dir, output_dir))

            summary = json.loads((result.run_dir / "summary.json").read_text(encoding="utf-8"))
            self.assertEqual(result.status, "failed")
            self.assertEqual(summary["failed_files"], 1)
            self.assertEqual(summary["failed_reasons"], {"DocxSecurityError": 1})

    def test_runner_rejects_docx_with_excessive_uncompressed_size(self) -> None:
        with tempfile.TemporaryDirectory() as input_dir, tempfile.TemporaryDirectory() as output_dir:
            source_path = Path(input_dir) / "oversized-bytes.docx"
            document = Document()
            document.add_paragraph("normal docx")
            document.save(str(source_path))

            with patch("doc_converter.converters.docx.pipeline.MAX_DOCX_UNCOMPRESSED_BYTES", 1):
                result = run_convert_folder(_docx_converter_config(input_dir, output_dir))

            summary = json.loads((result.run_dir / "summary.json").read_text(encoding="utf-8"))
            self.assertEqual(result.status, "failed")
            self.assertEqual(summary["failed_files"], 1)
            self.assertEqual(summary["failed_reasons"], {"DocxSecurityError": 1})

    def test_docx_table_cells_preserve_multiline_text_and_formula_metadata(self) -> None:
        with tempfile.TemporaryDirectory() as input_dir, tempfile.TemporaryDirectory() as output_dir:
            source_path = Path(input_dir) / "table-formula.docx"
            document = Document()
            table = document.add_table(rows=2, cols=2)
            table.rows[0].cells[0].text = "Показатель"
            table.rows[0].cells[1].text = "Значение"
            table.rows[1].cells[0].text = "Расчёт"

            formula_cell = table.rows[1].cells[1]
            formula_cell.paragraphs[0].text = "S = a * b"
            formula_cell.add_paragraph("Примечание к формуле")
            document.save(str(source_path))

            result = run_convert_folder(_docx_converter_config(input_dir, output_dir))

            self.assertEqual(result.status, "success")
            manifest_records = [
                json.loads(line)
                for line in (result.run_dir / "manifest.jsonl").read_text(encoding="utf-8").splitlines()
                if line.strip()
            ]
            document_dir = result.run_dir / manifest_records[0]["output_dir"]
            payload = json.loads((document_dir / "document.v1.json").read_text(encoding="utf-8"))

            table_cells = [unit for unit in payload["units"] if unit["type"] == "table_cell"]
            formula_cell_unit = next(unit for unit in table_cells if unit["text"].startswith("S = a * b"))

            self.assertEqual(formula_cell_unit["text"], "S = a * b\nПримечание к формуле")
            self.assertEqual(formula_cell_unit["formula"]["linear_text"], "S = a * b")
            self.assertEqual(formula_cell_unit["formula"]["calc_expr"], "S = a * b")

            search_text = (document_dir / "search_text.txt").read_text(encoding="utf-8")
            self.assertIn("S = a * b Примечание к формуле", search_text)

    def test_docx_preserves_subscript_superscript_and_inline_drawing_placeholders(self) -> None:
        with tempfile.TemporaryDirectory() as input_dir, tempfile.TemporaryDirectory() as output_dir:
            source_path = Path(input_dir) / "formula-markup.docx"
            image_path = Path(input_dir) / "inline.png"
            _write_tiny_png(image_path)

            document = Document()

            formula_with_subscripts = document.add_paragraph()
            formula_with_subscripts.add_run("С")
            run = formula_with_subscripts.add_run("НГ")
            run.font.subscript = True
            formula_with_subscripts.add_run(" = НГ x К")
            run = formula_with_subscripts.add_run("в")
            run.font.subscript = True
            formula_with_subscripts.add_run(" x L (6),")

            paragraph_with_superscript = document.add_paragraph()
            paragraph_with_superscript.add_run("P")
            run = paragraph_with_superscript.add_run("j")
            run.font.superscript = True
            paragraph_with_superscript.add_run(" - описание ресурса")

            formula_with_inline_drawing = document.add_paragraph()
            formula_with_inline_drawing.add_run("j = 1 ")
            formula_with_inline_drawing.add_run().add_picture(str(image_path))
            formula_with_inline_drawing.add_run(" J, где:")

            document.save(str(source_path))

            INLINE_GLYPH_CACHE.clear()
            with patch("doc_converter.converters.docx.inline_glyph._recognize_inline_glyph", return_value="÷"):
                result = run_convert_folder(_docx_converter_config(input_dir, output_dir))

            self.assertEqual(result.status, "success")
            manifest_records = [
                json.loads(line)
                for line in (result.run_dir / "manifest.jsonl").read_text(encoding="utf-8").splitlines()
                if line.strip()
            ]
            document_dir = result.run_dir / manifest_records[0]["output_dir"]
            payload = json.loads((document_dir / "document.v1.json").read_text(encoding="utf-8"))

            formula_units = [unit for unit in payload["units"] if unit["type"] == "formula"]
            paragraph_units = [unit for unit in payload["units"] if unit["type"] == "paragraph"]

            self.assertIn("С_(НГ) = НГ x К_(в) x L (6),", [unit["text"] for unit in formula_units])
            self.assertIn("P^(j) - описание ресурса", [unit["text"] for unit in paragraph_units])

            self.assertIn("j = 1 ÷ J, где:", [unit["text"] for unit in formula_units])

            search_text = (document_dir / "search_text.txt").read_text(encoding="utf-8")
            self.assertIn("С_(НГ) = НГ x К_(в) x L (6),", search_text)
            self.assertIn("P^(j) - описание ресурса", search_text)
            self.assertIn("j = 1 ÷ J, где:", search_text)
            self.assertNotIn("[INLINE_DRAWING:", search_text)

    def test_docx_recognizes_inline_symbol_drawings(self) -> None:
        with tempfile.TemporaryDirectory() as input_dir, tempfile.TemporaryDirectory() as output_dir:
            source_path = Path(input_dir) / "formula-symbol.docx"
            image_path = Path(input_dir) / "division.png"
            _write_symbol_png(image_path, "÷")

            document = Document()
            formula = document.add_paragraph()
            formula.add_run("j = 1 ")
            formula.add_run().add_picture(str(image_path))
            formula.add_run(" J, где:")
            document.save(str(source_path))

            INLINE_GLYPH_CACHE.clear()
            with patch("doc_converter.converters.docx.inline_glyph._recognize_inline_glyph", return_value="÷"):
                result = run_convert_folder(_docx_converter_config(input_dir, output_dir))

            self.assertEqual(result.status, "success")
            manifest_records = [
                json.loads(line)
                for line in (result.run_dir / "manifest.jsonl").read_text(encoding="utf-8").splitlines()
                if line.strip()
            ]
            document_dir = result.run_dir / manifest_records[0]["output_dir"]
            payload = json.loads((document_dir / "document.v1.json").read_text(encoding="utf-8"))

            formula_texts = [unit["text"] for unit in payload["units"] if unit["type"] == "formula"]
            self.assertIn("j = 1 ÷ J, где:", formula_texts)

            search_text = (document_dir / "search_text.txt").read_text(encoding="utf-8")
            self.assertIn("j = 1 ÷ J, где:", search_text)

    def test_docx_merges_formula_continuation_lines_split_into_adjacent_paragraphs(self) -> None:
        with tempfile.TemporaryDirectory() as input_dir, tempfile.TemporaryDirectory() as output_dir:
            source_path = Path(input_dir) / "formula-continuation.docx"

            document = Document()
            document.add_paragraph("A = B +")
            document.add_paragraph("+ C (1),")
            document.add_paragraph("D = E x")
            document.add_paragraph("x (F - 1) (2),")
            document.save(str(source_path))

            result = run_convert_folder(_docx_converter_config(input_dir, output_dir))

            self.assertEqual(result.status, "success")
            manifest_records = [
                json.loads(line)
                for line in (result.run_dir / "manifest.jsonl").read_text(encoding="utf-8").splitlines()
                if line.strip()
            ]
            document_dir = result.run_dir / manifest_records[0]["output_dir"]
            payload = json.loads((document_dir / "document.v1.json").read_text(encoding="utf-8"))

            formula_texts = [unit["text"] for unit in payload["units"] if unit["type"] == "formula"]
            paragraph_texts = [unit["text"] for unit in payload["units"] if unit["type"] == "paragraph"]

            self.assertIn("A = B + C (1),", formula_texts)
            self.assertIn("D = E x (F - 1) (2),", formula_texts)
            self.assertNotIn("+ C (1),", paragraph_texts)
            self.assertNotIn("x (F - 1) (2),", paragraph_texts)

            search_text = (document_dir / "search_text.txt").read_text(encoding="utf-8")
            self.assertIn("A = B + C (1),", search_text)
            self.assertIn("D = E x (F - 1) (2),", search_text)

    def test_docx_recognizes_additional_inline_math_symbol_drawings(self) -> None:
        with tempfile.TemporaryDirectory() as input_dir, tempfile.TemporaryDirectory() as output_dir:
            source_path = Path(input_dir) / "formula-extra-symbols.docx"
            symbols = ["∏", "∂", "∇", "∅", "∀", "∃", "<", ">", "+", "-", "="]

            document = Document()
            expected_texts: list[str] = []
            for index, symbol in enumerate(symbols, start=1):
                image_path = Path(input_dir) / f"symbol-{index}.png"
                _write_symbol_png(image_path, symbol)
                formula = document.add_paragraph()
                formula.add_run(f"k{index} = A ")
                formula.add_run().add_picture(str(image_path))
                formula.add_run(" B")
                expected_texts.append(f"k{index} = A {symbol} B")

            document.save(str(source_path))

            INLINE_GLYPH_CACHE.clear()
            symbol_by_asset = {
                f"image{index}.png": symbol
                for index, symbol in enumerate(symbols, start=1)
            }
            with patch(
                "doc_converter.converters.docx.inline_glyph._recognize_inline_glyph",
                side_effect=lambda _blob, asset_name, _drawing_extent: symbol_by_asset.get(asset_name),
            ):
                result = run_convert_folder(_docx_converter_config(input_dir, output_dir))

            self.assertEqual(result.status, "success")
            manifest_records = [
                json.loads(line)
                for line in (result.run_dir / "manifest.jsonl").read_text(encoding="utf-8").splitlines()
                if line.strip()
            ]
            document_dir = result.run_dir / manifest_records[0]["output_dir"]
            payload = json.loads((document_dir / "document.v1.json").read_text(encoding="utf-8"))

            formula_texts = [unit["text"] for unit in payload["units"] if unit["type"] == "formula"]
            expected_formula_options = [
                ("k1 = A ∏ B",),
                ("k2 = A ∂ B",),
                ("k3 = A ∇ B",),
                ("k4 = A ∅ B",),
                ("k5 = A ∀ B",),
                ("k6 = A ∃ B",),
                ("k7 = A < B",),
                ("k8 = A > B",),
                ("k9 = A + B",),
                ("k10 = A - B",),
                ("k11 = A = B",),
            ]
            search_text = (document_dir / "search_text.txt").read_text(encoding="utf-8")
            for options in expected_formula_options:
                with self.subTest(expected_formula_options=options):
                    self.assertTrue(any(option in formula_texts for option in options), formula_texts)
            for expected_search in expected_texts:
                with self.subTest(expected_search=expected_search):
                    self.assertIn(expected_search, search_text)

    def test_inline_glyph_fonts_prefer_bundled_assets(self) -> None:
        bundled_fonts = bundled_font_paths()

        self.assertTrue(bundled_fonts)

        fonts = _available_inline_glyph_fonts()

        self.assertTrue(fonts)
        self.assertEqual(Path(fonts[0]), bundled_fonts[0])
        self.assertEqual(Path(fonts[0]).parent.name, "fonts")
        self.assertEqual(Path(fonts[0]).parent.parent.name, "assets")
        self.assertNotIn("DejaVuSans.ttf", fonts)

    def test_mathtype_wmf_formula_assembly_interleaves_symbol_chunks(self) -> None:
        chunks = [
            WmfTextChunk(text="n1N", face="Times New Roman", charset=0, height=-384, order=0),
            WmfTextChunk(text="=÷", face="Symbol", charset=1, height=-384, order=1),
        ]

        self.assertEqual(_assemble_mathtype_wmf_formula(chunks), "n = 1 ÷ N")

    def test_extract_wmf_text_chunks_rejects_oversized_blob(self) -> None:
        with patch("doc_converter.converters.docx.formulas.wmf.MAX_WMF_BYTES", 8):
            with self.assertRaisesRegex(WmfParseLimitError, "maximum allowed size"):
                _extract_wmf_text_chunks(b"123456789")

    def test_extract_wmf_text_chunks_rejects_excessive_record_count(self) -> None:
        with patch("doc_converter.converters.docx.formulas.wmf.MAX_WMF_RECORDS", 3):
            with self.assertRaisesRegex(WmfParseLimitError, "record count exceeds"):
                _extract_wmf_text_chunks(_build_test_wmf_blob())

    def test_mathtype_wmf_formula_assembly_restores_scripts(self) -> None:
        chunks = [
            WmfTextChunk(text="(", face="Times New Roman", charset=0, height=-384, order=0),
            WmfTextChunk(text="СЦ", face="Times New Roman", charset=204, height=-384, order=1),
            WmfTextChunk(text=")", face="Times New Roman", charset=0, height=-384, order=2),
            WmfTextChunk(text="тек", face="Times New Roman", charset=204, height=-222, order=3),
            WmfTextChunk(text="k", face="Times New Roman", charset=0, height=-222, order=4),
        ]

        self.assertEqual(_assemble_mathtype_wmf_formula(chunks), "(СЦ)_(тек)^(k)")

    def test_mathtype_wmf_formula_ir_classifies_simple_scripts(self) -> None:
        chunks = [
            WmfTextChunk(text="(", face="Times New Roman", charset=0, height=-384, order=0),
            WmfTextChunk(text="СЦ", face="Times New Roman", charset=204, height=-384, order=1),
            WmfTextChunk(text=")", face="Times New Roman", charset=0, height=-384, order=2),
            WmfTextChunk(text="тек", face="Times New Roman", charset=204, height=-222, order=3),
            WmfTextChunk(text="k", face="Times New Roman", charset=0, height=-222, order=4),
        ]

        formula_ir = _build_wmf_formula_ir(chunks)

        self.assertEqual(formula_ir.layout_class, "base_with_scripts")
        self.assertEqual(formula_ir.base_text, "(СЦ)")
        self.assertEqual(formula_ir.script_text, "текk")

    def test_mathtype_wmf_formula_assembly_restores_double_sum_text(self) -> None:
        chunks = [
            WmfTextChunk(text="тек", face="Times New Roman", charset=204, height=-222, order=0),
            WmfTextChunk(text="11", face="Times New Roman", charset=0, height=-222, order=1),
            WmfTextChunk(text="ОТЗТСЦV", face="Times New Roman", charset=204, height=-384, order=2),
            WmfTextChunk(text="IN", face="Times New Roman", charset=0, height=-222, order=3),
            WmfTextChunk(text="nini", face="Times New Roman", charset=0, height=-222, order=4),
            WmfTextChunk(text="in", face="Times New Roman", charset=0, height=-222, order=5),
            WmfTextChunk(text="==", face="Symbol", charset=1, height=-222, order=6),
            WmfTextChunk(text="=", face="Symbol", charset=1, height=-384, order=7),
        ]

        self.assertEqual(
            _assemble_mathtype_wmf_formula(chunks),
            "ОТ_(тек) = sum_(i=1)^I sum_(n=1)^N ЗТ_(ni) × СЦ_(n) × V_(i)",
        )

    def test_mathtype_wmf_formula_assembly_coalesces_split_known_tokens(self) -> None:
        chunks = [
            WmfTextChunk(text="тек", face="Times New Roman", charset=204, height=-222, order=0),
            WmfTextChunk(text="11", face="Times New Roman", charset=0, height=-222, order=1),
            WmfTextChunk(text="ОТ", face="Times New Roman", charset=204, height=-384, order=2),
            WmfTextChunk(text="ЗТСЦV", face="Times New Roman", charset=204, height=-384, order=3),
            WmfTextChunk(text="IN", face="Times New Roman", charset=0, height=-222, order=4),
            WmfTextChunk(text="nini", face="Times New Roman", charset=0, height=-222, order=5),
            WmfTextChunk(text="in", face="Times New Roman", charset=0, height=-222, order=6),
            WmfTextChunk(text="==", face="Symbol", charset=1, height=-222, order=7),
            WmfTextChunk(text="=", face="Symbol", charset=1, height=-384, order=8),
        ]

        self.assertEqual(
            _assemble_mathtype_wmf_formula(chunks),
            "ОТ_(тек) = sum_(i=1)^I sum_(n=1)^N ЗТ_(ni) × СЦ_(n) × V_(i)",
        )

    def test_mathtype_wmf_formula_assembly_restores_known_421pr_formulas(self) -> None:
        cases = [
            (
                [
                    WmfTextChunk(text="тек", face="Times New Roman", charset=204, height=-222, order=0),
                    WmfTextChunk(text="11", face="Times New Roman", charset=0, height=-222, order=1),
                    WmfTextChunk(text="ОТмЗТСЦV", face="Times New Roman", charset=204, height=-384, order=2),
                    WmfTextChunk(text="IK", face="Times New Roman", charset=0, height=-222, order=3),
                    WmfTextChunk(text="kiki", face="Times New Roman", charset=0, height=-222, order=4),
                    WmfTextChunk(text="ik", face="Times New Roman", charset=0, height=-222, order=5),
                    WmfTextChunk(text="==", face="Symbol", charset=1, height=-222, order=6),
                    WmfTextChunk(text="=", face="Symbol", charset=1, height=-384, order=7),
                ],
                "ОТм_(тек) = sum_(i=1)^I sum_(k=1)^K ЗТ_(ki) × СЦ_(k) × V_(i)",
            ),
            (
                [
                    WmfTextChunk(text="t", face="Times New Roman", charset=0, height=-222, order=0),
                    WmfTextChunk(text="1", face="Times New Roman", charset=0, height=-222, order=1),
                    WmfTextChunk(text="PPV", face="Times New Roman", charset=0, height=-384, order=2),
                    WmfTextChunk(text="T", face="Times New Roman", charset=0, height=-222, order=3),
                    WmfTextChunk(text="t", face="Times New Roman", charset=0, height=-222, order=4),
                    WmfTextChunk(text="ii", face="Times New Roman", charset=0, height=-222, order=5),
                    WmfTextChunk(text="t", face="Times New Roman", charset=0, height=-222, order=6),
                    WmfTextChunk(text="=", face="Symbol", charset=1, height=-222, order=7),
                    WmfTextChunk(text="=", face="Symbol", charset=1, height=-384, order=8),
                ],
                "P^(t) = sum_(i=1)^I P_(i)^(t) × V_(i)",
            ),
            (
                [
                    WmfTextChunk(text="с", face="Times New Roman", charset=204, height=-222, order=0),
                    WmfTextChunk(text="Ц", face="Times New Roman", charset=204, height=-384, order=1),
                    WmfTextChunk(text="С", face="Times New Roman", charset=204, height=-384, order=2),
                    WmfTextChunk(text="Т", face="Times New Roman", charset=204, height=-384, order=3),
                    WmfTextChunk(text="=", face="Symbol", charset=1, height=-384, order=4),
                ],
                "С_(маш.р) = Ц_(а) / Т_(с)",
            ),
            (
                [
                    WmfTextChunk(text="k1", face="Times New Roman", charset=0, height=-384, order=0),
                    WmfTextChunk(text="К", face="Times New Roman", charset=204, height=-384, order=1),
                    WmfTextChunk(text="=÷", face="Symbol", charset=1, height=-384, order=2),
                ],
                "k = 1 ÷ K",
            ),
        ]

        for chunks, expected in cases:
            with self.subTest(expected=expected):
                self.assertEqual(_assemble_mathtype_wmf_formula(chunks), expected)

    def test_mathtype_wmf_formula_assembly_restores_known_521pr_formulas(self) -> None:
        cases = [
            (
                [
                    WmfTextChunk(text="ср1", face="Times New Roman", charset=204, height=-222, order=0),
                    WmfTextChunk(text="З = З  , (2)", face="Times New Roman", charset=204, height=-384, order=1),
                    WmfTextChunk(text="смр", face="Times New Roman", charset=204, height=-222, order=2),
                    WmfTextChunk(text="Т", face="Times New Roman", charset=204, height=-222, order=3),
                    WmfTextChunk(text="К", face="Times New Roman", charset=204, height=-384, order=4),
                ],
                "З_(ср) = З_(1) × К_(смрТ)",
            ),
            (
                [
                    WmfTextChunk(text="(", face="Symbol", charset=1, height=-496, order=0),
                    WmfTextChunk(text=")", face="Symbol", charset=1, height=-496, order=1),
                    WmfTextChunk(text=" =   , (3)", face="Times New Roman", charset=0, height=-384, order=2),
                    WmfTextChunk(text="пнр", face="Times New Roman", charset=204, height=-222, order=3),
                    WmfTextChunk(text="ii", face="Times New Roman", charset=0, height=-222, order=4),
                    WmfTextChunk(text="ЗТЗ", face="Times New Roman", charset=204, height=-384, order=5),
                ],
                "З_(пнр) = sum_(i) Т_(i) × З_(i)",
            ),
            (
                [
                    WmfTextChunk(text="пнр", face="Times New Roman", charset=204, height=-222, order=0),
                    WmfTextChunk(text="1", face="Times New Roman", charset=0, height=-222, order=1),
                    WmfTextChunk(text="Т", face="Times New Roman", charset=204, height=-222, order=2),
                    WmfTextChunk(text="З = З  К, (4)", face="Times New Roman", charset=204, height=-384, order=3),
                    WmfTextChunk(text="i", face="Times New Roman", charset=0, height=-222, order=4),
                ],
                "З_(i) = З_(1) × К_(пнрТ)^(i)",
            ),
            (
                [
                    WmfTextChunk(text="(", face="Symbol", charset=1, height=-503, order=0),
                    WmfTextChunk(text=")", face="Symbol", charset=1, height=-503, order=1),
                    WmfTextChunk(text="эмэм", face="Times New Roman", charset=204, height=-222, order=2),
                    WmfTextChunk(text="С = Э  Ц, (5)", face="Times New Roman", charset=204, height=-384, order=3),
                    WmfTextChunk(text="ii", face="Times New Roman", charset=0, height=-222, order=4),
                ],
                "С_(эм) = sum_(i) Э_(i) × Ц_(эмi)",
            ),
            (
                [
                    WmfTextChunk(text="(", face="Symbol", charset=1, height=-496, order=0),
                    WmfTextChunk(text=")", face="Symbol", charset=1, height=-496, order=1),
                    WmfTextChunk(text="мат", face="Times New Roman", charset=204, height=-222, order=2),
                    WmfTextChunk(text="С = М  Ц, (6)", face="Times New Roman", charset=204, height=-384, order=3),
                    WmfTextChunk(text="ii", face="Times New Roman", charset=0, height=-222, order=4),
                ],
                "С_(мат) = sum_(i) М_(i) × Ц_(i)",
            ),
        ]

        for chunks, expected in cases:
            with self.subTest(expected=expected):
                self.assertEqual(_assemble_mathtype_wmf_formula(chunks), expected)

    def test_mathtype_wmf_formula_assembly_restores_known_904pr_formulas(self) -> None:
        cases = [
            (
                [
                    WmfTextChunk(text="ОЦ(1),", face="Times New Roman", charset=204, height=-384, order=0),
                    WmfTextChunk(text="min(,2)", face="Times New Roman", charset=0, height=-222, order=1),
                    WmfTextChunk(text="nmn", face="Times New Roman", charset=0, height=-222, order=2),
                    WmfTextChunk(text="nmn", face="Times New Roman", charset=0, height=-222, order=3),
                    WmfTextChunk(text="*+*", face="Symbol", charset=1, height=-384, order=4),
                    WmfTextChunk(text="=", face="Symbol", charset=1, height=-384, order=5),
                    WmfTextChunk(text="+", face="Symbol", charset=1, height=-384, order=6),
                ],
                "ОЦ_(а) = (Х_(св) × n + Х_(сп) × m) / (n + m)",
            ),
            (
                [
                    WmfTextChunk(text="ОЦ(2),", face="Times New Roman", charset=204, height=-384, order=0),
                    WmfTextChunk(text="nm", face="Times New Roman", charset=0, height=-384, order=1),
                    WmfTextChunk(text="nm", face="Times New Roman", charset=0, height=-384, order=2),
                    WmfTextChunk(text="*+*", face="Symbol", charset=1, height=-384, order=3),
                    WmfTextChunk(text="=", face="Symbol", charset=1, height=-384, order=4),
                    WmfTextChunk(text="+", face="Symbol", charset=1, height=-384, order=5),
                ],
                "ОЦ_(а) = (Х_(св) × n + Х_(сп) × m) / (n + m)",
            ),
            (
                [
                    WmfTextChunk(text="ОЦ(3),", face="Times New Roman", charset=204, height=-384, order=0),
                    WmfTextChunk(text="2", face="Times New Roman", charset=0, height=-384, order=1),
                    WmfTextChunk(text="+", face="Symbol", charset=1, height=-384, order=2),
                    WmfTextChunk(text="=", face="Symbol", charset=1, height=-384, order=3),
                ],
                "ОЦ_(а) = (Х_(св) + Х_(сп)) / 2",
            ),
            (
                [
                    WmfTextChunk(text="12", face="Times New Roman", charset=0, height=-222, order=0),
                    WmfTextChunk(text="()()()", face="Symbol", charset=1, height=-503, order=1),
                    WmfTextChunk(text="Х(4),", face="Times New Roman", charset=204, height=-384, order=2),
                    WmfTextChunk(text="nn", face="Times New Roman", charset=0, height=-222, order=3),
                    WmfTextChunk(text="n", face="Times New Roman", charset=0, height=-222, order=4),
                    WmfTextChunk(text="xvxvxv", face="Times New Roman", charset=0, height=-384, order=5),
                    WmfTextChunk(text="vvv", face="Times New Roman", charset=0, height=-222, order=6),
                    WmfTextChunk(text="*+*++*", face="Symbol", charset=1, height=-384, order=7),
                    WmfTextChunk(text="=", face="Symbol", charset=1, height=-384, order=8),
                    WmfTextChunk(text="++", face="Symbol", charset=1, height=-384, order=9),
                ],
                "Х_(св) = (x_(1) × v_(1) + x_(2) × v_(2) + ... + x_(n) × v_(n)) / (v_(1) + v_(2) + ... + v_(n))",
            ),
            (
                [
                    WmfTextChunk(text="Х(5),", face="Times New Roman", charset=204, height=-384, order=0),
                    WmfTextChunk(text="m", face="Times New Roman", charset=0, height=-222, order=1),
                    WmfTextChunk(text="xxx", face="Times New Roman", charset=0, height=-384, order=2),
                    WmfTextChunk(text="m", face="Times New Roman", charset=0, height=-222, order=3),
                    WmfTextChunk(text="++", face="Symbol", charset=1, height=-384, order=4),
                    WmfTextChunk(text="=", face="Symbol", charset=1, height=-384, order=5),
                ],
                "Х_(сп) = (x_(1) + x_(2) + ... + x_(m)) / m",
            ),
            (
                [
                    WmfTextChunk(text="ср", face="Times New Roman", charset=204, height=-222, order=0),
                    WmfTextChunk(text="0,25", face="Times New Roman", charset=0, height=-384, order=1),
                    WmfTextChunk(text="Х(7),", face="Times New Roman", charset=204, height=-384, order=2),
                    WmfTextChunk(text="s", face="Times New Roman", charset=0, height=-384, order=3),
                    WmfTextChunk(text="<=*", face="Symbol", charset=1, height=-384, order=4),
                ],
                "s <= 0,25 × Х_(ср)",
            ),
        ]

        for chunks, expected in cases:
            with self.subTest(expected=expected):
                self.assertEqual(_assemble_mathtype_wmf_formula(chunks), expected)

    def test_mathtype_wmf_formula_assembly_restores_known_534pr_formulas(self) -> None:
        cases = [
            (
                [
                    WmfTextChunk(text="зт", face="Times New Roman", charset=204, height=-222, order=0),
                    WmfTextChunk(text="текTnинф", face="Times New Roman", charset=204, height=-222, order=1),
                    WmfTextChunk(text="ср", face="Times New Roman", charset=204, height=-222, order=2),
                    WmfTextChunk(text="C", face="Times New Roman", charset=0, height=-384, order=3),
                    WmfTextChunk(text="СЦ = К    К (1),", face="Times New Roman", charset=204, height=-384, order=4),
                    WmfTextChunk(text="t", face="Times New Roman", charset=0, height=-384, order=5),
                ],
                "СЦ_(зттек) = КТ_(n) × С_(1ср) / t_(ср) × К_(инф)",
            ),
            (
                [
                    WmfTextChunk(text="n", face="Times New Roman", charset=0, height=-222, order=0),
                    WmfTextChunk(text="зт", face="Times New Roman", charset=204, height=-222, order=1),
                    WmfTextChunk(text="1", face="Times New Roman", charset=0, height=-222, order=2),
                    WmfTextChunk(text="текii", face="Times New Roman", charset=204, height=-222, order=3),
                    WmfTextChunk(text="i=1", face="Times New Roman", charset=0, height=-222, order=4),
                    WmfTextChunk(text="ОТ = СЦ  T (2),", face="Times New Roman", charset=204, height=-384, order=5),
                ],
                "ОТ_(1) = sum_(i=1)^n СЦ_(зтiтек) × Т_(i)",
            ),
            (
                [
                    WmfTextChunk(text="зт", face="Times New Roman", charset=204, height=-222, order=0),
                    WmfTextChunk(text="21 ", face="Times New Roman", charset=0, height=-222, order=1),
                    WmfTextChunk(text="текТ", face="Times New Roman", charset=204, height=-222, order=2),
                    WmfTextChunk(text="ОТ = СЦ  Т  К (3),", face="Times New Roman", charset=204, height=-384, order=3),
                ],
                "ОТ_(2) = СЦ_(зт1тек) × Т × КТ",
            ),
            (
                [
                    WmfTextChunk(text="n", face="Times New Roman", charset=0, height=-222, order=0),
                    WmfTextChunk(text="1 ", face="Times New Roman", charset=0, height=-222, order=1),
                    WmfTextChunk(text="ср1ip", face="Times New Roman", charset=204, height=-222, order=2),
                    WmfTextChunk(text="i=1", face="Times New Roman", charset=0, height=-222, order=3),
                    WmfTextChunk(text="С = С  1 + K K + ПВ (4),", face="Times New Roman", charset=204, height=-384, order=4),
                ],
                "С_(1ср) = С_(1) × (1 + sum_(i=1)^n К_(i) + К_(р)) + ПВ",
            ),
        ]

        for chunks, expected in cases:
            with self.subTest(expected=expected):
                self.assertEqual(_assemble_mathtype_wmf_formula(chunks), expected)

    def test_mathtype_wmf_formula_assembly_restores_known_1pr_formulas(self) -> None:
        cases = [
            (
                [
                    WmfTextChunk(text="ВрПВрЭ", face="Times New Roman", charset=204, height=-222, order=0),
                    WmfTextChunk(text="НН", face="Times New Roman", charset=204, height=-384, order=1),
                    WmfTextChunk(text="=", face="Symbol", charset=1, height=-384, order=2),
                ],
                "Н_(ВрП) = sum Н_(ВрЭ)",
            ),
            (
                [
                    WmfTextChunk(text="факт", face="Times New Roman", charset=204, height=-222, order=0),
                    WmfTextChunk(text="ВрЭ", face="Times New Roman", charset=204, height=-222, order=1),
                    WmfTextChunk(text="пзротп", face="Times New Roman", charset=204, height=-222, order=2),
                    WmfTextChunk(text="ЗТ", face="Times New Roman", charset=204, height=-384, order=3),
                    WmfTextChunk(text="Ч100", face="Times New Roman", charset=204, height=-384, order=4),
                    WmfTextChunk(text="Н", face="Times New Roman", charset=204, height=-384, order=5),
                    WmfTextChunk(text="100", face="Times New Roman", charset=0, height=-384, order=6),
                    WmfTextChunk(text="ННН60", face="Times New Roman", charset=204, height=-384, order=7),
                    WmfTextChunk(text="=", face="Symbol", charset=1, height=-384, order=8),
                    WmfTextChunk(text="-++", face="Symbol", charset=1, height=-384, order=9),
                ],
                "Н_(ВрЭ) = ЗТ_(эСР) × 100 / (Ч_(факт) × [100 - (Н_(пзр) + Н_(о) + Н_(тп))] × 60)",
            ),
            (
                [
                    WmfTextChunk(text="ЗТ", face="Times New Roman", charset=204, height=-384, order=0),
                    WmfTextChunk(text="ЗТ", face="Times New Roman", charset=204, height=-384, order=1),
                    WmfTextChunk(text="n", face="Times New Roman", charset=0, height=-384, order=2),
                    WmfTextChunk(text="=", face="Symbol", charset=1, height=-384, order=3),
                ],
                "ЗТ_(эСР) = sum_(i=1)^n ЗТ_(э) / n",
            ),
            (
                [
                    WmfTextChunk(text="ЗТ", face="Times New Roman", charset=204, height=-384, order=0),
                    WmfTextChunk(text="ЗТ", face="Times New Roman", charset=204, height=-384, order=1),
                    WmfTextChunk(text="V", face="Times New Roman", charset=0, height=-384, order=2),
                    WmfTextChunk(text="=", face="Symbol", charset=1, height=-384, order=3),
                ],
                "ЗТ_(э) = ЗТ / V",
            ),
            (
                [
                    WmfTextChunk(text="min", face="Times New Roman", charset=0, height=-222, order=0),
                    WmfTextChunk(text="t", face="Times New Roman", charset=0, height=-384, order=1),
                    WmfTextChunk(text="К1,5", face="Times New Roman", charset=204, height=-384, order=2),
                    WmfTextChunk(text="t", face="Times New Roman", charset=0, height=-384, order=3),
                    WmfTextChunk(text="=", face="Symbol", charset=1, height=-384, order=4),
                ],
                "К_(уст) = t_(max) / t_(min) <= 1,5",
            ),
        ]

        for chunks, expected in cases:
            with self.subTest(expected=expected):
                self.assertEqual(_assemble_mathtype_wmf_formula(chunks), expected)

    def test_formula_representation_adds_latex_and_calc_expr_for_double_sum(self) -> None:
        representation = _formula_representation_from_text(
            "ОТ_(тек) = sum_(i=1)^I sum_(n=1)^N ЗТ_(ni) × СЦ_(n) × V_(i) (1.1),"
        )

        self.assertIsNotNone(representation)
        assert representation is not None
        self.assertEqual(representation["source_format"], "mathtype_wmf_text_records")
        self.assertIn(r"\sum_{i=1}^{I}", representation["display_latex"])
        self.assertIn("OT_tek = sum", representation["calc_expr"])

    def test_formula_representation_normalizes_known_sums_and_ranges(self) -> None:
        sum_representation = _formula_representation_from_text(
            "ОТм_(тек) = sum_(i=1)^I sum_(k=1)^K ЗТ_(ki) × СЦ_(k) × V_(i) (1.2)"
        )
        range_representation = _formula_representation_from_text("k = 1 ÷ K, где:")

        self.assertIsNotNone(sum_representation)
        assert sum_representation is not None
        self.assertIn(r"\sum_{i=1}^{I}", sum_representation["display_latex"])
        self.assertIn(r"\sum_{k=1}^{K}", sum_representation["display_latex"])
        self.assertIn("OTm_tek = sum", sum_representation["calc_expr"])

        self.assertIsNotNone(range_representation)
        assert range_representation is not None
        self.assertEqual(range_representation["linear_text"], "k = 1 ÷ K")
        self.assertEqual(range_representation["display_latex"], r"k = 1 \div K")
        self.assertEqual(range_representation["provenance"]["parser_path"], "mathtype_wmf_text_records")

    def test_formula_representation_normalizes_known_521pr_native_formulas(self) -> None:
        product_representation = _formula_representation_from_text("З_(ср) = З_(1) × К_(смрТ) (2)")
        sum_representation = _formula_representation_from_text("С_(эм) = sum_(i) Э_(i) × Ц_(эмi) (5)")

        self.assertIsNotNone(product_representation)
        assert product_representation is not None
        self.assertEqual(product_representation["source_format"], "mathtype_wmf_text_records")
        self.assertEqual(product_representation["calc_expr"], "Z_sr = Z_1 * K_smrT")
        self.assertIn("formula_native_wmf_pattern_recovered", product_representation["warnings"])

        self.assertIsNotNone(sum_representation)
        assert sum_representation is not None
        self.assertEqual(sum_representation["source_format"], "mathtype_wmf_text_records")
        self.assertIn(r"\sum_{i}", sum_representation["display_latex"])
        self.assertEqual(sum_representation["calc_expr"], "S_em = sum(E[i] * C_em[i] for i in I)")

    def test_formula_representation_normalizes_known_1pr_native_formulas(self) -> None:
        process_representation = _formula_representation_from_text("Н_(ВрП) = sum Н_(ВрЭ) (3)")
        element_representation = _formula_representation_from_text(
            "Н_(ВрЭ) = ЗТ_(эСР) × 100 / (Ч_(факт) × [100 - (Н_(пзр) + Н_(о) + Н_(тп))] × 60) (4)"
        )
        process_data_representation = _formula_representation_from_text(
            "Н_(ВрИ) = ЗТ_(Иср) × 100 / (Ч_(общ) × [100 - (Н_(пзр) + Н_(о))] × 60) (23)"
        )
        lab_element_representation = _formula_representation_from_text(
            "Н_(ВрЭл) = ЗТ_(эСРл) × 100 / (Ч_(общ) × [100 - (Н_(пзр) + Н_(о))] × 60) (31)"
        )
        average_representation = _formula_representation_from_text("ЗТ_(эСР) = sum_(i=1)^n ЗТ_(э) / n (5)")
        unit_representation = _formula_representation_from_text("ЗТ_(э) = ЗТ / V (6)")
        stability_representation = _formula_representation_from_text("К_(уст) = t_(max) / t_(min) <= 1,5 (8)")

        self.assertIsNotNone(process_representation)
        assert process_representation is not None
        self.assertEqual(process_representation["source_format"], "mathtype_wmf_text_records")
        self.assertEqual(process_representation["calc_expr"], "N_VrP = sum(N_VrE)")
        self.assertIn(r"\sum", process_representation["display_latex"])
        self.assertIn("formula_native_wmf_pattern_recovered", process_representation["warnings"])
        self.assertIn("calculation_expression_requires_domain_variable_binding", process_representation["warnings"])

        self.assertIsNotNone(element_representation)
        assert element_representation is not None
        self.assertEqual(element_representation["source_format"], "mathtype_wmf_text_records")
        self.assertEqual(
            element_representation["calc_expr"],
            "N_VrE = ZT_eSR * 100 / (Ch_fact * (100 - (N_pzr + N_o + N_tp)) * 60)",
        )
        self.assertIn(r"\frac{\mathrm{ЗТ}_{\text{эСР}} \times 100}", element_representation["display_latex"])
        self.assertIn("formula_native_wmf_pattern_recovered", element_representation["warnings"])

        self.assertIsNotNone(process_data_representation)
        assert process_data_representation is not None
        self.assertEqual(process_data_representation["source_format"], "mathtype_wmf_text_records")
        self.assertEqual(
            process_data_representation["calc_expr"],
            "N_VrI = ZT_Isr * 100 / (Ch_obsh * (100 - (N_pzr + N_o)) * 60)",
        )
        self.assertIn(r"\frac{\mathrm{ЗТ}_{\text{Иср}} \times 100}", process_data_representation["display_latex"])
        self.assertIn("formula_native_wmf_pattern_recovered", process_data_representation["warnings"])

        self.assertIsNotNone(lab_element_representation)
        assert lab_element_representation is not None
        self.assertEqual(lab_element_representation["source_format"], "mathtype_wmf_text_records")
        self.assertEqual(
            lab_element_representation["calc_expr"],
            "N_VrEl = ZT_eSRl * 100 / (Ch_obsh * (100 - (N_pzr + N_o)) * 60)",
        )
        self.assertIn(r"\frac{\mathrm{ЗТ}_{\text{эСРл}} \times 100}", lab_element_representation["display_latex"])
        self.assertIn("formula_native_wmf_pattern_recovered", lab_element_representation["warnings"])

        self.assertIsNotNone(average_representation)
        assert average_representation is not None
        self.assertEqual(average_representation["source_format"], "mathtype_wmf_text_records")
        self.assertIn(r"\sum_{i=1}^{n}", average_representation["display_latex"])
        self.assertEqual(average_representation["calc_expr"], "ZT_eSR = sum(ZT_e[i] for i in range(1, n + 1)) / n")
        self.assertIn("formula_native_wmf_pattern_recovered", average_representation["warnings"])

        self.assertIsNotNone(unit_representation)
        assert unit_representation is not None
        self.assertEqual(unit_representation["source_format"], "mathtype_wmf_text_records")
        self.assertEqual(unit_representation["calc_expr"], "ZT_e = ZT / V")
        self.assertIn(r"\frac{\mathrm{ЗТ}}{V}", unit_representation["display_latex"])
        self.assertIn("formula_native_wmf_pattern_recovered", unit_representation["warnings"])

        self.assertIsNotNone(stability_representation)
        assert stability_representation is not None
        self.assertEqual(stability_representation["source_format"], "mathtype_wmf_text_records")
        self.assertEqual(stability_representation["calc_expr"], "K_ust = t_max / t_min")
        self.assertIn(r"\le 1{,}5", stability_representation["display_latex"])
        self.assertIn("formula_native_wmf_pattern_recovered", stability_representation["warnings"])
        self.assertIn("formula_constraint_not_encoded_in_calc_expr", stability_representation["warnings"])

    def test_formula_representation_recovers_noisy_1pr_fraction_family(self) -> None:
        process_data_representation = _formula_representation_from_text("ЗТЧ100Н[100(НН)]60=-+_(ВрИпзро) (23),")
        lab_element_representation = _formula_representation_from_text("ЗТЧ100Н[100(НН)]60=-+_(ВрЭлпзро) (31),")

        self.assertIsNotNone(process_data_representation)
        assert process_data_representation is not None
        self.assertEqual(
            process_data_representation["linear_text"],
            "Н_(ВрИ) = ЗТ_(Иср) × 100 / (Ч_(общ) × [100 - (Н_(пзр) + Н_(о))] × 60)",
        )
        self.assertEqual(process_data_representation["source_format"], "mathtype_wmf_text_records")
        self.assertEqual(
            process_data_representation["calc_expr"],
            "N_VrI = ZT_Isr * 100 / (Ch_obsh * (100 - (N_pzr + N_o)) * 60)",
        )

        self.assertIsNotNone(lab_element_representation)
        assert lab_element_representation is not None
        self.assertEqual(
            lab_element_representation["linear_text"],
            "Н_(ВрЭл) = ЗТ_(эСРл) × 100 / (Ч_(общ) × [100 - (Н_(пзр) + Н_(о))] × 60)",
        )
        self.assertEqual(lab_element_representation["source_format"], "mathtype_wmf_text_records")
        self.assertEqual(
            lab_element_representation["calc_expr"],
            "N_VrEl = ZT_eSRl * 100 / (Ch_obsh * (100 - (N_pzr + N_o)) * 60)",
        )

    def test_formula_representation_recovers_noisy_1pr_average_family(self) -> None:
        process_average_representation = _formula_representation_from_text("ЗТЗТn=_(Исрф) (24),")
        unit_process_representation = _formula_representation_from_text("ЗЗТV=_(фактф)^(1) (25),")

        self.assertIsNotNone(process_average_representation)
        assert process_average_representation is not None
        self.assertEqual(
            process_average_representation["linear_text"],
            "ЗТ_(Иср) = sum_(ф=1)^n_(ф) ЗТ_(1факт) / n_(ф)",
        )
        self.assertEqual(process_average_representation["source_format"], "mathtype_wmf_text_records")
        self.assertEqual(
            process_average_representation["calc_expr"],
            "ZT_Isr = sum(ZT_1fakt[f] for f in range(1, n_f + 1)) / n_f",
        )
        self.assertIn(r"\sum_{\text{ф}=1}^{n_{\text{ф}}}", process_average_representation["display_latex"])

        self.assertIsNotNone(unit_process_representation)
        assert unit_process_representation is not None
        self.assertEqual(
            unit_process_representation["linear_text"],
            "ЗТ_(1факт) = ЗТ_(Vфакт) / V_(ф)",
        )
        self.assertEqual(unit_process_representation["source_format"], "mathtype_wmf_text_records")
        self.assertEqual(unit_process_representation["calc_expr"], "ZT_1fakt = ZT_Vfakt / V_f")
        self.assertIn(r"\frac{\mathrm{ЗТ}_{\text{Vфакт}}}{V_{\text{ф}}}", unit_process_representation["display_latex"])

    def test_formula_representation_recovers_noisy_1pr_tech_break_formula(self) -> None:
        tech_break_representation = _formula_representation_from_text("ТН100Т=_(вр) (9),")

        self.assertIsNotNone(tech_break_representation)
        assert tech_break_representation is not None
        self.assertEqual(tech_break_representation["linear_text"], "Н_(тп) = Т_(тп) × 100 / Т_(вр)")
        self.assertEqual(tech_break_representation["source_format"], "mathtype_wmf_text_records")
        self.assertEqual(tech_break_representation["calc_expr"], "N_tp = T_tp * 100 / T_vr")
        self.assertIn(r"\frac{\mathrm{Т}_{\text{тп}} \times 100}{\mathrm{Т}_{\text{вр}}}", tech_break_representation["display_latex"])

    def test_formula_representation_recovers_noisy_1pr_wage_and_worker_time_formulas(self) -> None:
        wage_representation = _formula_representation_from_text("ЗЗt=_(м) (10),")
        worker_time_representation = _formula_representation_from_text("ТН=_(рабВрЭ1раб)^(1) (13),")

        self.assertIsNotNone(wage_representation)
        assert wage_representation is not None
        self.assertEqual(wage_representation["linear_text"], "З_(ч) = З_(срм) / t_(м)")
        self.assertEqual(wage_representation["source_format"], "mathtype_wmf_text_records")
        self.assertEqual(wage_representation["calc_expr"], "Z_ch = Z_srm / t_m")
        self.assertIn(
            r"\frac{\mathrm{З}_{\text{срм}}}{t_{\text{м}}}",
            wage_representation["display_latex"],
        )

        self.assertIsNotNone(worker_time_representation)
        assert worker_time_representation is not None
        self.assertEqual(worker_time_representation["linear_text"], "Т_(1раб) = sum Н_(ВрЭ1раб)")
        self.assertEqual(worker_time_representation["source_format"], "mathtype_wmf_text_records")
        self.assertEqual(worker_time_representation["calc_expr"], "T_1rab = sum(N_VrE1rab)")
        self.assertIn(r"\sum", worker_time_representation["display_latex"])
        self.assertIn("formula_native_wmf_pattern_recovered", worker_time_representation["warnings"])
        self.assertIn(
            "calculation_expression_requires_domain_variable_binding",
            worker_time_representation["warnings"],
        )

    def test_formula_representation_recovers_noisy_1pr_participation_average_formula(self) -> None:
        participation_average_representation = _formula_representation_from_text("ККЧ=_(факт) (11),")

        self.assertIsNotNone(participation_average_representation)
        assert participation_average_representation is not None
        self.assertEqual(participation_average_representation["linear_text"], "К_(ср) = К_(уч) / Ч_(факт)")
        self.assertEqual(participation_average_representation["source_format"], "mathtype_wmf_text_records")
        self.assertEqual(participation_average_representation["calc_expr"], "K_sr = K_uch / Ch_fact")
        self.assertIn(
            r"\frac{\mathrm{К}_{\text{уч}}}{\mathrm{Ч}_{\text{факт}}}",
            participation_average_representation["display_latex"],
        )

    def test_formula_representation_recovers_noisy_1pr_participation_formula(self) -> None:
        participation_representation = _formula_representation_from_text("ТКЧТКН=_(учiВрП) (12),")

        self.assertIsNotNone(participation_representation)
        assert participation_representation is not None
        self.assertEqual(participation_representation["linear_text"], "К_(уч) = ТК × Ч_(i) × Т_(1раб) / Н_(ВрП)")
        self.assertEqual(participation_representation["source_format"], "mathtype_wmf_text_records")
        self.assertEqual(participation_representation["calc_expr"], "K_uch = TK * Ch_i * T_1rab / N_VrP")
        self.assertIn(
            r"\frac{\mathrm{ТК} \times \mathrm{Ч}_{i} \times \mathrm{Т}_{1\text{раб}}}{\mathrm{Н}_{\text{ВрП}}}",
            participation_representation["display_latex"],
        )

    def test_formula_representation_recovers_noisy_1pr_cameral_participation_family(self) -> None:
        participation_average_representation = _formula_representation_from_text("ККЧ=_(учКсрКобщ) (35),")
        participation_representation = _formula_representation_from_text("ТКЧТКТ=_(КАМобщ) (36),")

        self.assertIsNotNone(participation_average_representation)
        assert participation_average_representation is not None
        self.assertEqual(participation_average_representation["linear_text"], "К_(срК) = К_(учК) / Ч_(общ)")
        self.assertEqual(participation_average_representation["source_format"], "mathtype_wmf_text_records")
        self.assertEqual(participation_average_representation["calc_expr"], "K_srK = K_uchK / Ch_obsh")
        self.assertIn(
            r"\frac{\mathrm{К}_{\text{учК}}}{\mathrm{Ч}_{\text{общ}}}",
            participation_average_representation["display_latex"],
        )

        self.assertIsNotNone(participation_representation)
        assert participation_representation is not None
        self.assertEqual(participation_representation["linear_text"], "К_(учК) = ТК × Ч_(i) × Т_(КАМ1) / Т_(КАМобщ)")
        self.assertEqual(participation_representation["source_format"], "mathtype_wmf_text_records")
        self.assertEqual(participation_representation["calc_expr"], "K_uchK = TK * Ch_i * T_KAM1 / T_KAMobsh")
        self.assertIn(
            r"\frac{\mathrm{ТК} \times \mathrm{Ч}_{i} \times \mathrm{Т}_{\text{КАМ1}}}{\mathrm{Т}_{\text{КАМобщ}}}",
            participation_representation["display_latex"],
        )

    def test_formula_representation_recovers_noisy_1pr_additional_cost_formula(self) -> None:
        additional_cost_representation = _formula_representation_from_text("ДЗН100С= (37),")

        self.assertIsNotNone(additional_cost_representation)
        assert additional_cost_representation is not None
        self.assertEqual(additional_cost_representation["linear_text"], "Н_(ДЗ) = ДЗ × 100 / С_(р)")
        self.assertEqual(additional_cost_representation["source_format"], "mathtype_wmf_text_records")
        self.assertEqual(additional_cost_representation["calc_expr"], "N_DZ = DZ * 100 / S_r")
        self.assertIn(
            r"\frac{\mathrm{ДЗ} \times 100}{\mathrm{С}_{\text{р}}}",
            additional_cost_representation["display_latex"],
        )

    def test_formula_representation_recovers_noisy_1pr_estimated_work_participation_family(self) -> None:
        participation_representation = _formula_representation_from_text("ТКЧТКТ=_(Ообщ) (38),")
        participation_average_representation = _formula_representation_from_text("ККЧ=_(Побщ) (39),")

        self.assertIsNotNone(participation_representation)
        assert participation_representation is not None
        self.assertEqual(participation_representation["linear_text"], "К_(учО) = ТК_(о) × Ч_(Оi) × Т_(Оi) / Т_(Ообщ)")
        self.assertEqual(participation_representation["source_format"], "mathtype_wmf_text_records")
        self.assertEqual(participation_representation["calc_expr"], "K_uchO = TK_o * Ch_Oi * T_Oi / T_Oobsh")
        self.assertIn(
            r"\frac{\mathrm{ТК}_{\text{о}} \times \mathrm{Ч}_{\text{Оi}} \times \mathrm{Т}_{\text{Оi}}}{\mathrm{Т}_{\text{Ообщ}}}",
            participation_representation["display_latex"],
        )

        self.assertIsNotNone(participation_average_representation)
        assert participation_average_representation is not None
        self.assertEqual(participation_average_representation["linear_text"], "К_(срО) = К_(учО) / Ч_(Ообщ)")
        self.assertEqual(participation_average_representation["source_format"], "mathtype_wmf_text_records")
        self.assertEqual(participation_average_representation["calc_expr"], "K_srO = K_uchO / Ch_Oobsh")
        self.assertIn(
            r"\frac{\mathrm{К}_{\text{учО}}}{\mathrm{Ч}_{\text{Ообщ}}}",
            participation_average_representation["display_latex"],
        )

    def test_formula_representation_recovers_noisy_1pr_estimated_work_cost_family(self) -> None:
        technical_means_representation = _formula_representation_from_text("ССИ=_(ТСоТС) (42),")
        machine_representation = _formula_representation_from_text("С(ЦЭ)=_(МоММо) (43),")

        self.assertIsNotNone(technical_means_representation)
        assert technical_means_representation is not None
        self.assertEqual(technical_means_representation["linear_text"], "С_(ТСо) = СИ_(ТС)")
        self.assertEqual(technical_means_representation["source_format"], "mathtype_wmf_text_records")
        self.assertEqual(technical_means_representation["calc_expr"], "S_TSo = SI_TS")
        self.assertIn(
            r"\mathrm{С}_{\text{ТСо}} = \mathrm{СИ}_{\text{ТС}}",
            technical_means_representation["display_latex"],
        )

        self.assertIsNotNone(machine_representation)
        assert machine_representation is not None
        self.assertEqual(machine_representation["linear_text"], "С_(Мо) = Ц_(М) × Э_(Мо)")
        self.assertEqual(machine_representation["source_format"], "mathtype_wmf_text_records")
        self.assertEqual(machine_representation["calc_expr"], "S_Mo = C_M * E_Mo")
        self.assertIn(
            r"\mathrm{С}_{\text{Мо}} = \mathrm{Ц}_{\text{М}} \times \mathrm{Э}_{\text{Мо}}",
            machine_representation["display_latex"],
        )

    def test_formula_representation_recovers_noisy_1pr_technical_cost_family(self) -> None:
        technical_support_representation = _formula_representation_from_text("С(ИЦ)=_(ТСТСiТСi) (15),")
        amortization_representation = _formula_representation_from_text("ВАН=_(с) (17),")
        repair_representation = _formula_representation_from_text("НРАС100= (19),")
        machine_cost_representation = _formula_representation_from_text("С(ЭЦ)=_(ММiМi) (20),")
        material_cost_representation = _formula_representation_from_text("С(МЦ)=_(матiматi) (22),")

        self.assertIsNotNone(technical_support_representation)
        assert technical_support_representation is not None
        self.assertEqual(
            technical_support_representation["linear_text"],
            "С_(ТС) = sum_(i) И_(ТСi) × Ц_(ТСi)",
        )
        self.assertEqual(technical_support_representation["source_format"], "mathtype_wmf_text_records")
        self.assertEqual(
            technical_support_representation["calc_expr"],
            "S_TS = sum(I_TS[i] * C_TS[i] for i in I)",
        )
        self.assertIn(r"\sum_{i}", technical_support_representation["display_latex"])

        self.assertIsNotNone(amortization_representation)
        assert amortization_representation is not None
        self.assertEqual(amortization_representation["linear_text"], "А_(ТС) = В_(с) / Н_(с)")
        self.assertEqual(amortization_representation["source_format"], "mathtype_wmf_text_records")
        self.assertEqual(amortization_representation["calc_expr"], "A_TS = V_s / N_s")
        self.assertIn(r"\frac{\mathrm{В}_{\text{с}}}{\mathrm{Н}_{\text{с}}}", amortization_representation["display_latex"])

        self.assertIsNotNone(repair_representation)
        assert repair_representation is not None
        self.assertEqual(repair_representation["linear_text"], "Р_(ТС) = Н_(р) × А_(ТС) × С_(с) / 100")
        self.assertEqual(repair_representation["source_format"], "mathtype_wmf_text_records")
        self.assertEqual(repair_representation["calc_expr"], "R_TS = N_r * A_TS * S_s / 100")
        self.assertIn(r"\frac{\mathrm{Н}_{\text{р}} \times \mathrm{А}_{\text{ТС}} \times \mathrm{С}_{\text{с}}}{100}", repair_representation["display_latex"])

        self.assertIsNotNone(machine_cost_representation)
        assert machine_cost_representation is not None
        self.assertEqual(machine_cost_representation["linear_text"], "С_(М) = sum_(i) Э_(Мi) × Ц_(Мi)")
        self.assertEqual(machine_cost_representation["source_format"], "mathtype_wmf_text_records")
        self.assertEqual(
            machine_cost_representation["calc_expr"],
            "S_M = sum(E_M[i] * C_M[i] for i in I)",
        )
        self.assertIn(r"\sum_{i}", machine_cost_representation["display_latex"])

        self.assertIsNotNone(material_cost_representation)
        assert material_cost_representation is not None
        self.assertEqual(material_cost_representation["linear_text"], "С_(мат) = sum_(i) М_(i) × Ц_(матi)")
        self.assertEqual(material_cost_representation["source_format"], "mathtype_wmf_text_records")
        self.assertEqual(
            material_cost_representation["calc_expr"],
            "S_mat = sum(M[i] * C_mat[i] for i in I)",
        )
        self.assertIn(r"\sum_{i}", material_cost_representation["display_latex"])

    def test_formula_representation_normalizes_known_904pr_native_formulas(self) -> None:
        weighted_representation = _formula_representation_from_text(
            "ОЦ_(а) = (Х_(св) × n + Х_(сп) × m) / (n + m) (1)"
        )
        volume_representation = _formula_representation_from_text(
            "Х_(св) = (x_(1) × v_(1) + x_(2) × v_(2) + ... + x_(n) × v_(n)) / (v_(1) + v_(2) + ... + v_(n)) (4)"
        )
        inequality_representation = _formula_representation_from_text("s <= 0,25 × Х_(ср) (7)")

        self.assertIsNotNone(weighted_representation)
        assert weighted_representation is not None
        self.assertEqual(weighted_representation["source_format"], "mathtype_wmf_text_records")
        self.assertEqual(weighted_representation["calc_expr"], "OTs_a = ( X_sv * n + X_sp * m ) / ( n + m )")
        self.assertIn("formula_native_wmf_pattern_recovered", weighted_representation["warnings"])

        self.assertIsNotNone(volume_representation)
        assert volume_representation is not None
        self.assertEqual(volume_representation["source_format"], "mathtype_wmf_text_records")
        self.assertIn(r"\ldots", volume_representation["display_latex"])
        self.assertIn("sum(x[i] * v[i]", volume_representation["calc_expr"])
        self.assertIn("calculation_expression_requires_domain_variable_binding", volume_representation["warnings"])

        self.assertIsNotNone(inequality_representation)
        assert inequality_representation is not None
        self.assertEqual(inequality_representation["source_format"], "mathtype_wmf_text_records")
        self.assertEqual(inequality_representation["calc_expr"], "s <= 0.25 * X_sr")
        self.assertIn(r"\le", inequality_representation["display_latex"])

    def test_formula_representation_normalizes_known_534pr_native_formulas(self) -> None:
        rate_representation = _formula_representation_from_text(
            "СЦ_(зттек) = КТ_(n) × С_(1ср) / t_(ср) × К_(инф) (1)"
        )
        sum_representation = _formula_representation_from_text(
            "ОТ_(1) = sum_(i=1)^n СЦ_(зтiтек) × Т_(i) (2)"
        )
        wage_representation = _formula_representation_from_text(
            "С_(1ср) = С_(1) × (1 + sum_(i=1)^n К_(i) + К_(р)) + ПВ (4)"
        )

        self.assertIsNotNone(rate_representation)
        assert rate_representation is not None
        self.assertEqual(rate_representation["source_format"], "mathtype_wmf_text_records")
        self.assertEqual(rate_representation["calc_expr"], "SC_zt_tek = KT_n * S_1sr / t_sr * K_inf")
        self.assertIn("formula_native_wmf_pattern_recovered", rate_representation["warnings"])

        self.assertIsNotNone(sum_representation)
        assert sum_representation is not None
        self.assertEqual(sum_representation["source_format"], "mathtype_wmf_text_records")
        self.assertIn(r"\sum_{i=1}^{n}", sum_representation["display_latex"])
        self.assertIn("sum(SC_zt_i_tek[i] * T[i]", sum_representation["calc_expr"])
        self.assertIn("calculation_expression_requires_domain_variable_binding", sum_representation["warnings"])

        self.assertIsNotNone(wage_representation)
        assert wage_representation is not None
        self.assertEqual(wage_representation["source_format"], "mathtype_wmf_text_records")
        self.assertIn(r"\left(1 + \sum_{i=1}^{n}", wage_representation["display_latex"])
        self.assertIn("sum(K[i] for i in range(1, n + 1))", wage_representation["calc_expr"])

    def test_formula_representation_adds_calc_expr_for_linearized_assignment(self) -> None:
        representation = _formula_representation_from_text("С_(ИГДИ) = С_(П) + С_(К) (1),")

        self.assertIsNotNone(representation)
        assert representation is not None
        self.assertEqual(representation["display_latex"], r"\mathrm{С}_{\text{ИГДИ}} = \mathrm{С}_{\text{П}} + \mathrm{С}_{\text{К}}")
        self.assertEqual(representation["calc_expr"], "S_IGDI = S_P + S_K")
        self.assertEqual(
            representation["variables"],
            {
                "С_ИГДИ": "S_IGDI",
                "С_П": "S_P",
                "С_К": "S_K",
            },
        )
        self.assertIn("formula_calc_expr_is_heuristic", representation["warnings"])
        self.assertIn("calculation_expression_requires_domain_variable_binding", representation["warnings"])

    def test_formula_representation_adds_calc_expr_for_nested_linearized_formula(self) -> None:
        representation = _formula_representation_from_text(
            "ДЗ_(вП) = (С_(Ппз) + ДЗ_(НП) + ДЗ_(режим) + ДЗ_(ноч) + ДЗ_(орг)) x (Д_(ЗПП) x П_(ДЗвыс) + Д_(прочП) - 1) (9),"
        )

        self.assertIsNotNone(representation)
        assert representation is not None
        self.assertEqual(
            representation["calc_expr"],
            "DZ_vP = ( S_Ppz + DZ_NP + DZ_rezhim + DZ_noch + DZ_org ) * ( D_ZPP * P_DZvys + D_prochP - 1 )",
        )
        self.assertEqual(representation["variables"]["ДЗ_вП"], "DZ_vP")
        self.assertEqual(representation["variables"]["П_ДЗвыс"], "P_DZvys")

    def test_formula_representation_supports_identifiers_with_digits(self) -> None:
        representation = _formula_representation_from_text("С_(Свлс) = ПЗ1_(п) + ПЗ2_(п) x S_(влс) (18),")

        self.assertIsNotNone(representation)
        assert representation is not None
        self.assertEqual(
            representation["display_latex"],
            r"\mathrm{С}_{\text{Свлс}} = \mathrm{ПЗ1}_{\text{п}} + \mathrm{ПЗ2}_{\text{п}}  \times  S_{\text{влс}}",
        )
        self.assertEqual(representation["calc_expr"], "S_Svls = PZ1_p + PZ2_p * S_vls")
        self.assertEqual(representation["variables"]["ПЗ1_п"], "PZ1_p")
        self.assertEqual(representation["variables"]["ПЗ2_п"], "PZ2_p")

    def test_formula_representation_supports_percent_and_power(self) -> None:
        representation = _formula_representation_from_text("К_(рост) = (1 + П_(рост)%)^2 (5),")

        self.assertIsNotNone(representation)
        assert representation is not None
        self.assertEqual(representation["calc_expr"], "K_rost = ( 1 + P_rost / 100 ) ** 2")
        self.assertIn(r"\%", representation["display_latex"])
        self.assertEqual(representation["variables"]["П_рост"], "P_rost")

    def test_formula_representation_collapses_repeated_multiply_at_line_wrap(self) -> None:
        representation = _formula_representation_from_text(
            "С_(Свлс) = ПЗ1_(п) + ПЗ2_(п) x\n x S_(влс) (18),"
        )

        self.assertIsNotNone(representation)
        assert representation is not None
        self.assertEqual(representation["linear_text"], "С_(Свлс) = ПЗ1_(п) + ПЗ2_(п) x S_(влс)")
        self.assertEqual(representation["calc_expr"], "S_Svls = PZ1_p + PZ2_p * S_vls")
        self.assertEqual(
            representation["display_latex"],
            r"\mathrm{С}_{\text{Свлс}} = \mathrm{ПЗ1}_{\text{п}} + \mathrm{ПЗ2}_{\text{п}}  \times  S_{\text{влс}}",
        )

    def test_formula_representation_strips_trailing_reference_before_calc_expr(self) -> None:
        representation = _formula_representation_from_text("З^(смр) = Т x З_(ср), (1)")

        self.assertIsNotNone(representation)
        assert representation is not None
        self.assertEqual(representation["linear_text"], "З^(смр) = Т x З_(ср)")
        self.assertEqual(representation["calc_expr"], "Z_smr = T * Z_sr")
        self.assertEqual(
            representation["variables"],
            {
                "З_смр": "Z_smr",
                "Т": "T",
                "З_ср": "Z_sr",
            },
        )

    def test_formula_representation_extracts_numeric_expression_from_chain(self) -> None:
        representation = _formula_representation_from_text(
            "a + b x x = 30 + 0,35 x 200 = 100,0 тыс. руб."
        )

        self.assertIsNotNone(representation)
        assert representation is not None
        self.assertEqual(representation["calc_expr"], "30 + 0.35 * 200")
        self.assertEqual(representation["variables"], {})
        self.assertIn("formula_calc_expr_is_heuristic", representation["warnings"])

    def test_formula_representation_extracts_expression_after_narrative_prefix(self) -> None:
        representation = _formula_representation_from_text(
            "Коэффициент изменения цены разработки проектной и рабочей документации будет равен: (140 + 45) : 140 = 1,32"
        )

        self.assertIsNotNone(representation)
        assert representation is not None
        self.assertEqual(representation["calc_expr"], "( 140 + 45 ) / 140")
        self.assertEqual(representation["variables"], {})
        self.assertIn("formula_calc_expr_is_heuristic", representation["warnings"])

    def test_formula_representation_extracts_first_assignment_from_semicolon_clause(self) -> None:
        representation = _formula_representation_from_text(
            "Производственный корпус объекта цветной металлургии мощностью 200 ед.: a = 30; b = 0,35"
        )

        self.assertIsNotNone(representation)
        assert representation is not None
        self.assertEqual(representation["calc_expr"], "a = 30")
        self.assertEqual(representation["variables"], {"a": "a"})

    def test_formula_representation_extracts_parenthetical_assignment(self) -> None:
        representation = _formula_representation_from_text(
            "k_(з.п) - коэффициент, устанавливающий долю заработной платы производственного персонала в общих затратах на проектирование (k_(з.п) = 0,3 - 0,65)."
        )

        self.assertIsNotNone(representation)
        assert representation is not None
        self.assertEqual(representation["calc_expr"], "k_z_p = 0.3 - 0.65")
        self.assertEqual(representation["variables"]["k_з_п"], "k_z_p")

    def test_formula_representation_supports_square_bracket_grouping(self) -> None:
        representation = _formula_representation_from_text(
            "НЗ_(п) = [С_(ФОТпТН) x (1 + НР) + С_(возТН) + С_(ТС) + С_(М) + С_(авто) + С_(мат)] x (1 + П) (1),"
        )

        self.assertIsNotNone(representation)
        assert representation is not None
        self.assertEqual(
            representation["calc_expr"],
            "NZ_p = ( S_FOTpTN * ( 1 + NR ) + S_vozTN + S_TS + S_M + S_avto + S_mat ) * ( 1 + P )",
        )
        self.assertEqual(representation["variables"]["НЗ_п"], "NZ_p")
        self.assertEqual(representation["variables"]["С_ФОТпТН"], "S_FOTpTN")
        self.assertEqual(representation["variables"]["С_возТН"], "S_vozTN")
        self.assertEqual(representation["variables"]["П"], "P")

    def test_runner_skips_duplicate_docx_and_can_copy_originals(self) -> None:
        with tempfile.TemporaryDirectory() as input_dir, tempfile.TemporaryDirectory() as output_dir:
            first_path = Path(input_dir) / "a.docx"
            second_path = Path(input_dir) / "nested" / "b.docx"
            second_path.parent.mkdir(parents=True, exist_ok=True)

            document = Document()
            document.add_paragraph("Одинаковый текст")
            document.save(str(first_path))
            second_path.write_bytes(first_path.read_bytes())

            result = run_convert_folder(
                ConverterConfig(
                    input_dir=Path(input_dir),
                    output_dir=Path(output_dir),
                    options=ConverterOptions(include_originals=True),
                )
            )

            manifest_records = [
                json.loads(line)
                for line in (result.run_dir / "manifest.jsonl").read_text(encoding="utf-8").splitlines()
                if line.strip()
            ]
            self.assertEqual([record["status"] for record in manifest_records], ["success", "skipped_duplicate"])
            shared_output_dir = manifest_records[0]["output_dir"]
            self.assertEqual(manifest_records[1]["output_dir"], shared_output_dir)
            document_dir = result.run_dir / shared_output_dir
            self.assertTrue((document_dir / "originals" / "a.docx").exists())
            self.assertTrue((document_dir / "originals" / "nested" / "b.docx").exists())


def _add_footnotes_xml(source_path: Path, footnote_text: str) -> None:
    temp_path = source_path.with_suffix(".tmp.docx")
    footnotes_xml = f"""<?xml version="1.0" encoding="UTF-8" standalone="yes"?>
<w:footnotes xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">
  <w:footnote w:type="separator" w:id="-1"><w:p><w:r><w:separator/></w:r></w:p></w:footnote>
  <w:footnote w:id="1"><w:p><w:r><w:t>{footnote_text}</w:t></w:r></w:p></w:footnote>
</w:footnotes>
"""
    with zipfile.ZipFile(source_path, "r") as source_archive, zipfile.ZipFile(temp_path, "w") as target_archive:
        for item in source_archive.infolist():
            target_archive.writestr(item, source_archive.read(item.filename))
        target_archive.writestr("word/footnotes.xml", footnotes_xml.encode("utf-8"))
    temp_path.replace(source_path)


def _build_test_wmf_blob() -> bytes:
    font_params = bytearray(50)
    font_params[0:2] = (-384).to_bytes(2, byteorder="little", signed=True)
    font_params[13] = 0
    font_params[18:34] = b"Times New Roman\x00"

    records = [
        _wmf_record(0x02FB, bytes(font_params)),
        _wmf_record(0x012D, (0).to_bytes(2, byteorder="little", signed=False)),
        _wmf_record(0x0521, (2).to_bytes(2, byteorder="little", signed=True) + b"AB"),
        _wmf_record(0x0521, (2).to_bytes(2, byteorder="little", signed=True) + b"CD"),
        _wmf_record(0x0000, b""),
    ]
    return b"\x00" * 18 + b"".join(records)


def _wmf_record(func: int, params: bytes) -> bytes:
    if len(params) % 2 != 0:
        params += b"\x00"
    size_words = (6 + len(params)) // 2
    return size_words.to_bytes(4, byteorder="little", signed=False) + func.to_bytes(2, byteorder="little", signed=False) + params


def _write_tiny_png(path: Path) -> None:
    png_bytes = base64.b64decode(
        "iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAQAAAC1HAwCAAAAC0lEQVR42mP8/x8AAwMCAO+XlGQAAAAASUVORK5CYII="
    )
    path.write_bytes(png_bytes)


def _write_symbol_png(path: Path, symbol: str) -> None:
    image = Image.new("RGBA", (128, 128), (255, 255, 255, 0))
    draw = ImageDraw.Draw(image)
    font = _load_test_symbol_font(72)
    bbox = draw.textbbox((0, 0), symbol, font=font)
    x = (128 - (bbox[2] - bbox[0])) // 2 - bbox[0]
    y = (128 - (bbox[3] - bbox[1])) // 2 - bbox[1]
    draw.text((x, y), symbol, font=font, fill=(0, 0, 0, 255))
    image.save(path)


def _load_test_symbol_font(size: int) -> ImageFont.ImageFont | ImageFont.FreeTypeFont:
    for bundled_font in bundled_font_paths():
        try:
            return ImageFont.truetype(str(bundled_font), size=size)
        except OSError:
            continue
    for candidate in (
        "C:/Windows/Fonts/cambria.ttc",
        "C:/Windows/Fonts/cambriai.ttf",
        "C:/Windows/Fonts/seguisym.ttf",
        "C:/Windows/Fonts/times.ttf",
        "C:/Windows/Fonts/arial.ttf",
    ):
        try:
            return ImageFont.truetype(candidate, size=size)
        except OSError:
            continue
    return ImageFont.load_default()


def _docx_converter_config(input_dir: str, output_dir: str) -> ConverterConfig:
    return ConverterConfig(
        input_dir=Path(input_dir),
        output_dir=Path(output_dir),
        options=ConverterOptions(formula_recognition=FormulaRecognitionConfig()),
    )


if __name__ == "__main__":
    unittest.main()