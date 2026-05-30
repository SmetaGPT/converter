from __future__ import annotations

import json
import os
import tempfile
import unittest
from contextlib import redirect_stdout
from io import StringIO
from pathlib import Path
from typing import Any, cast
from urllib.parse import unquote, urlparse
from unittest.mock import Mock, patch

from docx import Document
from openpyxl import Workbook, load_workbook

from doc_converter.cli import build_parser, main
from doc_converter.config import AgentRunMetadata, ConverterConfig, ConverterOptions, FormulaRecognitionConfig
from doc_converter.formula_recognition import FormulaRecognitionPostprocessResult
from doc_converter.ocr_runtime import find_ocrmypdf_executable
from doc_converter.runner import ConverterError, run_convert_folder
from doc_converter import schema_validation
from doc_converter.schema_validation import validate_payload


class CliSmokeTests(unittest.TestCase):
    def test_help_parser_contains_convert_folder(self) -> None:
        help_text = build_parser().format_help()
        self.assertIn("convert-folder", help_text)
        self.assertIn("check-ocr", help_text)
        self.assertIn("doctor", help_text)
        self.assertIn("dry-run", help_text)
        self.assertIn("evaluate-formula", help_text)
        self.assertIn("evaluate-document-formulas", help_text)
        self.assertNotIn("--workers", help_text)

    def test_check_ocr_outputs_machine_readable_status(self) -> None:
        buffer = StringIO()
        with redirect_stdout(buffer):
            exit_code = main(["check-ocr", "--output-format=json"])

        self.assertIn(exit_code, {0, 40})
        payload = _load_cli_result(buffer)
        self.assertEqual(payload["command"], "check-ocr")
        self.assertIn(payload["status"], {"ok", "environment_invalid"})
        validate_payload(payload["data"], "ocr-runtime.v1.schema.json")
        self.assertEqual(payload["data"]["schema_version"], "ocr-runtime.v1")
        self.assertIn(payload["data"]["status"], {"ready", "missing"})
        self.assertIn("tools", payload["data"])
        self.assertIn("languages", payload["data"])

    def test_doctor_outputs_machine_readable_status(self) -> None:
        buffer = StringIO()
        with patch("doc_converter.cli.detect_ocr_runtime", return_value=_ready_ocr_payload()):
            with patch("doc_converter.cli._check_font_bundle", return_value=_ready_font_bundle_payload()):
                with patch("doc_converter.cli._check_schema_contracts", return_value=_ready_schema_checks_payload()):
                    with patch(
                        "doc_converter.cli._check_openrouter_reachability",
                        return_value={"provider": None, "model": None, "status": "not_configured", "warnings": []},
                    ):
                        with redirect_stdout(buffer):
                            exit_code = main(["doctor", "--output-format=json"])

        self.assertEqual(exit_code, 0)
        payload = _load_cli_result(buffer)
        self.assertEqual(payload["command"], "doctor")
        self.assertEqual(payload["status"], "ok")
        self.assertEqual(payload["data"]["ocr_runtime"]["status"], "ready")
        self.assertEqual(payload["data"]["font_bundle"]["status"], "ready")
        self.assertEqual(payload["data"]["schemas"]["status"], "ready")

    def test_doctor_reports_environment_invalid(self) -> None:
        buffer = StringIO()
        with patch("doc_converter.cli.detect_ocr_runtime", return_value=_missing_ocr_payload()):
            with patch("doc_converter.cli._check_font_bundle", return_value=_ready_font_bundle_payload()):
                with patch("doc_converter.cli._check_schema_contracts", return_value=_ready_schema_checks_payload()):
                    with patch(
                        "doc_converter.cli._check_openrouter_reachability",
                        return_value={"provider": None, "model": None, "status": "not_configured", "warnings": []},
                    ):
                        with redirect_stdout(buffer):
                            exit_code = main(["doctor", "--output-format=json"])

        self.assertEqual(exit_code, 40)
        payload = _load_cli_result(buffer)
        self.assertEqual(payload["status"], "environment_invalid")
        self.assertEqual(payload["data"]["ocr_runtime"]["status"], "missing")

    def test_dry_run_reports_review_required_for_unsupported_inputs(self) -> None:
        with tempfile.TemporaryDirectory() as input_dir:
            document = Document()
            document.add_paragraph("Dry run candidate")
            document.save(str(Path(input_dir) / "sample.docx"))
            (Path(input_dir) / "notes.txt").write_text("unsupported", encoding="utf-8")

            buffer = StringIO()
            with redirect_stdout(buffer):
                exit_code = main(["dry-run", input_dir, "--output-format=json"])

        self.assertEqual(exit_code, 20)
        payload = _load_cli_result(buffer)
        self.assertEqual(payload["command"], "dry-run")
        self.assertEqual(payload["status"], "review_required")
        self.assertEqual(payload["data"]["supported_files"], 1)
        self.assertEqual(payload["data"]["unsupported_files"], 1)
        self.assertEqual(payload["data"]["route_counts"], {"docx_native": 1})

    def test_convert_folder_reports_input_invalid_exit_code(self) -> None:
        with tempfile.TemporaryDirectory() as output_dir:
            missing = Path(output_dir) / "missing"
            buffer = StringIO()
            with redirect_stdout(buffer):
                exit_code = main(["convert-folder", str(missing), output_dir, "--output-format=json"])

        self.assertEqual(exit_code, 30)
        payload = _load_cli_result(buffer)
        self.assertEqual(payload["status"], "input_invalid")

    def test_evaluate_formula_outputs_machine_readable_result(self) -> None:
        buffer = StringIO()
        with redirect_stdout(buffer):
            exit_code = main(
                [
                    "evaluate-formula",
                    "--calc-expr",
                    "K_rost = ( 1 + P_rost / 100 ) ** 2",
                    "--values",
                    '{"P_rost": 15}',
                    "--output-format=json",
                ]
            )

        self.assertEqual(exit_code, 0)
        payload = _load_cli_result(buffer)
        self.assertEqual(payload["status"], "ok")
        self.assertEqual(payload["data"]["status"], "ok")
        self.assertEqual(payload["data"]["target"], "K_rost")
        self.assertAlmostEqual(payload["data"]["value"], 1.3225)
        self.assertEqual(payload["data"]["used_variables"], {"P_rost": 15.0})

    def test_evaluate_formula_reports_missing_variables(self) -> None:
        buffer = StringIO()
        with redirect_stdout(buffer):
            exit_code = main(
                [
                    "evaluate-formula",
                    "--calc-expr",
                    "R = A + B",
                    "--values",
                    '{"A": 1}',
                    "--output-format=json",
                ]
            )

        self.assertEqual(exit_code, 20)
        payload = _load_cli_result(buffer)
        self.assertEqual(payload["status"], "review_required")
        self.assertEqual(payload["data"]["status"], "missing_variables")
        self.assertEqual(payload["data"]["target"], "R")
        self.assertEqual(payload["data"]["missing_variables"], ["B"])

    def test_evaluate_formula_accepts_utf8_bom_values_file(self) -> None:
        with tempfile.TemporaryDirectory() as temp_dir:
            values_path = Path(temp_dir) / "values.json"
            values_path.write_text('{"P_rost": 15}', encoding="utf-8-sig")

            buffer = StringIO()
            with redirect_stdout(buffer):
                exit_code = main(
                    [
                        "evaluate-formula",
                        "--calc-expr",
                        "K_rost = ( 1 + P_rost / 100 ) ** 2",
                        "--values-file",
                        str(values_path),
                        "--output-format=json",
                    ]
                )

        self.assertEqual(exit_code, 0)
        payload = _load_cli_result(buffer)
        self.assertEqual(payload["status"], "ok")
        self.assertAlmostEqual(payload["data"]["value"], 1.3225)

    def test_evaluate_document_formulas_outputs_machine_readable_results(self) -> None:
        with tempfile.TemporaryDirectory() as temp_dir:
            document_path = Path(temp_dir) / "document.v1.json"
            values_path = Path(temp_dir) / "values.json"
            document_path.write_text(
                json.dumps(_minimal_formula_document_payload(second_calc_expr="S = R + C"), ensure_ascii=False),
                encoding="utf-8",
            )
            values_path.write_text('{"A": 2, "B": 3}', encoding="utf-8")

            buffer = StringIO()
            with redirect_stdout(buffer):
                exit_code = main(
                    [
                        "evaluate-document-formulas",
                        str(document_path),
                        "--values-file",
                        str(values_path),
                        "--output-format=json",
                    ]
                )

        self.assertEqual(exit_code, 10)
        payload = _load_cli_result(buffer)
        self.assertEqual(payload["status"], "partial")
        self.assertEqual(payload["data"]["formula_units"], 2)
        self.assertEqual(payload["data"]["calc_expr_units"], 2)
        self.assertEqual(payload["data"]["evaluated"], 1)
        self.assertEqual(payload["data"]["missing_variables"], 1)
        self.assertEqual(payload["data"]["results"][0]["target"], "R")
        self.assertAlmostEqual(payload["data"]["results"][0]["value"], 5.0)
        self.assertEqual(payload["data"]["results"][1]["status"], "missing_variables")
        self.assertEqual(payload["data"]["results"][1]["missing_variables"], ["C"])

    def test_evaluate_document_formulas_resolves_dependencies(self) -> None:
        with tempfile.TemporaryDirectory() as temp_dir:
            document_path = Path(temp_dir) / "document.v1.json"
            values_path = Path(temp_dir) / "values.json"
            document_path.write_text(
                json.dumps(
                    _minimal_formula_document_payload(
                        first_text="S = R * 2",
                        first_calc_expr="S = R * 2",
                        second_text="R = A + B",
                        second_calc_expr="R = A + B",
                    ),
                    ensure_ascii=False,
                ),
                encoding="utf-8",
            )
            values_path.write_text('{"A": 2, "B": 3}', encoding="utf-8")

            buffer = StringIO()
            with redirect_stdout(buffer):
                exit_code = main(
                    [
                        "evaluate-document-formulas",
                        str(document_path),
                        "--values-file",
                        str(values_path),
                        "--output-format=json",
                    ]
                )

        self.assertEqual(exit_code, 0)
        payload = _load_cli_result(buffer)
        self.assertEqual(payload["status"], "ok")
        self.assertEqual(payload["data"]["evaluated"], 2)
        self.assertEqual(payload["data"]["missing_variables"], 0)
        self.assertAlmostEqual(payload["data"]["results"][0]["value"], 10.0)
        self.assertEqual(payload["data"]["results"][0]["used_variables"], {"R": 5.0})
        self.assertAlmostEqual(payload["data"]["results"][1]["value"], 5.0)

    def test_find_ocrmypdf_prefers_active_python_directory(self) -> None:
        with tempfile.TemporaryDirectory() as temp_dir:
            scripts_dir = Path(temp_dir)
            executable_name = "ocrmypdf.exe" if os.name == "nt" else "ocrmypdf"
            python_name = "python.exe" if os.name == "nt" else "python"
            ocrmypdf_path = scripts_dir / executable_name
            ocrmypdf_path.write_text("", encoding="utf-8")

            with patch("doc_converter.ocr_runtime.sys.executable", str(scripts_dir / python_name)):
                with patch("doc_converter.ocr_runtime.shutil.which", return_value=None):
                    self.assertEqual(find_ocrmypdf_executable(), str(ocrmypdf_path.resolve()))

    def test_schema_lookup_prefers_frozen_bundle_layout(self) -> None:
        source_schema = Path(__file__).resolve().parents[1] / "schemas" / "run.v1.schema.json"

        with tempfile.TemporaryDirectory() as temp_dir:
            bundle_dir = Path(temp_dir)
            bundled_schemas_dir = bundle_dir / "_internal" / "schemas"
            bundled_schemas_dir.mkdir(parents=True)
            (bundled_schemas_dir / "run.v1.schema.json").write_text(
                source_schema.read_text(encoding="utf-8"),
                encoding="utf-8",
            )

            schema_validation._schemas_dir.cache_clear()
            schema_validation._load_validator.cache_clear()
            try:
                with patch("doc_converter.schema_validation.sys.frozen", True, create=True):
                    with patch(
                        "doc_converter.schema_validation.sys.executable",
                        str(bundle_dir / "DocumentConverter.exe"),
                    ):
                        self.assertEqual(schema_validation._schemas_dir(), bundled_schemas_dir.resolve())
                        schema_validation._load_validator("run.v1.schema.json")
            finally:
                schema_validation._schemas_dir.cache_clear()
                schema_validation._load_validator.cache_clear()

    def test_empty_folder_creates_run_package(self) -> None:
        with tempfile.TemporaryDirectory() as input_dir, tempfile.TemporaryDirectory() as output_dir:
            result = run_convert_folder(ConverterConfig(input_dir=Path(input_dir), output_dir=Path(output_dir)))

            self.assertEqual(result.status, "success")
            self.assertEqual(result.discovered_files, 0)
            self.assertEqual(result.supported_files, 0)
            self.assertTrue((result.run_dir / "run.json").exists())
            self.assertTrue((result.run_dir / "manifest.jsonl").exists())
            self.assertTrue((result.run_dir / "processing-log.jsonl").exists())
            self.assertTrue((result.run_dir / "summary.json").exists())
            self.assertTrue((result.run_dir / "errors.jsonl").exists())
            self.assertTrue((result.run_dir / "processed-documents-catalog.json").exists())
            self.assertTrue((result.run_dir / "processed-documents-catalog.xlsx").exists())
            self.assertTrue((result.run_dir / "queue-state.json").exists())
            self.assertTrue((result.run_dir / "documents").is_dir())

            run_payload = json.loads((result.run_dir / "run.json").read_text(encoding="utf-8"))
            summary = json.loads((result.run_dir / "summary.json").read_text(encoding="utf-8"))
            queue_state = json.loads((result.run_dir / "queue-state.json").read_text(encoding="utf-8"))
            catalog = json.loads((result.run_dir / "processed-documents-catalog.json").read_text(encoding="utf-8"))

            validate_payload(run_payload, "run.v1.schema.json")
            validate_payload(summary, "summary.v1.schema.json")
            validate_payload(queue_state, "queue-state.v1.schema.json")
            validate_payload(catalog, "processed-documents-catalog.v1.schema.json")

            self.assertEqual(summary["schema_version"], "summary.v1")
            self.assertEqual(summary["status"], "success")
            self.assertEqual(summary["supported_files"], 0)
            self.assertEqual(catalog["documents"], [])
            self.assertNotIn("workers", run_payload["options"])
            self.assertEqual(run_payload["agent_run_metadata"]["agent_id"], "manual")
            self.assertEqual(run_payload["agent_run_metadata"]["task_id"], result.run_id)

    def test_run_metadata_records_agent_run_metadata(self) -> None:
        with tempfile.TemporaryDirectory() as input_dir, tempfile.TemporaryDirectory() as output_dir:
            result = run_convert_folder(
                ConverterConfig(
                    input_dir=Path(input_dir),
                    output_dir=Path(output_dir),
                    agent_run_metadata=AgentRunMetadata(
                        agent_id="roadmap-agent",
                        agent_version="2026.05",
                        task_id="S1.3",
                        parent_run_id="parent-001",
                    ),
                )
            )

            run_payload = json.loads((result.run_dir / "run.json").read_text(encoding="utf-8"))
            validate_payload(run_payload, "run.v1.schema.json")
            self.assertEqual(
                run_payload["agent_run_metadata"],
                {
                    "agent_id": "roadmap-agent",
                    "agent_version": "2026.05",
                    "task_id": "S1.3",
                    "parent_run_id": "parent-001",
                },
            )

    def test_convert_folder_cli_accepts_agent_metadata_flags(self) -> None:
        with tempfile.TemporaryDirectory() as input_dir, tempfile.TemporaryDirectory() as output_dir:
            buffer = StringIO()
            with redirect_stdout(buffer):
                exit_code = main(
                    [
                        "convert-folder",
                        input_dir,
                        output_dir,
                        "--agent-id",
                        "autonomous-agent",
                        "--agent-version",
                        "v1",
                        "--task-id",
                        "task-123",
                        "--parent-run-id",
                        "parent-123",
                        "--output-format=json",
                    ]
                )

            self.assertEqual(exit_code, 0)
            cli_payload = _load_cli_result(buffer)
            run_payload = json.loads((Path(cli_payload["data"]["run_dir"]) / "run.json").read_text(encoding="utf-8"))
            validate_payload(run_payload, "run.v1.schema.json")
            self.assertEqual(
                run_payload["agent_run_metadata"],
                {
                    "agent_id": "autonomous-agent",
                    "agent_version": "v1",
                    "task_id": "task-123",
                    "parent_run_id": "parent-123",
                },
            )

    def test_dry_run_internal_error_maps_to_exit_code_50(self) -> None:
        with tempfile.TemporaryDirectory() as input_dir:
            buffer = StringIO()
            with patch("doc_converter.cli.build_inventory", side_effect=RuntimeError("boom")):
                with redirect_stdout(buffer):
                    exit_code = main(["dry-run", input_dir, "--output-format=json"])

        self.assertEqual(exit_code, 50)
        payload = _load_cli_result(buffer)
        self.assertEqual(payload["status"], "internal_error")
        self.assertEqual(payload["data"]["error_type"], "RuntimeError")

    def test_run_metadata_omits_formula_recognition_api_key(self) -> None:
        with tempfile.TemporaryDirectory() as env_dir, tempfile.TemporaryDirectory() as input_dir, tempfile.TemporaryDirectory() as output_dir:
            root = Path(env_dir)
            (root / ".env.local").write_text(
                "LLM_PROVIDER=openrouter\n"
                "OPENROUTER_MODEL=deepseek/deepseek-v4-pro\n"
                "FORMULA_MODEL=openai/gpt-4o\n"
                "OPENROUTER_API_KEY=super-secret\n",
                encoding="utf-8",
            )

            with patch("doc_converter.config.Path.cwd", return_value=root):
                options = ConverterOptions()

            result = run_convert_folder(
                ConverterConfig(input_dir=Path(input_dir), output_dir=Path(output_dir), options=options)
            )

            run_payload = json.loads((result.run_dir / "run.json").read_text(encoding="utf-8"))
            validate_payload(run_payload, "run.v1.schema.json")

            self.assertEqual(
                run_payload["options"].get("formula_recognition"),
                {
                    "provider": "openrouter",
                    "model": "openai/gpt-4o",
                    "configured": True,
                },
            )
            self.assertNotIn("api_key", json.dumps(run_payload, ensure_ascii=False))

    def test_runner_invokes_formula_recognition_postprocess_when_configured(self) -> None:
        with tempfile.TemporaryDirectory() as input_dir, tempfile.TemporaryDirectory() as output_dir:
            source_path = Path(input_dir) / "formula.docx"
            document = Document()
            document.add_paragraph("formula postprocess smoke")
            document.save(str(source_path))

            options = ConverterOptions(
                formula_recognition=FormulaRecognitionConfig(
                    provider="openrouter",
                    model="openai/gpt-4o",
                    api_key="secret",
                )
            )

            with patch(
                "doc_converter.run.postprocess.run_formula_recognition_postprocess",
                return_value=FormulaRecognitionPostprocessResult(
                    attempted=2,
                    recognized=1,
                    provider_calls=1,
                    warnings=("formula_recognition_provider_failed",),
                    artifact_path="formula-recognition.jsonl",
                ),
            ) as mocked_postprocess:
                result = run_convert_folder(
                    ConverterConfig(
                        input_dir=Path(input_dir),
                        output_dir=Path(output_dir),
                        options=options,
                    )
                )

            mocked_postprocess.assert_called_once()
            manifest_record = json.loads((result.run_dir / "manifest.jsonl").read_text(encoding="utf-8").splitlines()[0])
            self.assertEqual(manifest_record["formula_recognition_attempted"], 2)
            self.assertEqual(manifest_record["formula_recognition_recognized"], 1)
            self.assertEqual(manifest_record["formula_recognition_provider_calls"], 1)
            self.assertEqual(manifest_record["formula_recognition_results_path"], "formula-recognition.jsonl")
            self.assertIn("formula_recognition_provider_failed", manifest_record["warnings"])

    def test_run_metadata_serializes_local_formula_backend_without_api_key(self) -> None:
        with tempfile.TemporaryDirectory() as input_dir, tempfile.TemporaryDirectory() as output_dir:
            source_path = Path(input_dir) / "formula.docx"
            document = Document()
            document.add_paragraph("formula local backend smoke")
            document.save(str(source_path))

            options = ConverterOptions(
                formula_recognition=FormulaRecognitionConfig(
                    local_backend="tesseract",
                )
            )

            result = run_convert_folder(
                ConverterConfig(input_dir=Path(input_dir), output_dir=Path(output_dir), options=options)
            )

            run_payload = json.loads((result.run_dir / "run.json").read_text(encoding="utf-8"))
            validate_payload(run_payload, "run.v1.schema.json")
            self.assertEqual(
                run_payload["options"].get("formula_recognition"),
                {
                    "local_backend": "tesseract",
                    "configured": True,
                },
            )

    def test_processed_documents_catalog_describes_output_folder_and_status(self) -> None:
        with tempfile.TemporaryDirectory() as input_dir, tempfile.TemporaryDirectory() as output_dir:
            source_path = Path(input_dir) / "source.docx"
            document = Document()
            document.add_paragraph("catalog smoke")
            document.save(str(source_path))

            result = run_convert_folder(ConverterConfig(input_dir=Path(input_dir), output_dir=Path(output_dir)))

            catalog = json.loads((result.run_dir / "processed-documents-catalog.json").read_text(encoding="utf-8"))
            validate_payload(catalog, "processed-documents-catalog.v1.schema.json")
            workbook = load_workbook(result.run_dir / "processed-documents-catalog.xlsx")
            sheet = workbook["Документы"]

            self.assertEqual(len(catalog["documents"]), 1)
            entry = catalog["documents"][0]
            self.assertEqual(entry["original_filename"], "source.docx")
            self.assertEqual(entry["relative_input_path"], "source.docx")
            self.assertEqual(entry["status"], "success")
            self.assertEqual(entry["status_label"], "Успешная обработка")
            self.assertTrue(str(entry["output_dir"]).startswith("documents/sha256_"))
            self.assertTrue(str(entry["output_folder_name"]).startswith("sha256_"))
            self.assertIsNone(entry["issue"])
            self.assertEqual(sheet["A2"].value, "source.docx")
            self.assertEqual(sheet["C2"].value, entry["output_folder_name"])
            self.assertEqual(sheet["F2"].value, "document.v1.json")
            self.assertEqual(sheet["G2"].value, "search_text.txt")
            self.assertIsNotNone(sheet["C2"].hyperlink)
            self.assertIsNotNone(sheet["F2"].hyperlink)
            self.assertIsNotNone(sheet["G2"].hyperlink)
            self.assertEqual(_file_uri_to_path(sheet["C2"].hyperlink.target), result.run_dir / entry["output_dir"])
            self.assertEqual(
                _file_uri_to_path(sheet["F2"].hyperlink.target),
                result.run_dir / entry["output_dir"] / "document.v1.json",
            )
            self.assertEqual(
                _file_uri_to_path(sheet["G2"].hyperlink.target),
                result.run_dir / entry["output_dir"] / "search_text.txt",
            )

    def test_failed_document_writes_review_required_file_and_failed_reason(self) -> None:
        with tempfile.TemporaryDirectory() as input_dir, tempfile.TemporaryDirectory() as output_dir:
            source_path = Path(input_dir) / "sample.docx"
            document = Document()
            document.add_paragraph("boom")
            document.save(str(source_path))

            converter = Mock()
            converter.convert.side_effect = OSError("broken docx extractor")
            with patch("doc_converter.run.orchestration.get_converter", return_value=converter):
                result = run_convert_folder(
                    ConverterConfig(input_dir=Path(input_dir), output_dir=Path(output_dir), options=ConverterOptions())
                )

            summary = json.loads((result.run_dir / "summary.json").read_text(encoding="utf-8"))
            review_required_records = [
                json.loads(line)
                for line in (result.run_dir / "review-required.jsonl").read_text(encoding="utf-8").splitlines()
                if line.strip()
            ]
            catalog = json.loads((result.run_dir / "processed-documents-catalog.json").read_text(encoding="utf-8"))

            self.assertEqual(result.status, "failed")
            self.assertEqual(summary["failed_files"], 1)
            self.assertEqual(summary["review_required_files"], 1)
            self.assertEqual(summary["failed_reasons"]["OSError"], 1)
            self.assertEqual(review_required_records[0]["status"], "failed")
            self.assertEqual(catalog["documents"][0]["status"], "failed")
            self.assertEqual(catalog["documents"][0]["status_label"], "Ошибка")
            self.assertEqual(catalog["documents"][0]["issue"], "broken docx extractor")
            validate_payload(review_required_records[0], "review-required.v1.schema.json")

    def test_mixed_folder_reports_unsupported_inputs(self) -> None:
        with tempfile.TemporaryDirectory() as input_dir, tempfile.TemporaryDirectory() as output_dir:
            document = Document()
            document.add_paragraph("supported")
            document.save(str(Path(input_dir) / "source.docx"))
            (Path(input_dir) / "ignored.txt").write_text("ignored", encoding="utf-8")

            result = run_convert_folder(ConverterConfig(input_dir=Path(input_dir), output_dir=Path(output_dir)))

            summary = json.loads((result.run_dir / "summary.json").read_text(encoding="utf-8"))
            log_records = [
                json.loads(line)
                for line in (result.run_dir / "processing-log.jsonl").read_text(encoding="utf-8").splitlines()
                if line.strip()
            ]

            self.assertEqual(result.discovered_files, 2)
            self.assertEqual(result.supported_files, 1)
            self.assertEqual(summary["discovered_files"], 2)
            self.assertEqual(summary["supported_files"], 1)
            self.assertEqual(summary["unsupported_files"], 1)
            unsupported_records = [
                record for record in log_records if record.get("event") == "document_skipped_unsupported"
            ]
            self.assertEqual(len(unsupported_records), 1)
            self.assertEqual(unsupported_records[0]["relative_path"], "ignored.txt")
            self.assertEqual(unsupported_records[0]["status"], "skipped_unsupported")

    def test_xlsx_input_creates_machine_readable_cell_units(self) -> None:
        with tempfile.TemporaryDirectory() as input_dir, tempfile.TemporaryDirectory() as output_dir:
            workbook = Workbook()
            sheet = workbook.active
            if sheet is None:
                self.fail("Workbook must have an active worksheet")
            sheet.title = "Расчет"
            sheet["A1"] = "Цена"
            sheet["B1"] = 10
            sheet["A2"] = "Итого"
            sheet["B2"] = "=B1*2"
            workbook.save(Path(input_dir) / "sample.xlsx")

            result = run_convert_folder(ConverterConfig(input_dir=Path(input_dir), output_dir=Path(output_dir)))

            manifest = [
                json.loads(line)
                for line in (result.run_dir / "manifest.jsonl").read_text(encoding="utf-8").splitlines()
                if line.strip()
            ]
            document_path = result.run_dir / str(manifest[0]["output_dir"]) / "document.v1.json"
            payload = json.loads(document_path.read_text(encoding="utf-8"))

            self.assertEqual(result.status, "success")
            self.assertEqual(result.supported_files, 1)
            self.assertEqual(manifest[0]["format"], "xlsx")
            self.assertEqual(manifest[0]["route"], "xlsx_native")
            self.assertEqual(manifest[0]["formula_cells"], 1)
            self.assertEqual(payload["source"]["format"], "xlsx")
            self.assertEqual(payload["processing"]["route"], "xlsx_native")
            formula_cells = [unit for unit in payload["units"] if (unit.get("cell") or {}).get("formula")]
            self.assertEqual(len(formula_cells), 1)
            self.assertEqual(formula_cells[0]["cell"]["address"], "B2")
            self.assertEqual(formula_cells[0]["cell"]["formula"], "=B1*2")
            self.assertIn("B2: =B1*2", (document_path.parent / "search_text.txt").read_text(encoding="utf-8"))

    def test_repeated_run_reuses_previous_output_for_unchanged_input(self) -> None:
        with tempfile.TemporaryDirectory() as input_dir, tempfile.TemporaryDirectory() as output_dir:
            source_path = Path(input_dir) / "sample.docx"
            document = Document()
            document.add_paragraph("Resume candidate")
            document.save(str(source_path))

            first_result = run_convert_folder(
                ConverterConfig(input_dir=Path(input_dir), output_dir=Path(output_dir), options=ConverterOptions())
            )
            first_manifest = [
                json.loads(line)
                for line in (first_result.run_dir / "manifest.jsonl").read_text(encoding="utf-8").splitlines()
                if line.strip()
            ]

            with patch("doc_converter.run.orchestration.get_converter") as get_converter:
                second_result = run_convert_folder(
                    ConverterConfig(input_dir=Path(input_dir), output_dir=Path(output_dir), options=ConverterOptions())
                )
                get_converter.assert_not_called()

            second_manifest = [
                json.loads(line)
                for line in (second_result.run_dir / "manifest.jsonl").read_text(encoding="utf-8").splitlines()
                if line.strip()
            ]

            self.assertEqual(first_manifest[0]["status"], "success")
            self.assertEqual(second_manifest[0]["status"], "success")
            self.assertTrue(second_manifest[0]["reused_previous_output"])
            self.assertEqual(second_manifest[0]["resumed_from_run_id"], first_result.run_id)

    def test_missing_input_directory_fails(self) -> None:
        with tempfile.TemporaryDirectory() as output_dir:
            missing = Path(output_dir) / "missing"
            with self.assertRaises(ConverterError):
                run_convert_folder(ConverterConfig(input_dir=missing, output_dir=Path(output_dir)))

    def test_equal_input_and_output_directory_fails_before_run_starts(self) -> None:
        with tempfile.TemporaryDirectory() as temp_dir:
            shared_dir = Path(temp_dir) / "shared"
            shared_dir.mkdir()

            with self.assertRaises(ConverterError):
                run_convert_folder(ConverterConfig(input_dir=shared_dir, output_dir=shared_dir))

            self.assertFalse((shared_dir / "runs").exists())

    def test_output_directory_inside_input_fails_before_run_starts(self) -> None:
        with tempfile.TemporaryDirectory() as temp_dir:
            input_dir = Path(temp_dir) / "input"
            input_dir.mkdir()
            output_dir = input_dir / "out"

            with self.assertRaises(ConverterError):
                run_convert_folder(ConverterConfig(input_dir=input_dir, output_dir=output_dir))

            self.assertFalse(output_dir.exists())

    def test_input_directory_inside_output_fails_before_run_starts(self) -> None:
        with tempfile.TemporaryDirectory() as temp_dir:
            output_dir = Path(temp_dir) / "output"
            input_dir = output_dir / "input"
            input_dir.mkdir(parents=True)

            with self.assertRaises(ConverterError):
                run_convert_folder(ConverterConfig(input_dir=input_dir, output_dir=output_dir))

            self.assertFalse((output_dir / "runs").exists())


if __name__ == "__main__":
    unittest.main()


def _file_uri_to_path(target: str) -> Path:
    parsed = urlparse(target)
    if parsed.scheme != "file":
        return Path(target).resolve()
    raw_path = unquote(parsed.path)
    if raw_path.startswith("/") and len(raw_path) >= 3 and raw_path[2] == ":":
        raw_path = raw_path[1:]
    return Path(raw_path).resolve()


def _minimal_formula_document_payload(
    *,
    first_text: str = "R = A + B",
    first_calc_expr: str = "R = A + B",
    second_text: str = "S = R * 2",
    second_calc_expr: str = "S = R * 2",
) -> dict[str, object]:
    document_id = "sha256:" + ("0" * 64)
    return {
        "schema_version": "document.v1",
        "document_id": document_id,
        "source": {
            "original_path": "input.docx",
            "relative_input_path": None,
            "filename": "input.docx",
            "format": "docx",
            "sha256": "0" * 64,
            "size_bytes": 0,
        },
        "processing": {
            "route": "docx_native",
            "status": "success",
            "ocr_applied": False,
            "warnings": [],
        },
        "metadata": {
            "title": "Тестовый документ",
            "document_type": "методика",
            "short_summary": "Тестовый документ с формулами",
            "confidence": "high",
            "method": "filename_fallback",
        },
        "units": [
            _minimal_formula_unit(
                document_id=document_id,
                unit_id="u_000001",
                order=1,
                text=first_text,
                calc_expr=first_calc_expr,
            ),
            _minimal_formula_unit(
                document_id=document_id,
                unit_id="u_000002",
                order=2,
                text=second_text,
                calc_expr=second_calc_expr,
            ),
        ],
        "assets": [],
        "quality": {"flags": [], "warnings": []},
    }


def _minimal_formula_unit(
    *,
    document_id: str,
    unit_id: str,
    order: int,
    text: str,
    calc_expr: str,
) -> dict[str, object]:
    return {
        "unit_id": unit_id,
        "parent_id": None,
        "type": "formula",
        "order": order,
        "text": text,
        "asset_ref": None,
        "formula": {
            "source_format": "docx_text_linearized",
            "linear_text": text,
            "display_latex": text,
            "calc_expr": calc_expr,
            "variables": {},
            "confidence": "low",
            "warnings": [],
        },
        "cell": None,
        "source_ref": {
            "document_id": document_id,
            "page": None,
            "bbox": None,
            "docx_path": "/w:document/w:body/w:p[1]",
            "coordinate_system": None,
            "page_width": None,
            "page_height": None,
        },
        "quality": {"flags": [], "warnings": []},
    }


def _load_cli_result(buffer: StringIO) -> dict[str, Any]:
    payload = cast(dict[str, Any], json.loads(buffer.getvalue()))
    validate_payload(payload, "cli-result.v1.schema.json")
    return payload


def _ready_ocr_payload() -> dict[str, Any]:
    payload: dict[str, Any] = {
        "schema_version": "ocr-runtime.v1",
        "status": "ready",
        "tools": {
            "ocrmypdf": {"available": True, "path": "C:/ocrmypdf.exe"},
            "tesseract": {"available": True, "path": "C:/tesseract.exe"},
            "ghostscript": {"available": True, "path": "C:/gswin64c.exe"},
        },
        "languages": {
            "requested": ["rus", "eng"],
            "installed": ["rus", "eng"],
            "missing": [],
        },
        "warnings": [],
    }
    validate_payload(payload, "ocr-runtime.v1.schema.json")
    return payload


def _missing_ocr_payload() -> dict[str, Any]:
    payload: dict[str, Any] = {
        "schema_version": "ocr-runtime.v1",
        "status": "missing",
        "tools": {
            "ocrmypdf": {"available": False, "path": None},
            "tesseract": {"available": True, "path": "C:/tesseract.exe"},
            "ghostscript": {"available": True, "path": "C:/gswin64c.exe"},
        },
        "languages": {
            "requested": ["rus", "eng"],
            "installed": ["rus", "eng"],
            "missing": [],
        },
        "warnings": ["OCRmyPDF CLI is not available in the active runtime."],
    }
    validate_payload(payload, "ocr-runtime.v1.schema.json")
    return payload


def _ready_font_bundle_payload() -> dict[str, Any]:
    return {
        "status": "ready",
        "directory": "D:/converter/assets/fonts",
        "fonts": ["LiberationSerif-Regular.ttf"],
        "warnings": [],
    }


def _ready_schema_checks_payload() -> dict[str, Any]:
    return {
        "status": "ready",
        "directory": "D:/converter/schemas",
        "schemas": [{"name": "cli-result.v1.schema.json", "available": True}],
        "warnings": [],
    }