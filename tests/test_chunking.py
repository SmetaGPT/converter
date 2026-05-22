from __future__ import annotations

import shutil
import tempfile
import unittest
from pathlib import Path
from typing import Any

from docx import Document

from doc_converter.chunking import build_chunks_from_document
from doc_converter.config import ConverterConfig
from doc_converter.runner import run_convert_folder
from doc_converter.schema_validation import validate_json_file, validate_payload


class ChunkingTests(unittest.TestCase):
    def test_build_chunks_from_document_produces_schema_valid_records(self) -> None:
        payload: dict[str, Any] = {
            "schema_version": "document.v1",
            "document_id": "sha256:" + "1" * 64,
            "source": {
                "original_path": "C:/tmp/sample.docx",
                "relative_input_path": "sample.docx",
                "filename": "sample.docx",
                "format": "docx",
                "sha256": "1" * 64,
                "size_bytes": 100,
            },
            "processing": {"route": "docx_native", "status": "success", "warnings": []},
            "units": [
                {
                    "unit_id": "u_000000",
                    "parent_id": None,
                    "type": "document",
                    "order": 0,
                    "text": None,
                    "asset_ref": None,
                    "source_ref": {"document_id": "sha256:" + "1" * 64, "page": None, "bbox": None, "docx_path": None, "coordinate_system": None, "page_width": None, "page_height": None},
                    "quality": {"flags": [], "warnings": []}
                },
                {
                    "unit_id": "u_000001",
                    "parent_id": "u_000000",
                    "type": "section",
                    "order": 1,
                    "text": "Раздел 1",
                    "asset_ref": None,
                    "source_ref": {"document_id": "sha256:" + "1" * 64, "page": None, "bbox": None, "docx_path": "/doc/p[1]", "coordinate_system": None, "page_width": None, "page_height": None},
                    "quality": {"flags": [], "warnings": []}
                },
                {
                    "unit_id": "u_000002",
                    "parent_id": "u_000000",
                    "type": "paragraph",
                    "order": 2,
                    "text": "Первый абзац",
                    "asset_ref": None,
                    "source_ref": {"document_id": "sha256:" + "1" * 64, "page": None, "bbox": None, "docx_path": "/doc/p[2]", "coordinate_system": None, "page_width": None, "page_height": None},
                    "quality": {"flags": [], "warnings": []}
                }
            ],
            "assets": [],
            "quality": {"flags": [], "warnings": []}
        }

        chunks = build_chunks_from_document(payload, max_chars=50)
        self.assertEqual(len(chunks), 1)
        self.assertEqual(chunks[0]["section_path"], ["Раздел 1"])
        validate_payload(chunks[0], "chunks.v1.schema.json")

    def test_copied_run_package_remains_chunkable(self) -> None:
        with tempfile.TemporaryDirectory() as input_dir, tempfile.TemporaryDirectory() as output_dir, tempfile.TemporaryDirectory() as copy_root:
            source_path = Path(input_dir) / "sample.docx"
            document = Document()
            document.add_heading("Раздел", level=1)
            document.add_paragraph("Переносимый текст")
            document.save(str(source_path))

            result = run_convert_folder(ConverterConfig(input_dir=Path(input_dir), output_dir=Path(output_dir)))
            copied_run_dir = Path(copy_root) / result.run_dir.name
            shutil.copytree(result.run_dir, copied_run_dir)

            document_path = next((copied_run_dir / "documents").glob("*/document.v1.json"))
            payload = validate_json_file(document_path, "document.v1.schema.json")
            chunks = build_chunks_from_document(payload)

            self.assertTrue(chunks)
            for chunk in chunks:
                validate_payload(chunk, "chunks.v1.schema.json")


if __name__ == "__main__":
    unittest.main()